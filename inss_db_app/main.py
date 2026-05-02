from __future__ import annotations

from contextlib import asynccontextmanager
import csv
import io
import json
import hashlib
import hmac
import mimetypes
from pathlib import Path
import threading
import traceback
from types import SimpleNamespace
from urllib.parse import urlencode
from urllib.parse import quote_plus
from datetime import UTC, datetime

from fastapi import Body, Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session
from openpyxl import load_workbook

from .auth import hash_password, verify_password
from .c6_worker_loan import C6WorkerLoanError, c6_worker_loan_client
from .config import settings
from .database import Base, SessionLocal, engine, ensure_runtime_indexes, validate_runtime_schema
from .facta import FactaError, build_facta_client
from .importer import (
    backfill_missing_addresses,
    FilterSet,
    count_client_results,
    create_import_batch,
    delete_batch,
    export_clients,
    export_cpfs_for_enrichment,
    export_updated_clients_from_latest_enrichment,
    format_money,
    import_public_campaign_updates,
    import_phone_enrichments,
    normalize_base_segment,
    normalize_cpf,
    parse_decimal,
    process_import_batch,
    query_clients,
    save_uploaded_files,
)
from .models import AppUser, AuditLog, Client, ClientOccurrence, ExportHistory, ExportJob, ImportBatch, PhoneEnrichment, PublicCampaignUpdate, SavedFilter, SourceFile
from .schemas import ImportRequest


BASE_DIR = Path(__file__).resolve().parent
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))
AUTH_COOKIE_NAME = "inss_auth"
PERMISSION_LABELS = {
    "can_view_dashboard": "Dashboard",
    "can_import": "Importacao",
    "can_search": "Consulta",
    "can_export": "Gerar listas",
    "can_view_history": "Historico",
    "can_delete_batches": "Apagar campanhas",
    "can_manage_users": "Gerenciar usuarios",
}
BASE_SEGMENT_OPTIONS = [
    ("INSS", "INSS"),
    ("GOVERNO", "Governo"),
    ("PREFEITURA", "Prefeitura"),
]
CLIENTS_PAGE_SIZE = 25
HISTORY_PAGE_SIZE = 20


def format_reference(value: str) -> str:
    if not value or "-" not in value:
        return value
    year, month = value.split("-", 1)
    return f"{month}/{year}"


def format_phone(value: str) -> str:
    digits = "".join(char for char in (value or "") if char.isdigit())
    if len(digits) == 11:
        return f"({digits[:2]}) {digits[2:7]}-{digits[7:]}"
    if len(digits) == 10:
        return f"({digits[:2]}) {digits[2:6]}-{digits[6:]}"
    return value


def format_dashboard_datetime(value) -> str:
    if value is None:
        return "-"
    try:
        return value.strftime("%d/%m/%Y %H:%M")
    except AttributeError:
        return str(value)


def current_utc():
    return datetime.now(UTC)


def is_future_timestamp(value: datetime | None) -> bool:
    if value is None:
        return False
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    return value > current_utc()


templates.env.filters["brl"] = format_money
templates.env.filters["ref_br"] = format_reference
templates.env.filters["phone_br"] = format_phone


def get_session() -> Session:
    session = SessionLocal()
    try:
        yield session
    finally:
        session.close()


def user_can(user: AppUser | None, permission: str) -> bool:
    if user is None or not user.is_active:
        return False
    if user.is_admin:
        return True
    return bool(getattr(user, permission, False))


def build_permissions(user: AppUser | None) -> dict[str, bool]:
    permissions = {key: user_can(user, key) for key in PERMISSION_LABELS}
    permissions["is_admin"] = bool(user.is_admin) if user else False
    return permissions


def create_auth_token(user_id: int) -> str:
    payload = str(user_id)
    signature = hmac.new(settings.session_secret_key.encode("utf-8"), payload.encode("utf-8"), hashlib.sha256).hexdigest()
    return f"{payload}:{signature}"


def parse_auth_token(token: str) -> int | None:
    if not token or ":" not in token:
        return None
    payload, signature = token.split(":", 1)
    expected = hmac.new(settings.session_secret_key.encode("utf-8"), payload.encode("utf-8"), hashlib.sha256).hexdigest()
    if not hmac.compare_digest(signature, expected):
        return None
    try:
        return int(payload)
    except ValueError:
        return None


def is_safe_internal_path(target: str) -> bool:
    value = (target or "").strip()
    return value.startswith("/") and not value.startswith("//")


def safe_next_path(target: str, default: str = "/") -> str:
    return target if is_safe_internal_path(target) else default


def build_filter_query_params(filters: FilterSet, base_segment: str, selected_matricula: str = "", selected_servico: str = "", page: int | None = None) -> dict[str, str]:
    payload = {
        "base_segment": base_segment,
        "base_source": filters.base_source,
        "quick_mode": filters.quick_mode,
        "quick_value": filters.quick_value,
        "year": str(filters.year or ""),
        "month": str(filters.month or ""),
        "week": filters.week,
        "uf": filters.uf,
        "city": filters.city,
        "esp": filters.esp,
        "margin_min": str(filters.margin_min or ""),
        "margin_max": str(filters.margin_max or ""),
        "age_min": str(filters.age_min or ""),
        "age_max": str(filters.age_max or ""),
        "has_phone": "sim" if filters.has_phone is True else "nao" if filters.has_phone is False else "",
        "source_file": filters.source_file,
        "cpf": filters.cpf,
        "phone": filters.phone,
        "ddd": filters.ddd,
        "name": filters.name,
        "selected_matricula": selected_matricula,
        "selected_servico": selected_servico,
    }
    if page is not None:
        payload["page"] = str(page)
    return payload


def build_pagination(total_items: int, page: int, page_size: int) -> dict[str, int | bool]:
    total_pages = max((total_items + page_size - 1) // page_size, 1)
    safe_page = min(max(page, 1), total_pages)
    return {
        "page": safe_page,
        "page_size": page_size,
        "total_items": total_items,
        "total_pages": total_pages,
        "offset": (safe_page - 1) * page_size,
        "has_prev": safe_page > 1,
        "has_next": safe_page < total_pages,
        "prev_page": max(safe_page - 1, 1),
        "next_page": min(safe_page + 1, total_pages),
    }


def log_audit(session: Session, actor_username: str, action: str, target_type: str = "", target_id: str = "", message: str = "", metadata: dict[str, object] | None = None) -> None:
    session.add(
        AuditLog(
            actor_username=actor_username or "sistema",
            action=action,
            target_type=target_type,
            target_id=str(target_id or ""),
            message=message,
            metadata_json=json.dumps(metadata or {}, ensure_ascii=False, default=str),
        )
    )


def log_audit_with_new_session(actor_username: str, action: str, target_type: str = "", target_id: str = "", message: str = "", metadata: dict[str, object] | None = None) -> None:
    session = SessionLocal()
    try:
        log_audit(session, actor_username, action, target_type=target_type, target_id=target_id, message=message, metadata=metadata)
        session.commit()
    finally:
        session.close()


def serialize_filters(filters: FilterSet) -> str:
    return json.dumps(filters.__dict__, ensure_ascii=False, default=str)


def deserialize_filters(payload: str) -> FilterSet:
    try:
        raw = json.loads(payload or "{}")
    except json.JSONDecodeError:
        raw = {}
    return FilterSet(**raw)


def summarize_export_filters(filters: FilterSet) -> str:
    parts: list[str] = []
    if filters.base_segment:
        parts.append(f"segmento={filters.base_segment}")
    if filters.year:
        parts.append(f"ano={filters.year}")
    if filters.month:
        parts.append(f"mes={filters.month}")
    if filters.uf:
        parts.append(f"uf={filters.uf.upper()}")
    if filters.city:
        parts.append(f"cidade={filters.city}")
    if filters.ddd:
        parts.append(f"ddd={filters.ddd}")
    if filters.base_source:
        parts.append(f"fonte={filters.base_source}")
    if filters.campaign_profile:
        parts.append(f"perfil={filters.campaign_profile}")
    if filters.has_phone is True:
        parts.append("com_telefone=sim")
    if filters.has_phone is False:
        parts.append("com_telefone=nao")
    return ", ".join(parts) if parts else "sem filtros"


def enqueue_export_job(
    session: Session,
    requested_by: str,
    job_type: str,
    filters: FilterSet | None = None,
    file_format: str = "xlsx",
    include_audit: bool = False,
    message: str = "Exportacao colocada em fila.",
    metadata: dict[str, object] | None = None,
) -> ExportJob:
    job = ExportJob(
        requested_by=requested_by,
        job_type=job_type,
        status="PENDENTE",
        filters_json=serialize_filters(filters) if filters is not None else "{}",
        include_audit=include_audit,
        file_format=file_format,
    )
    session.add(job)
    session.flush()
    log_audit(
        session,
        requested_by,
        "export_job_created",
        target_type="export_job",
        target_id=str(job.id),
        message=message,
        metadata=metadata or {},
    )
    session.commit()
    threading.Thread(target=run_export_job, args=(job.id,), daemon=True).start()
    return job


def run_export_job(job_id: int) -> None:
    session = SessionLocal()
    try:
        job = session.get(ExportJob, job_id)
        if job is None:
            return
        job.status = "PROCESSANDO"
        job.started_at = current_utc()
        session.commit()

        filters = deserialize_filters(job.filters_json)
        if job.job_type == "campaign":
            total_rows = count_client_results(session, filters)
            if total_rows <= 0:
                raise ValueError(f"Nenhum registro encontrado para os filtros informados ({summarize_export_filters(filters)}).")
            export_path = export_clients(session, filters, include_audit=job.include_audit, file_format=job.file_format)
        elif job.job_type == "cpf_enrichment":
            total_rows = count_client_results(session, filters)
            if total_rows <= 0:
                raise ValueError(f"Nenhum CPF encontrado para os filtros informados ({summarize_export_filters(filters)}).")
            export_path = export_cpfs_for_enrichment(session, filters, file_format=job.file_format)
        elif job.job_type == "updated_base":
            export_path = export_updated_clients_from_latest_enrichment(
                session,
                base_segment=filters.base_segment or "INSS",
                file_format=job.file_format,
            )
        else:
            raise ValueError(f"Tipo de job de exportacao nao suportado: {job.job_type}")

        job.output_file = str(export_path)
        job.total_records = session.scalar(select(ExportHistory.total_registros).where(ExportHistory.arquivo_saida == str(export_path)).order_by(ExportHistory.id.desc()).limit(1)) or 0
        job.status = "CONCLUIDO"
        job.finished_at = current_utc()
        log_audit(
            session,
            job.requested_by,
            "export_job_completed",
            target_type="export_job",
            target_id=str(job.id),
            message="Exportacao assincrona concluida.",
            metadata={"arquivo": export_path.name, "registros": job.total_records, "job_type": job.job_type},
        )
        session.commit()
    except Exception as exc:
        session.rollback()
        failed_job = session.get(ExportJob, job_id)
        if failed_job is not None:
            failed_job.status = "ERRO"
            failed_job.error_message = str(exc)
            failed_job.finished_at = current_utc()
            log_audit(session, failed_job.requested_by, "export_job_failed", target_type="export_job", target_id=str(failed_job.id), message="Falha na exportacao assincrona.", metadata={"erro": str(exc)})
            session.commit()
    finally:
        session.close()


def render_template(request: Request, template_name: str, context: dict[str, object], status_code: int = 200) -> HTMLResponse:
    payload = {
        "current_user": getattr(request.state, "current_user", None),
        "permissions": build_permissions(getattr(request.state, "current_user", None)),
        "base_segment_options": BASE_SEGMENT_OPTIONS,
    }
    payload.update(context)
    return templates.TemplateResponse(request, template_name, payload, status_code=status_code)


def base_segment_label(value: str) -> str:
    normalized = normalize_base_segment(value)
    return dict(BASE_SEGMENT_OPTIONS).get(normalized, normalized)


def _norm_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().replace("\ufeff", "")


def _parse_extras_json(payload: str) -> dict[str, object]:
    if not payload:
        return {}
    try:
        data = json.loads(payload)
    except json.JSONDecodeError:
        return {}
    return data if isinstance(data, dict) else {}


def _first_extra_value(extras: dict[str, object], *keys: str) -> str:
    for key in keys:
        value = extras.get(key)
        if _norm_text(value):
            return _norm_text(value)
    return ""


def _build_public_filter_options(session: Session, base_segment: str, base_source: str) -> dict[str, list[str]]:
    options = {
        "entidades": [],
        "convenios": [],
        "servicos": [],
        "situacoes": [],
        "cargos": [],
        "tipos_vinculo": [],
    }
    if base_segment not in {"GOVERNO", "PREFEITURA"} or not _norm_text(base_source):
        return options

    rows = session.execute(
        select(
            ClientOccurrence.entidade,
            ClientOccurrence.secretaria,
            ClientOccurrence.convenio,
            ClientOccurrence.servico,
            ClientOccurrence.situacao,
            ClientOccurrence.cbo_titulo,
            ClientOccurrence.extras_json,
        ).where(
            ClientOccurrence.base_segment == base_segment,
            ClientOccurrence.base_source == base_source,
        )
    ).all()

    entidades: set[str] = set()
    convenios: set[str] = set()
    servicos: set[str] = set()
    situacoes: set[str] = set()
    cargos: set[str] = set()
    tipos_vinculo: set[str] = set()

    for entidade, secretaria, convenio, servico, situacao, cbo_titulo, extras_json in rows:
        if _norm_text(entidade):
            entidades.add(_norm_text(entidade))
        if _norm_text(secretaria):
            entidades.add(_norm_text(secretaria))
        if _norm_text(convenio):
            convenios.add(_norm_text(convenio))
        if _norm_text(servico):
            servicos.add(_norm_text(servico))
        if _norm_text(situacao):
            situacoes.add(_norm_text(situacao))
        if _norm_text(cbo_titulo):
            cargos.add(_norm_text(cbo_titulo))

        extras = _parse_extras_json(extras_json or "")
        entidade_extra = _first_extra_value(
            extras,
            "ENTIDADE",
            "ORGAO",
            "ORGÃO",
            "SECRETARIA",
            "DESCRIÇAO CNPJ",
            "DESCRIÇÃO CNPJ",
            "DESCRICAO CNPJ",
        )
        convenio_extra = _first_extra_value(extras, "CONVENIO", "CONVÊNIO")
        servico_extra = _first_extra_value(
            extras,
            "SERVICO",
            "SERVIÇO",
            "SERVICO (SERVIDOR)",
            "SERVIÇO (SERVIDOR)",
            "TIPO SERVICO (SERVIDOR)",
            "TIPO SERVIÇO (SERVIDOR)",
        )
        situacao_extra = _first_extra_value(extras, "SITUACAO", "SITUAÇÃO", "STATUS")
        cargo_extra = _first_extra_value(extras, "CARGO", "FUNCAO", "FUNÇÃO", "FUNÃƒÆ’Ã¢â‚¬Â¡ÃƒÆ’Ã†â€™O", "CARGO/FUNCAO", "CARGO FUNCAO", "CBO TITULO", "CBO_TITULO")
        tipo_vinculo = _first_extra_value(extras, "REGIME_CONTRATACAO", "TIPO DE VINCULO", "TIPO_VINCULO", "VINCULO", "REGIME DE CONTRATACAO")
        if entidade_extra:
            entidades.add(entidade_extra)
        if convenio_extra:
            convenios.add(convenio_extra)
        if servico_extra:
            servicos.add(servico_extra)
        if situacao_extra:
            situacoes.add(situacao_extra)
        if cargo_extra:
            cargos.add(cargo_extra)
        if tipo_vinculo:
            tipos_vinculo.add(tipo_vinculo)

    options["entidades"] = sorted(entidades)
    options["convenios"] = sorted(convenios)
    options["servicos"] = sorted(servicos)
    options["situacoes"] = sorted(situacoes)
    options["cargos"] = sorted(cargos)
    options["tipos_vinculo"] = sorted(tipos_vinculo)
    return options


def ensure_default_admin() -> None:
    session = SessionLocal()
    try:
        existing_user = session.scalar(select(func.count()).select_from(AppUser)) or 0
        if existing_user:
            return
        password_hash, password_salt = hash_password(settings.default_admin_password)
        session.add(
            AppUser(
                username=settings.default_admin_username,
                full_name="Administrador",
                password_hash=password_hash,
                password_salt=password_salt,
                is_active=True,
                is_admin=True,
                can_view_dashboard=True,
                can_import=True,
                can_search=True,
                can_export=True,
                can_view_history=True,
                can_delete_batches=True,
                can_manage_users=True,
            )
        )
        session.commit()
    finally:
        session.close()


def run_import_job(batch_id: int, request_data: ImportRequest) -> None:
    session = SessionLocal()
    try:
        batch = session.get(ImportBatch, batch_id)
        if batch is None:
            return
        summary = process_import_batch(session, batch, request_data)
        log_audit_with_new_session(
            request_data.user_name,
            "import_batch_completed",
            target_type="batch",
            target_id=str(batch_id),
            message="Importacao concluida.",
            metadata={"rows_imported": summary.rows_imported, "files_imported": summary.files_imported},
        )
    except Exception as exc:
        traceback.print_exc()
        log_audit_with_new_session(
            request_data.user_name,
            "import_batch_failed",
            target_type="batch",
            target_id=str(batch_id),
            message="Falha na importacao.",
            metadata={"erro": str(exc), "traceback": traceback.format_exc(limit=20)},
        )
    finally:
        session.close()


def create_schema() -> None:
    Base.metadata.create_all(bind=engine)
    ensure_runtime_indexes()
    ensure_default_admin()


def run_startup_address_backfill() -> None:
    session = SessionLocal()
    try:
        summary = backfill_missing_addresses(session)
        print(
            "Startup address backfill:",
            f"scanned={summary.scanned_rows}",
            f"updated={summary.updated_rows}",
            f"rebuilt_clients={summary.rebuilt_clients}",
            f"skipped_invalid_cep={summary.skipped_invalid_cep}",
        )
    except Exception as exc:
        print(f"Startup address backfill failed: {exc}")
    finally:
        session.close()


def get_authenticated_user(request: Request, session: Session) -> AppUser:
    user_id = parse_auth_token(request.cookies.get(AUTH_COOKIE_NAME, ""))
    if not user_id:
        raise HTTPException(status_code=401)
    user = session.get(AppUser, user_id)
    if user is None or not user.is_active:
        raise HTTPException(status_code=401)
    if user.locked_until and user.locked_until > current_utc():
        raise HTTPException(status_code=403)
    return user


def require_permission(permission: str):
    def dependency(request: Request, session: Session = Depends(get_session)) -> AppUser:
        user = get_authenticated_user(request, session)
        if not user_can(user, permission):
            raise HTTPException(status_code=403)
        return user

    return dependency


@asynccontextmanager
async def lifespan(_: FastAPI):
    if settings.enable_startup_schema_maintenance:
        create_schema()
        print("Startup maintenance: schema ready")
    else:
        validate_runtime_schema()
        ensure_default_admin()
        print("Startup maintenance: schema validated")
    if settings.enable_startup_address_backfill:
        threading.Thread(target=run_startup_address_backfill, daemon=True).start()
    yield


app = FastAPI(title="Banco de Clientes INSS", lifespan=lifespan)
app.mount("/static", StaticFiles(directory=str(BASE_DIR / "static")), name="static")


@app.middleware("http")
async def attach_current_user(request: Request, call_next):
    request.state.current_user = None
    user_id = parse_auth_token(request.cookies.get(AUTH_COOKIE_NAME, ""))
    if user_id:
        session = SessionLocal()
        try:
            user = session.get(AppUser, user_id)
            if user is not None and user.is_active:
                request.state.current_user = user
        finally:
            session.close()
    response = await call_next(request)
    return response


@app.exception_handler(401)
async def unauthorized_handler(request: Request, _exc: HTTPException):
    target = quote_plus(str(request.url.path))
    return RedirectResponse(url=f"/login?next={target}", status_code=303)


@app.exception_handler(403)
async def forbidden_handler(request: Request, _exc: HTTPException):
    return render_template(
        request,
        "error.html",
        {"title": "Acesso negado", "message": "Voce nao tem permissao para acessar esta funcionalidade."},
        status_code=403,
    )


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, next: str = "/") -> HTMLResponse:
    if getattr(request.state, "current_user", None):
        return RedirectResponse(url=safe_next_path(next), status_code=303)
    return render_template(request, "login.html", {"error": "", "next": safe_next_path(next), "title": "Login"})


@app.post("/login", response_class=HTMLResponse, response_model=None)
def login_submit(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    next: str = Form("/"),
    session: Session = Depends(get_session),
) -> HTMLResponse:
    user = session.scalar(select(AppUser).where(AppUser.username == username.strip()))
    redirect_target = safe_next_path(next)
    if user is not None and is_future_timestamp(user.locked_until):
        return render_template(
            request,
            "login.html",
            {"error": "Usuario temporariamente bloqueado.", "next": redirect_target, "title": "Login"},
            status_code=403,
        )
    if user is None or not user.is_active or not verify_password(password, user.password_hash, user.password_salt):
        return render_template(
            request,
            "login.html",
            {"error": "Usuario ou senha invalidos.", "next": redirect_target, "title": "Login"},
            status_code=400,
        )
    user.last_login_at = current_utc()
    log_audit(session, user.username, "login_success", target_type="user", target_id=str(user.id), message="Login efetuado com sucesso.")
    session.commit()
    response = RedirectResponse(url=redirect_target, status_code=303)
    response.set_cookie(AUTH_COOKIE_NAME, create_auth_token(user.id), httponly=True, samesite="lax")
    return response


@app.get("/logout")
def logout() -> RedirectResponse:
    response = RedirectResponse(url="/login", status_code=303)
    response.delete_cookie(AUTH_COOKIE_NAME)
    return response


def _extract_cpfs_from_sheet(content: bytes, filename: str, cpf_column: str = "cpf") -> list[str]:
    name = (filename or "").lower()
    values: list[str] = []
    if name.endswith(".csv"):
        text = content.decode("utf-8", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            raw = str(row.get(cpf_column, "")).strip()
            if raw:
                values.append(raw)
    elif name.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not header:
            return []
        index_by_name = {str(col).strip().lower(): idx for idx, col in enumerate(header) if col is not None}
        idx = index_by_name.get(cpf_column.lower())
        if idx is None:
            return []
        for row in ws.iter_rows(min_row=2, values_only=True):
            raw = row[idx] if idx < len(row) else None
            if raw is not None:
                values.append(str(raw).strip())
    return values


@app.get("/clt", response_class=HTMLResponse)
def clt_page(
    request: Request,
    _current_user: AppUser = Depends(require_permission("can_search")),
) -> HTMLResponse:
    return render_template(request, "clt.html", {"title": "Consulta CLT FACTA", "result": None, "error": ""})


@app.post("/clt/consulta-individual", response_class=HTMLResponse)
def clt_consulta_individual(
    request: Request,
    cpf: str = Form(...),
    nome: str = Form(""),
    celular: str = Form(""),
    tipo_envio: str = Form("WHATSAPP"),
    _current_user: AppUser = Depends(require_permission("can_search")),
) -> HTMLResponse:
    cpf_limpo = "".join(ch for ch in cpf if ch.isdigit())
    if len(cpf_limpo) != 11:
        return render_template(request, "clt.html", {"title": "Consulta CLT FACTA", "result": None, "error": "CPF invalido. Informe 11 digitos."}, status_code=400)
    try:
        client = build_facta_client()
        offline = client.consulta_offline(cpf_limpo)
        mensagem_offline = str(offline.get("mensagem") or "")
        if "Dados retornados com sucesso" in mensagem_offline:
            return render_template(request, "clt.html", {"title": "Consulta CLT FACTA", "result": {"cpf": cpf_limpo, "origem": "offline", "offline": offline}, "error": ""})

        if not nome or not celular:
            return render_template(
                request,
                "clt.html",
                {
                    "title": "Consulta CLT FACTA",
                    "result": {"cpf": cpf_limpo, "origem": "sem_dado_offline", "offline": offline, "proximo_passo": "Informe nome e celular para solicitar autorizacao online."},
                    "error": "",
                },
            )

        autorizacao = client.solicita_autorizacao(cpf=cpf_limpo, nome=nome, celular=celular, tipo_envio=tipo_envio)
        consulta_online = client.consulta_online(cpf_limpo)
        ofertas = client.consulta_ofertas(cpf_limpo)
        return render_template(
            request,
            "clt.html",
            {"title": "Consulta CLT FACTA", "result": {"cpf": cpf_limpo, "origem": "online", "autorizacao": autorizacao, "consulta_online": consulta_online, "ofertas": ofertas}, "error": ""},
        )
    except FactaError as exc:
        return render_template(request, "clt.html", {"title": "Consulta CLT FACTA", "result": None, "error": str(exc)}, status_code=502)


@app.post("/clt/consulta-lote", response_class=JSONResponse)
def clt_consulta_lote(
    cpfs: str = Form(...),
    _current_user: AppUser = Depends(require_permission("can_search")),
) -> JSONResponse:
    try:
        client = build_facta_client()
    except FactaError as exc:
        return JSONResponse({"total": 0, "resultados": [], "erro": str(exc)}, status_code=500)

    linhas = [item.strip() for item in cpfs.replace("\r", "").split("\n") if item.strip()]
    resultados: list[dict[str, str]] = []
    for item in linhas:
        cpf_limpo = "".join(ch for ch in item if ch.isdigit())
        if len(cpf_limpo) != 11:
            resultados.append({"cpf": item, "status": "erro", "mensagem": "CPF invalido"})
            continue
        try:
            offline = client.consulta_offline(cpf_limpo)
            mensagem = str(offline.get("mensagem") or "")
            status = "ok_offline" if "Dados retornados com sucesso" in mensagem else "sem_dado_offline"
            resultados.append({"cpf": cpf_limpo, "status": status, "mensagem": mensagem})
        except FactaError as exc:
            resultados.append({"cpf": cpf_limpo, "status": "erro", "mensagem": str(exc)})

    return JSONResponse({"total": len(linhas), "resultados": resultados})


@app.post("/clt/consulta-lote-planilha", response_class=HTMLResponse)
async def clt_consulta_lote_planilha(
    request: Request,
    file: UploadFile = File(...),
    cpf_column: str = Form("cpf"),
    _current_user: AppUser = Depends(require_permission("can_search")),
) -> HTMLResponse:
    try:
        client = build_facta_client()
    except FactaError as exc:
        return render_template(request, "clt.html", {"title": "Consulta CLT FACTA", "result": None, "error": str(exc)}, status_code=500)

    if not file.filename:
        return render_template(request, "clt.html", {"title": "Consulta CLT FACTA", "result": None, "error": "Arquivo sem nome."}, status_code=400)
    if not (file.filename.lower().endswith(".csv") or file.filename.lower().endswith(".xlsx")):
        return render_template(request, "clt.html", {"title": "Consulta CLT FACTA", "result": None, "error": "Formato invalido. Envie .csv ou .xlsx"}, status_code=400)

    content = await file.read()
    cpfs_raw = _extract_cpfs_from_sheet(content=content, filename=file.filename, cpf_column=cpf_column.strip() or "cpf")
    if not cpfs_raw:
        return render_template(
            request,
            "clt.html",
            {"title": "Consulta CLT FACTA", "result": None, "error": f"Nenhum CPF encontrado na coluna '{cpf_column}'."},
            status_code=400,
        )

    resultados: list[dict[str, str]] = []
    for item in cpfs_raw:
        cpf_limpo = "".join(ch for ch in item if ch.isdigit())
        if len(cpf_limpo) != 11:
            resultados.append({"cpf": item, "status": "erro", "mensagem": "CPF invalido"})
            continue
        try:
            offline = client.consulta_offline(cpf_limpo)
            mensagem = str(offline.get("mensagem") or "")
            status = "ok_offline" if "Dados retornados com sucesso" in mensagem else "sem_dado_offline"
            resultados.append({"cpf": cpf_limpo, "status": status, "mensagem": mensagem})
        except FactaError as exc:
            resultados.append({"cpf": cpf_limpo, "status": "erro", "mensagem": str(exc)})

    total = len(cpfs_raw)
    ok_offline = sum(1 for row in resultados if row["status"] == "ok_offline")
    sem_dado = sum(1 for row in resultados if row["status"] == "sem_dado_offline")
    erros = sum(1 for row in resultados if row["status"] == "erro")
    result = {
        "modo": "lote_planilha",
        "arquivo": file.filename,
        "coluna_cpf": cpf_column,
        "total": total,
        "ok_offline": ok_offline,
        "sem_dado_offline": sem_dado,
        "erros": erros,
        "resultados": resultados,
    }
    return render_template(request, "clt.html", {"title": "Consulta CLT FACTA", "result": result, "error": ""})


@app.get("/", response_class=HTMLResponse)
def home(
    request: Request,
    current_user: AppUser = Depends(require_permission("can_view_dashboard")),
    session: Session = Depends(get_session),
) -> HTMLResponse:
    total_batches = session.scalar(select(func.count()).select_from(ImportBatch)) or 0
    completed_batches = session.scalar(select(func.count()).select_from(ImportBatch).where(ImportBatch.status == "CONCLUIDO")) or 0
    total_files = session.scalar(select(func.count()).select_from(SourceFile)) or 0
    total_occurrences = session.scalar(select(func.count()).select_from(ClientOccurrence)) or 0
    total_clients = session.scalar(select(func.count()).select_from(Client)) or 0
    total_with_phone = session.scalar(select(func.count()).select_from(Client).where(Client.tem_telefone.is_(True))) or 0
    latest_exports = session.execute(select(ExportHistory).order_by(ExportHistory.criado_em.desc()).limit(5)).scalars().all()
    latest_audits = session.execute(select(AuditLog).order_by(AuditLog.created_at.desc()).limit(5)).scalars().all()
    segment_stats = [
        {
            "segment": segment,
            "label": base_segment_label(segment),
            "occurrences": session.scalar(select(func.count()).select_from(ClientOccurrence).where(ClientOccurrence.base_segment == segment)) or 0,
            "clients": session.scalar(select(func.count(func.distinct(ClientOccurrence.cpf))).where(ClientOccurrence.base_segment == segment)) or 0,
            "with_phone": session.scalar(
                select(func.count(func.distinct(ClientOccurrence.cpf))).where(
                    ClientOccurrence.base_segment == segment,
                    or_(ClientOccurrence.telefone1 != "", ClientOccurrence.telefone2 != "", ClientOccurrence.telefone3 != ""),
                )
            ) or 0,
        }
        for segment, _label in BASE_SEGMENT_OPTIONS
    ]
    top_states = session.execute(
        select(Client.uf_atual, func.count())
        .where(Client.uf_atual != "")
        .group_by(Client.uf_atual)
        .order_by(func.count().desc())
        .limit(5)
    ).all()

    if total_clients == 0 and total_occurrences > 0:
        has_phone_filter = or_(
            ClientOccurrence.telefone1 != "",
            ClientOccurrence.telefone2 != "",
            ClientOccurrence.telefone3 != "",
        )
        total_clients = session.scalar(select(func.count(func.distinct(ClientOccurrence.cpf))).select_from(ClientOccurrence)) or 0
        total_with_phone = (
            session.scalar(
                select(func.count(func.distinct(ClientOccurrence.cpf)))
                .select_from(ClientOccurrence)
                .where(has_phone_filter)
            )
            or 0
        )
        top_states = session.execute(
            select(ClientOccurrence.uf, func.count(func.distinct(ClientOccurrence.cpf)))
            .where(ClientOccurrence.uf != "")
            .group_by(ClientOccurrence.uf)
            .order_by(func.count(func.distinct(ClientOccurrence.cpf)).desc())
            .limit(5)
        ).all()

    latest_batches = session.execute(select(ImportBatch).order_by(ImportBatch.data_importacao.desc()).limit(4)).scalars().all()

    phone_ratio = round((total_with_phone / total_clients) * 100, 1) if total_clients else 0
    no_phone_ratio = round(100 - phone_ratio, 1) if total_clients else 0
    total_without_phone = max(total_clients - total_with_phone, 0)
    state_total = sum(item[1] for item in top_states) or 1
    states_chart = [
        {"uf": uf or "N/D", "count": count, "pct": round((count / state_total) * 100, 1)}
        for uf, count in top_states
    ]
    latest_batch = latest_batches[0] if latest_batches else None
    latest_export = latest_exports[0] if latest_exports else None
    top_state = states_chart[0] if states_chart else None
    avg_records_per_file = round(total_occurrences / total_files) if total_files else 0
    avg_clients_per_campaign = round(total_clients / total_batches) if total_batches else 0
    file_per_campaign = round(total_files / total_batches, 1) if total_batches else 0
    return render_template(
        request,
        "home.html",
        {
            "title": "Inicio",
            "current_user": current_user,
            "total_batches": total_batches,
            "total_files": total_files,
            "total_occurrences": total_occurrences,
            "total_clients": total_clients,
            "total_with_phone": total_with_phone,
            "total_without_phone": total_without_phone,
            "phone_ratio": phone_ratio,
            "no_phone_ratio": no_phone_ratio,
            "completed_batches": completed_batches,
            "avg_records_per_file": avg_records_per_file,
            "avg_clients_per_campaign": avg_clients_per_campaign,
            "file_per_campaign": file_per_campaign,
            "latest_batch_label": format_dashboard_datetime(latest_batch.data_importacao) if latest_batch else "-",
            "latest_export_label": format_dashboard_datetime(latest_export.criado_em) if latest_export else "-",
            "top_state": top_state,
            "states_chart": states_chart,
            "latest_batches": latest_batches,
            "latest_exports": latest_exports,
            "latest_audits": latest_audits,
            "segment_stats": segment_stats,
        },
    )


@app.get("/import", response_class=HTMLResponse)
def import_page(request: Request, _current_user: AppUser = Depends(require_permission("can_import"))) -> HTMLResponse:
    session = SessionLocal()
    try:
        available_public_sources = [
            value
            for (value,) in session.execute(
                select(ClientOccurrence.base_source)
                .where(
                    ClientOccurrence.base_segment.in_(["GOVERNO", "PREFEITURA"]),
                    ClientOccurrence.base_source != "",
                )
                .distinct()
                .order_by(ClientOccurrence.base_source.asc())
            ).all()
            if value
        ]
    finally:
        session.close()
    return render_template(
        request,
        "import.html",
        {
            "summary": None,
            "error": request.query_params.get("error", ""),
            "message": request.query_params.get("message", ""),
            "title": "Importacao",
            "selected_base_segment": request.query_params.get("base_segment", "INSS"),
            "available_public_sources": available_public_sources,
        },
    )


@app.post("/import", response_class=HTMLResponse, response_model=None)
async def run_import(
    request: Request,
    year: int = Form(...),
    month: int = Form(...),
    base_segment: str = Form("INSS"),
    base_source: str = Form(""),
    user_name: str = Form("operador"),
    origin_folder: str = Form("upload_web"),
    allowed_species: list[str] = Form(default=[]),
    custom_species: str = Form(""),
    server_folder: str = Form(""),
    files: list[UploadFile] = File(default=[]),
    current_user: AppUser = Depends(require_permission("can_import")),
    session: Session = Depends(get_session),
) -> HTMLResponse:
    normalized_segment = normalize_base_segment(base_segment)
    selected_paths: list[Path] = []
    if server_folder.strip():
        folder = Path(server_folder.strip())
        if not folder.exists() or not folder.is_dir():
            return render_template(request, "import.html", {"summary": None, "error": f"Pasta nao encontrada: {folder}", "title": "Importacao"}, status_code=400)
        selected_paths = [path for path in sorted(folder.iterdir()) if path.suffix.lower() in {".csv", ".xlsx", ".xlsm", ".xls"}]
        origin_folder = str(folder)
    elif files:
        selected_paths = save_uploaded_files([(item.filename or "arquivo.csv", await item.read()) for item in files])
    else:
        return render_template(request, "import.html", {"summary": None, "error": "Envie arquivos ou informe uma pasta do servidor.", "title": "Importacao"}, status_code=400)

    import_request = ImportRequest(
        year=year,
        month=month,
        user_name=current_user.username,
        origin_folder=origin_folder,
        files=selected_paths,
        base_segment=normalized_segment,
        base_source=base_source,
        allowed_species=[*allowed_species, custom_species] if normalized_segment == "INSS" else [],
    )
    batch = create_import_batch(session, import_request)
    log_audit(
        session,
        current_user.username,
        "import_batch_created",
        target_type="batch",
        target_id=str(batch.id),
        message="Importacao iniciada.",
        metadata={"segmento": normalized_segment, "fonte": base_source, "arquivos": [path.name for path in selected_paths]},
    )
    session.commit()
    threading.Thread(target=run_import_job, args=(batch.id, import_request), daemon=True).start()
    return RedirectResponse(url=f"/history/{batch.id}", status_code=303)


@app.post("/import/public-update-import")
async def run_public_update_import_from_import_page(
    base_segment: str = Form("GOVERNO"),
    source_name: str = Form(""),
    files: list[UploadFile] = File(default=[]),
    current_user: AppUser = Depends(require_permission("can_import")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    normalized_segment = normalize_base_segment(base_segment, default="GOVERNO")
    if normalized_segment not in {"GOVERNO", "PREFEITURA"}:
        return RedirectResponse(url="/import?error=" + quote_plus("A atualizacao complementar e exclusiva para GOV/PREF."), status_code=303)

    uploaded_files: list[tuple[str, bytes]] = []
    for item in files:
        if item.filename:
            uploaded_files.append((item.filename, await item.read()))
    if not uploaded_files:
        return RedirectResponse(
            url=f"/import?base_segment={normalized_segment}&error=" + quote_plus("Selecione ao menos um arquivo complementar."),
            status_code=303,
        )

    saved_paths = save_uploaded_files(uploaded_files)
    summary = import_public_campaign_updates(
        session,
        saved_paths,
        base_segment=normalized_segment,
        source_name=source_name,
        imported_by=current_user.username,
    )
    log_audit(
        session,
        current_user.username,
        "public_update_imported",
        target_type="public_update",
        target_id=normalized_segment,
        message="Atualizacao complementar GOV/PREF importada.",
        metadata={"segmento": normalized_segment, "fonte": source_name, "arquivos": [path.name for path in saved_paths], "clientes_atualizados": summary.clients_updated},
    )
    session.commit()
    message = (
        f"Atualizacao GOV/PREF importada: {summary.rows_imported} linha(s), "
        f"{summary.contracts_saved} contrato(s) gravado(s), "
        f"{summary.clients_updated} cliente(s) atualizados."
    )
    return RedirectResponse(
        url=f"/import?base_segment={normalized_segment}&message=" + quote_plus(message),
        status_code=303,
    )


@app.get("/clients", response_class=HTMLResponse)
def clients_page(
    request: Request,
    base_segment: str = "INSS",
    base_source: str = "",
    selected_matricula: str = "",
    selected_servico: str = "",
    page: int = 1,
    quick_mode: str = "cpf",
    quick_value: str = "",
    year: str = "",
    month: str = "",
    week: str = "",
    uf: str = "",
    city: str = "",
    esp: str = "",
    margin_min: str = "",
    margin_max: str = "",
    age_min: str = "",
    age_max: str = "",
    has_phone: str = "",
    source_file: str = "",
    cpf: str = "",
    phone: str = "",
    ddd: str = "",
    name: str = "",
    current_user: AppUser = Depends(require_permission("can_search")),
    session: Session = Depends(get_session),
) -> HTMLResponse:
    def has_value(value: object) -> bool:
        return bool(str(value or "").strip())

    normalized_segment = normalize_base_segment(base_segment)
    filters = FilterSet(
        base_segment=normalized_segment,
        quick_mode="telefone" if quick_mode == "telefone" else "cpf",
        quick_value=quick_value,
        year=int(year) if year else None,
        month=int(month) if month else None,
        week=week,
        uf=uf,
        city=city,
        esp=esp,
        margin_min=parse_decimal(margin_min),
        margin_max=parse_decimal(margin_max),
        age_min=int(age_min) if age_min else None,
        age_max=int(age_max) if age_max else None,
        has_phone=True if has_phone == "sim" else False if has_phone == "nao" else None,
        source_file=source_file,
        base_source=base_source,
        cpf=cpf,
        phone=phone,
        ddd=ddd,
        name=name,
    )
    total_results = count_client_results(session, filters)
    pagination = build_pagination(total_results, page, CLIENTS_PAGE_SIZE)
    results = query_clients(session, filters, limit=CLIENTS_PAGE_SIZE, offset=int(pagination["offset"]))
    has_active_lookup = any(
        [
            bool(filters.quick_value.strip()),
            filters.year is not None,
            filters.month is not None,
            bool(filters.week.strip()),
            bool(filters.uf.strip()),
            bool(filters.city.strip()),
            bool(filters.esp.strip()),
            filters.margin_min is not None,
            filters.margin_max is not None,
            filters.age_min is not None,
            filters.age_max is not None,
            filters.has_phone is not None,
            bool(filters.source_file.strip()),
            bool(filters.base_source.strip()),
            bool(filters.cpf.strip()),
            bool(filters.phone.strip()),
            bool(filters.ddd.strip()),
            bool(filters.name.strip()),
            bool(selected_matricula.strip()),
            bool(selected_servico.strip()),
        ]
    )
    highlight = results[0] if results and has_active_lookup else None
    public_rows = results if normalized_segment in {"GOVERNO", "PREFEITURA"} else []
    has_public_matriculas = False
    if public_rows:
        rows_with_matricula = [row for row in public_rows if has_value(row.get("matricula"))]
        unique_rows_with_matricula: list[dict[str, object]] = []
        seen_matricula_keys: set[tuple[str, str, str]] = set()
        for row in rows_with_matricula:
            row_key = (
                str(row.get("matricula", "") or "").strip().upper(),
                str(row.get("servico", "") or "").strip().upper(),
                str(row.get("situacao", "") or "").strip().upper(),
            )
            if row_key in seen_matricula_keys:
                continue
            seen_matricula_keys.add(row_key)
            unique_rows_with_matricula.append(row)
        if rows_with_matricula:
            has_public_matriculas = True
            public_rows = unique_rows_with_matricula
            results = unique_rows_with_matricula
            if has_active_lookup:
                highlight = results[0]
    if public_rows and (selected_matricula or selected_servico):
        highlight = next(
            (
                row
                for row in public_rows
                if (not selected_matricula or row.get("matricula") == selected_matricula)
                and (not selected_servico or row.get("servico") == selected_servico)
            ),
            highlight,
        )
    selection_query = urlencode(build_filter_query_params(filters, normalized_segment, page=int(pagination["page"])))
    if public_rows:
        for row in public_rows:
            row["selection_url"] = "/clients?" + selection_query + "&" + urlencode(
                {
                    "selected_matricula": row.get("matricula", "") or "",
                    "selected_servico": row.get("servico", "") or "",
                }
            )
    saved_filters = session.execute(
        select(SavedFilter).where(SavedFilter.user_id == current_user.id, SavedFilter.page_key == "clients").order_by(SavedFilter.created_at.desc()).limit(10)
    ).scalars().all()
    return render_template(
        request,
        "clients.html",
        {
            "title": f"Consulta {base_segment_label(normalized_segment)}",
            "results": results,
            "public_rows": public_rows,
            "highlight": highlight,
            "filters": filters,
            "result_count": total_results,
            "selected_base_segment": normalized_segment,
            "selected_matricula": selected_matricula,
            "selected_servico": selected_servico,
            "selection_query": selection_query,
            "has_public_matriculas": has_public_matriculas,
            "pagination": pagination,
            "saved_filters": saved_filters,
            "favorite_query": urlencode(build_filter_query_params(filters, normalized_segment, selected_matricula=selected_matricula, selected_servico=selected_servico)),
            "current_query_url": str(request.url.path) + (("?" + request.url.query) if request.url.query else ""),
            "message": request.query_params.get("message", ""),
            "error": request.query_params.get("error", ""),
        },
    )


@app.get("/exports", response_class=HTMLResponse)
def exports_page(
    request: Request,
    base_segment: str = "INSS",
    base_source: str = "",
    message: str = "",
    error: str = "",
    current_user: AppUser = Depends(require_permission("can_export")),
    session: Session = Depends(get_session),
) -> HTMLResponse:
    normalized_segment = normalize_base_segment(base_segment)
    normalized_source = base_source.strip()
    exports = session.execute(select(ExportHistory).order_by(ExportHistory.criado_em.desc()).limit(50)).scalars().all()
    export_jobs = (
        session.execute(
            select(ExportJob)
            .where(ExportJob.requested_by == current_user.username)
            .order_by(ExportJob.created_at.desc())
            .limit(20)
        )
        .scalars()
        .all()
    )
    phone_enrichments = session.execute(select(PhoneEnrichment).order_by(PhoneEnrichment.criado_em.desc()).limit(20)).scalars().all()
    available_sources = [
        value
        for (value,) in session.execute(
            select(ClientOccurrence.base_source)
            .where(
                ClientOccurrence.base_segment == normalized_segment,
                ClientOccurrence.base_source != "",
            )
            .distinct()
            .order_by(ClientOccurrence.base_source.asc())
        ).all()
        if value
    ]
    public_filter_options = _build_public_filter_options(session, normalized_segment, normalized_source)
    return render_template(
        request,
        "exports.html",
        {
            "exports": exports,
            "export_jobs": export_jobs,
            "phone_enrichments": phone_enrichments,
            "available_sources": available_sources,
            "public_filter_options": public_filter_options,
            "message": message,
            "error": error,
            "title": f"Listas {base_segment_label(normalized_segment)}",
            "selected_base_segment": normalized_segment,
            "selected_base_source": normalized_source,
        },
    )


@app.post("/exports")
def run_export(
    base_segment: str = Form("INSS"),
    base_source: str = Form(""),
    year: str = Form(""),
    month: str = Form(""),
    ddb_start: str = Form(""),
    ddb_end: str = Form(""),
    uf: str = Form(""),
    city: str = Form(""),
    entidade: list[str] = Form(default=[]),
    convenio: list[str] = Form(default=[]),
    servico_publico: list[str] = Form(default=[]),
    situacao: list[str] = Form(default=[]),
    cargo: list[str] = Form(default=[]),
    tipo_vinculo: list[str] = Form(default=[]),
    salary_min: str = Form(""),
    salary_max: str = Form(""),
    campaign_profile: str = Form(""),
    esp: list[str] = Form(default=[]),
    margin_min: str = Form(""),
    margin_max: str = Form(""),
    age_min: str = Form(""),
    age_max: str = Form(""),
    has_phone: str = Form(""),
    source_file: str = Form(""),
    cpf: str = Form(""),
    ddd: str = Form(""),
    name: str = Form(""),
    include_audit: bool = Form(False),
    file_format: str = Form("xlsx"),
    current_user: AppUser = Depends(require_permission("can_export")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    filters = FilterSet(
        base_segment=normalize_base_segment(base_segment),
        base_source=base_source.strip(),
        year=int(year) if year else None,
        month=int(month) if month else None,
        ddb_start=ddb_start,
        ddb_end=ddb_end,
        uf=uf,
        city=city,
        entidade="|".join(item.strip() for item in entidade if item.strip()),
        convenio="|".join(item.strip() for item in convenio if item.strip()),
        servico_publico="|".join(item.strip() for item in servico_publico if item.strip()),
        situacao="|".join(item.strip() for item in situacao if item.strip()),
        cargo="|".join(item.strip() for item in cargo if item.strip()),
        tipo_vinculo="|".join(item.strip() for item in tipo_vinculo if item.strip()),
        salary_min=parse_decimal(salary_min),
        salary_max=parse_decimal(salary_max),
        campaign_profile=campaign_profile,
        esp=",".join(item for item in esp if item),
        margin_min=parse_decimal(margin_min),
        margin_max=parse_decimal(margin_max),
        age_min=int(age_min) if age_min else None,
        age_max=int(age_max) if age_max else None,
        has_phone=True if has_phone == "sim" else False if has_phone == "nao" else None,
        source_file=source_file,
        cpf=cpf,
        ddd=ddd,
        name=name,
    )
    enqueue_export_job(
        session,
        current_user.username,
        "campaign",
        filters=filters,
        file_format=file_format,
        include_audit=include_audit,
        message="Exportacao colocada em fila.",
        metadata={"segmento": filters.base_segment, "formato": file_format, "auditoria": include_audit},
    )
    return RedirectResponse(url="/exports?message=" + quote_plus("Exportacao colocada em fila. Atualize a pagina para acompanhar."), status_code=303)


@app.post("/exports/cpf-enrichment")
def run_cpf_enrichment_export(
    base_segment: str = Form("INSS"),
    year: str = Form(""),
    month: str = Form(""),
    ddb_start: str = Form(""),
    ddb_end: str = Form(""),
    uf: str = Form(""),
    city: str = Form(""),
    esp: list[str] = Form(default=[]),
    margin_min: str = Form(""),
    margin_max: str = Form(""),
    age_min: str = Form(""),
    age_max: str = Form(""),
    has_phone: str = Form(""),
    source_file: str = Form(""),
    cpf: str = Form(""),
    ddd: str = Form(""),
    name: str = Form(""),
    file_format: str = Form("xlsx"),
    current_user: AppUser = Depends(require_permission("can_export")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    filters = FilterSet(
        base_segment=normalize_base_segment(base_segment),
        year=int(year) if year else None,
        month=int(month) if month else None,
        ddb_start=ddb_start,
        ddb_end=ddb_end,
        uf=uf,
        city=city,
        esp=",".join(item for item in esp if item),
        margin_min=parse_decimal(margin_min),
        margin_max=parse_decimal(margin_max),
        age_min=int(age_min) if age_min else None,
        age_max=int(age_max) if age_max else None,
        has_phone=True if has_phone == "sim" else False if has_phone == "nao" else None,
        source_file=source_file,
        cpf=cpf,
        ddd=ddd,
        name=name,
    )
    enqueue_export_job(
        session,
        current_user.username,
        "cpf_enrichment",
        filters=filters,
        file_format=file_format,
        message="Lista de CPFs colocada em fila.",
        metadata={"segmento": filters.base_segment, "formato": file_format},
    )
    return RedirectResponse(url="/exports?message=" + quote_plus("Lista de CPFs colocada em fila. Atualize a pagina para acompanhar."), status_code=303)


@app.post("/exports/updated-base")
def run_updated_base_export(
    base_segment: str = Form("INSS"),
    file_format: str = Form("xlsx"),
    current_user: AppUser = Depends(require_permission("can_export")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    filters = FilterSet(base_segment=normalize_base_segment(base_segment))
    enqueue_export_job(
        session,
        current_user.username,
        "updated_base",
        filters=filters,
        file_format=file_format,
        message="Base atualizada colocada em fila.",
        metadata={"segmento": filters.base_segment, "formato": file_format},
    )
    return RedirectResponse(url="/exports?message=" + quote_plus("Base atualizada colocada em fila. Atualize a pagina para acompanhar."), status_code=303)


@app.post("/exports/phone-enrichment-import")
async def run_phone_enrichment_import(
    source_name: str = Form("LEMIT"),
    files: list[UploadFile] = File(default=[]),
    current_user: AppUser = Depends(require_permission("can_export")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    uploaded_files: list[tuple[str, bytes]] = []
    for item in files:
        if item.filename:
            uploaded_files.append((item.filename, await item.read()))
    if not uploaded_files:
        return RedirectResponse(url="/exports?error=" + quote_plus("Selecione ao menos um arquivo de retorno."), status_code=303)

    saved_paths = save_uploaded_files(uploaded_files)
    summary = import_phone_enrichments(session, saved_paths, source_name=source_name, imported_by=current_user.username)
    log_audit(
        session,
        current_user.username,
        "phone_enrichment_imported",
        target_type="phone_enrichment",
        target_id=source_name,
        message="Retorno telefonico importado.",
        metadata={"fonte": source_name, "arquivos": [path.name for path in saved_paths], "clientes_atualizados": summary.clients_updated},
    )
    session.commit()
    message = (
        f"Retorno importado: {summary.clients_updated} clientes atualizados, "
        f"{summary.data_fields_updated} atualizacao(oes) de dados complementares, "
        f"{summary.phones_added} telefone(s) novo(s) e {summary.skipped_not_found} CPF(s) sem cadastro."
    )
    return RedirectResponse(url="/exports?message=" + quote_plus(message), status_code=303)


@app.post("/exports/public-update-import")
async def run_public_update_import(
    base_segment: str = Form("GOVERNO"),
    source_name: str = Form(""),
    files: list[UploadFile] = File(default=[]),
    current_user: AppUser = Depends(require_permission("can_export")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    normalized_segment = normalize_base_segment(base_segment, default="GOVERNO")
    if normalized_segment not in {"GOVERNO", "PREFEITURA"}:
        return RedirectResponse(url="/exports?error=" + quote_plus("A atualizacao complementar e exclusiva para GOV/PREF."), status_code=303)

    uploaded_files: list[tuple[str, bytes]] = []
    for item in files:
        if item.filename:
            uploaded_files.append((item.filename, await item.read()))
    if not uploaded_files:
        return RedirectResponse(url=f"/exports?base_segment={normalized_segment}&error=" + quote_plus("Selecione ao menos um arquivo complementar."), status_code=303)

    saved_paths = save_uploaded_files(uploaded_files)
    summary = import_public_campaign_updates(
        session,
        saved_paths,
        base_segment=normalized_segment,
        source_name=source_name,
        imported_by=current_user.username,
    )
    log_audit(
        session,
        current_user.username,
        "public_update_imported",
        target_type="public_update",
        target_id=normalized_segment,
        message="Atualizacao complementar GOV/PREF importada.",
        metadata={"segmento": normalized_segment, "fonte": source_name, "arquivos": [path.name for path in saved_paths], "clientes_atualizados": summary.clients_updated},
    )
    session.commit()
    message = (
        f"Atualizacao GOV/PREF importada: {summary.rows_imported} linha(s), "
        f"{summary.contracts_saved} contrato(s) gravado(s), "
        f"{summary.clients_updated} cliente(s) atualizados e "
        f"{summary.skipped_not_found} CPF(s) sem base correspondente."
    )
    return RedirectResponse(url=f"/exports?base_segment={normalized_segment}&message=" + quote_plus(message), status_code=303)


@app.get("/downloads/{filename}")
def download_export(filename: str, current_user: AppUser = Depends(require_permission("can_export")), session: Session = Depends(get_session)) -> FileResponse:
    path = settings.exports_dir / filename
    if not path.exists():
        raise HTTPException(status_code=404, detail="Arquivo nao encontrado.")
    media_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    log_audit(session, current_user.username, "download_export", target_type="export_file", target_id=filename, message="Arquivo de exportacao baixado.")
    session.commit()
    return FileResponse(path=path, filename=filename, media_type=media_type)


@app.get("/history", response_class=HTMLResponse)
def history_page(
    request: Request,
    message: str = "",
    page: int = 1,
    _current_user: AppUser = Depends(require_permission("can_view_history")),
    session: Session = Depends(get_session),
) -> HTMLResponse:
    total_batches = session.scalar(select(func.count()).select_from(ImportBatch)) or 0
    pagination = build_pagination(total_batches, page, HISTORY_PAGE_SIZE)
    batches = (
        session.execute(
            select(ImportBatch)
            .order_by(ImportBatch.data_importacao.desc())
            .offset(int(pagination["offset"]))
            .limit(HISTORY_PAGE_SIZE)
        )
        .scalars()
        .all()
    )
    batch_ids = [batch.id for batch in batches]
    file_counts = {batch_id: 0 for batch_id in batch_ids}
    if batch_ids:
        file_counts.update(
            {
                batch_id: total
                for batch_id, total in session.execute(
                    select(SourceFile.batch_id, func.count())
                    .where(SourceFile.batch_id.in_(batch_ids))
                    .group_by(SourceFile.batch_id)
                ).all()
            }
        )
    return render_template(request, "history.html", {"batches": batches, "file_counts": file_counts, "message": message, "title": "Historico", "pagination": pagination})


@app.get("/history/{batch_id}", response_class=HTMLResponse)
def history_batch_detail(
    batch_id: int,
    request: Request,
    _current_user: AppUser = Depends(require_permission("can_view_history")),
    session: Session = Depends(get_session),
) -> HTMLResponse:
    batch = session.get(ImportBatch, batch_id)
    if batch is None:
        raise HTTPException(status_code=404, detail="Lote nao encontrado.")

    file_rows = (
        session.execute(
            select(
                SourceFile.id,
                SourceFile.nome_arquivo,
                SourceFile.semana_label,
                SourceFile.tipo_arquivo,
                SourceFile.total_linhas,
                func.count(ClientOccurrence.id).label("importadas"),
            )
            .outerjoin(ClientOccurrence, ClientOccurrence.source_file_id == SourceFile.id)
            .where(SourceFile.batch_id == batch_id)
            .group_by(
                SourceFile.id,
                SourceFile.nome_arquivo,
                SourceFile.semana_label,
                SourceFile.tipo_arquivo,
                SourceFile.total_linhas,
            )
            .order_by(SourceFile.nome_arquivo.asc())
        )
        .all()
    )
    imported_rows = sum(row.importadas for row in file_rows)
    total_rows = sum(row.total_linhas for row in file_rows)
    skipped_rows = max(total_rows - imported_rows, 0)
    source_file_ids = [row.id for row in file_rows]
    unique_contacts = 0
    contacts_with_phone = 0
    updated_ratio = 0
    if source_file_ids:
        unique_contacts = session.scalar(
            select(func.count(func.distinct(ClientOccurrence.cpf))).where(ClientOccurrence.source_file_id.in_(source_file_ids))
        ) or 0
        contacts_with_phone = session.scalar(
            select(func.count(func.distinct(ClientOccurrence.cpf))).where(
                ClientOccurrence.source_file_id.in_(source_file_ids),
                or_(
                    ClientOccurrence.telefone1 != "",
                    ClientOccurrence.telefone2 != "",
                    ClientOccurrence.telefone3 != "",
                ),
            )
        ) or 0
    if total_rows:
        updated_ratio = round((imported_rows / total_rows) * 100)
    return render_template(
        request,
        "history_detail.html",
        {
            "title": "Detalhes do lote",
            "batch": batch,
            "file_rows": file_rows,
            "imported_rows": imported_rows,
            "total_rows": total_rows,
            "skipped_rows": skipped_rows,
            "unique_contacts": unique_contacts,
            "contacts_with_phone": contacts_with_phone,
            "updated_ratio": updated_ratio,
        },
    )


@app.post("/history/{batch_id}/delete")
def delete_history_batch(
    batch_id: int,
    current_user: AppUser = Depends(require_permission("can_delete_batches")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    ok, message = delete_batch(session, batch_id)
    log_audit(
        session,
        current_user.username,
        "batch_deleted" if ok else "batch_delete_failed",
        target_type="batch",
        target_id=str(batch_id),
        message=message or "",
    )
    session.commit()
    target = "/history?message=" + quote_plus(message if message else ("Campanha apagada com sucesso." if ok else "Falha ao apagar lote."))
    return RedirectResponse(url=target, status_code=303)


@app.get("/clients/{cpf}", response_class=HTMLResponse)
def client_detail_page(
    cpf: str,
    request: Request,
    current_user: AppUser = Depends(require_permission("can_search")),
    session: Session = Depends(get_session),
) -> HTMLResponse:
    normalized_cpf = "".join(char for char in cpf if char.isdigit())
    if len(normalized_cpf) != 11:
        raise HTTPException(status_code=404, detail="Cliente nao encontrado.")

    client = session.get(Client, normalized_cpf)
    occurrences = (
        session.execute(
            select(ClientOccurrence, SourceFile.nome_arquivo, ImportBatch.id, ImportBatch.ano_referencia, ImportBatch.mes_referencia)
            .join(SourceFile, SourceFile.id == ClientOccurrence.source_file_id)
            .join(ImportBatch, ImportBatch.id == SourceFile.batch_id)
            .where(ClientOccurrence.cpf == normalized_cpf)
            .order_by(
                ImportBatch.ano_referencia.desc(),
                ImportBatch.mes_referencia.desc(),
                SourceFile.id.desc(),
                ClientOccurrence.id.desc(),
            )
            .limit(100)
        )
        .all()
    )
    if client is None and not occurrences:
        raise HTTPException(status_code=404, detail="Cliente nao encontrado.")

    latest_occurrence = occurrences[0][0] if occurrences else None
    if client is None and latest_occurrence is not None:
        highest_margin = max((occurrence.vl_margem for occurrence, *_rest in occurrences if occurrence.vl_margem is not None), default=None)
        latest_reference = f"{occurrences[0][4]:02d}/{occurrences[0][3]}" if occurrences[0][4] and occurrences[0][3] else "-"
        client = SimpleNamespace(
            cpf=normalized_cpf,
            nome_atual=latest_occurrence.nome or "",
            cidade_atual=latest_occurrence.cidade or "",
            uf_atual=latest_occurrence.uf or "",
            melhor_telefone=latest_occurrence.telefone1 or latest_occurrence.telefone2 or latest_occurrence.telefone3 or "",
            tem_telefone=bool(latest_occurrence.telefone1 or latest_occurrence.telefone2 or latest_occurrence.telefone3),
            tem_inss=any(normalize_base_segment(item[0].base_segment) == "INSS" for item in occurrences),
            tem_governo=any(normalize_base_segment(item[0].base_segment) == "GOVERNO" for item in occurrences),
            pis_atual=latest_occurrence.pis or "",
            matricula_atual=latest_occurrence.matricula or "",
            entidade_atual=latest_occurrence.entidade or "",
            maior_margem=highest_margin,
            ultima_referencia=latest_reference,
            qtd_ocorrencias=len(occurrences),
            atualizado_em=latest_occurrence.criado_em,
        )

    segment_details: dict[str, dict[str, object]] = {}
    public_occurrences_by_key: dict[tuple[str, str], tuple[ClientOccurrence, str, int, int, int]] = {}
    public_fallback_by_segment: dict[str, tuple[ClientOccurrence, str, int, int, int]] = {}
    for occurrence, file_name, batch_id, ano_referencia, mes_referencia in occurrences:
        segment = normalize_base_segment(occurrence.base_segment)
        if segment in {"GOVERNO", "PREFEITURA"} and occurrence.matricula:
            public_occurrences_by_key[(segment, occurrence.matricula)] = (occurrence, file_name, batch_id, ano_referencia, mes_referencia)
            public_fallback_by_segment.setdefault(segment, (occurrence, file_name, batch_id, ano_referencia, mes_referencia))
        if segment in segment_details:
            continue
        address_parts = [part for part in [occurrence.logradouro, occurrence.numero, occurrence.complemento, occurrence.bairro] if part]
        address_text = ", ".join(str(part) for part in address_parts)
        extras = _parse_extras_json(occurrence.extras_json or "")
        segment_details[segment] = {
            "segment": segment,
            "label": base_segment_label(segment),
            "source": occurrence.base_source or "-",
            "phone": occurrence.telefone1 or occurrence.telefone2 or occurrence.telefone3 or "",
            "phone_1": occurrence.telefone1 or "",
            "phone_2": occurrence.telefone2 or "",
            "phone_3": occurrence.telefone3 or "",
            "email": occurrence.email or "",
            "address": address_text,
            "bairro": occurrence.bairro or "",
            "cep": occurrence.cep or "",
            "city": occurrence.cidade or client.cidade_atual or "-",
            "uf": occurrence.uf or client.uf_atual or "-",
            "reference": f"{mes_referencia:02d}/{ano_referencia}" if mes_referencia and ano_referencia else "-",
            "margin": occurrence.vl_margem,
            "rmi": occurrence.vl_rmi,
            "rmc": occurrence.vl_rmc,
            "margem_total": occurrence.margem_total,
            "margem_liquida": occurrence.margem_liquida,
            "margem_bruta": occurrence.margem_bruta,
            "margem_utilizada": occurrence.margem_utilizada,
            "benefit": occurrence.nu_nb or "",
            "species": occurrence.esp or "",
            "matricula": occurrence.matricula or "",
            "entity": occurrence.entidade or "",
            "cbo_title": occurrence.cbo_titulo or "",
            "service": occurrence.servico or "",
            "status": occurrence.situacao or "",
            "pis": occurrence.pis or client.pis_atual or "",
            "salary": occurrence.salario,
            "file_name": file_name,
            "updated_at": occurrence.criado_em,
            "regional_alert": _first_extra_value(extras, "ALERTA_REGIONAL"),
            "ddd_phone": _first_extra_value(extras, "DDD_TELEFONE"),
            "ddd_uf": _first_extra_value(extras, "DDD_UF_BRASILAPI"),
        }

    public_matricula_details: list[dict[str, object]] = []
    public_updates = (
        session.execute(
            select(PublicCampaignUpdate)
            .where(PublicCampaignUpdate.cpf == normalized_cpf)
            .order_by(
                PublicCampaignUpdate.base_segment.asc(),
                PublicCampaignUpdate.matricula.asc(),
                PublicCampaignUpdate.criado_em.desc(),
                PublicCampaignUpdate.id.desc(),
            )
        )
        .scalars()
        .all()
    )
    seen_public_keys: set[tuple[str, str]] = set()
    for update in public_updates:
        segment = normalize_base_segment(update.base_segment)
        if segment not in {"GOVERNO", "PREFEITURA"} or not update.matricula:
            continue
        key = (segment, update.matricula)
        if key in seen_public_keys:
            continue
        seen_public_keys.add(key)
        occurrence_bundle = public_occurrences_by_key.get(key) or public_fallback_by_segment.get(segment)
        occurrence = occurrence_bundle[0] if occurrence_bundle else None
        file_name = occurrence_bundle[1] if occurrence_bundle else "-"
        ano_referencia = occurrence_bundle[3] if occurrence_bundle else None
        mes_referencia = occurrence_bundle[4] if occurrence_bundle else None
        extras = _parse_extras_json(occurrence.extras_json or "") if occurrence is not None else {}
        public_matricula_details.append(
            {
                "segment": segment,
                "label": base_segment_label(segment),
                "matricula": update.matricula,
                "service": update.servico_servidor or (occurrence.servico if occurrence is not None else ""),
                "status": update.situacao or (occurrence.situacao if occurrence is not None else ""),
                "entity": occurrence.entidade if occurrence is not None else "",
                "secretaria": occurrence.secretaria if occurrence is not None else "",
                "convenio": occurrence.convenio if occurrence is not None else "",
                "pis": occurrence.pis if occurrence is not None else "",
                "salary": occurrence.salario if occurrence is not None else None,
                "margin": update.margem_disponivel if update.margem_disponivel is not None else (occurrence.vl_margem if occurrence is not None else None),
                "margem_total": update.margem_total if update.margem_total is not None else (occurrence.margem_total if occurrence is not None else None),
                "margem_liquida": occurrence.margem_liquida if occurrence is not None else None,
                "margem_bruta": occurrence.margem_bruta if occurrence is not None else None,
                "margem_utilizada": occurrence.margem_utilizada if occurrence is not None else None,
                "tipo_vinculo": _first_extra_value(extras, "REGIME_CONTRATACAO", "TIPO DE VINCULO", "TIPO_VINCULO", "VINCULO", "REGIME DE CONTRATACAO"),
                "cargo_funcao": _first_extra_value(extras, "CARGO", "FUNCAO", "FUNÇÃO", "FUNÃƒÆ’Ã¢â‚¬Â¡ÃƒÆ’Ã†â€™O", "CARGO/FUNCAO", "CARGO FUNCAO", "CBO TITULO", "CBO_TITULO") or (occurrence.cbo_titulo if occurrence is not None else ""),
                "phone": (occurrence.telefone1 or occurrence.telefone2 or occurrence.telefone3) if occurrence is not None else "",
                "email": occurrence.email if occurrence is not None else "",
                "city": occurrence.cidade if occurrence is not None else client.cidade_atual,
                "uf": occurrence.uf if occurrence is not None else client.uf_atual,
                "source": update.fonte or (occurrence.base_source if occurrence is not None else ""),
                "source_file": update.arquivo_origem or file_name,
                "reference": f"{mes_referencia:02d}/{ano_referencia}" if mes_referencia and ano_referencia else "-",
                "banks": _first_extra_value(extras, "PUBLIC_UPDATE_BANCOS"),
                "consignataria": _first_extra_value(extras, "PUBLIC_UPDATE_CONSIGNATARIA_PRINCIPAL") or update.consignataria,
                "contracts_count": _first_extra_value(extras, "PUBLIC_UPDATE_QTD_CONTRATOS"),
                "active_count": _first_extra_value(extras, "PUBLIC_UPDATE_QTD_ATIVOS"),
                "updated_at": update.criado_em or (occurrence.criado_em if occurrence is not None else None),
            }
        )

    for (segment, matricula), occurrence_bundle in public_occurrences_by_key.items():
        if (segment, matricula) in seen_public_keys:
            continue
        occurrence, file_name, _batch_id, ano_referencia, mes_referencia = occurrence_bundle
        extras = _parse_extras_json(occurrence.extras_json or "")
        public_matricula_details.append(
            {
                "segment": segment,
                "label": base_segment_label(segment),
                "matricula": matricula,
                "service": occurrence.servico or "",
                "status": occurrence.situacao or "",
                "entity": occurrence.entidade or "",
                "secretaria": occurrence.secretaria or "",
                "convenio": occurrence.convenio or "",
                "pis": occurrence.pis or "",
                "salary": occurrence.salario,
                "margin": occurrence.vl_margem,
                "margem_total": occurrence.margem_total,
                "margem_liquida": occurrence.margem_liquida,
                "margem_bruta": occurrence.margem_bruta,
                "margem_utilizada": occurrence.margem_utilizada,
                "tipo_vinculo": _first_extra_value(extras, "REGIME_CONTRATACAO", "TIPO DE VINCULO", "TIPO_VINCULO", "VINCULO", "REGIME DE CONTRATACAO"),
                "cargo_funcao": _first_extra_value(extras, "CARGO", "FUNCAO", "FUNÇÃO", "FUNÃƒÆ’Ã¢â‚¬Â¡ÃƒÆ’Ã†â€™O", "CARGO/FUNCAO", "CARGO FUNCAO", "CBO TITULO", "CBO_TITULO") or occurrence.cbo_titulo,
                "phone": occurrence.telefone1 or occurrence.telefone2 or occurrence.telefone3 or "",
                "email": occurrence.email or "",
                "city": occurrence.cidade or client.cidade_atual,
                "uf": occurrence.uf or client.uf_atual,
                "source": occurrence.base_source or "",
                "source_file": file_name,
                "reference": f"{mes_referencia:02d}/{ano_referencia}" if mes_referencia and ano_referencia else "-",
                "banks": _first_extra_value(extras, "PUBLIC_UPDATE_BANCOS"),
                "consignataria": _first_extra_value(extras, "PUBLIC_UPDATE_CONSIGNATARIA_PRINCIPAL"),
                "contracts_count": _first_extra_value(extras, "PUBLIC_UPDATE_QTD_CONTRATOS"),
                "active_count": _first_extra_value(extras, "PUBLIC_UPDATE_QTD_ATIVOS"),
                "updated_at": occurrence.criado_em,
            }
        )

    public_matricula_details.sort(
        key=lambda item: (
            {"GOVERNO": 0, "PREFEITURA": 1}.get(str(item.get("segment", "")), 9),
            str(item.get("matricula", "")),
        )
    )
    log_audit(session, current_user.username, "client_detail_viewed", target_type="client", target_id=normalized_cpf, message="Detalhe do cliente consultado.")
    session.commit()
    return render_template(
        request,
        "client_detail.html",
        {
            "title": f"Cliente {client.nome_atual or client.cpf}",
            "client": client,
            "latest_occurrence": latest_occurrence,
            "occurrences": occurrences,
            "segment_details": [segment_details[key] for key in ["INSS", "GOVERNO", "PREFEITURA"] if key in segment_details],
            "public_matricula_details": public_matricula_details,
        },
    )


@app.post("/clients/favorites")
def save_client_filter(
    name: str = Form(...),
    query_string: str = Form(""),
    current_user: AppUser = Depends(require_permission("can_search")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    favorite_name = name.strip()
    if not favorite_name:
        return RedirectResponse(url="/clients?error=" + quote_plus("Informe um nome para o filtro."), status_code=303)
    session.add(SavedFilter(user_id=current_user.id, page_key="clients", name=favorite_name[:120], query_string=query_string.strip()))
    log_audit(session, current_user.username, "saved_filter_created", target_type="saved_filter", message="Filtro favorito salvo.", metadata={"nome": favorite_name})
    session.commit()
    return RedirectResponse(url="/clients?message=" + quote_plus("Filtro favorito salvo."), status_code=303)


@app.post("/clients/{cpf}/do-not-call")
def set_client_do_not_call(
    cpf: str,
    enabled: bool = Form(False),
    return_url: str = Form("/clients"),
    current_user: AppUser = Depends(require_permission("can_search")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    normalized_cpf = normalize_cpf(cpf)
    client = session.get(Client, normalized_cpf)
    if client is None:
        return RedirectResponse(url="/clients?error=" + quote_plus("Cliente nao encontrado para marcar bloqueio de ligacao."), status_code=303)
    client.do_not_call = bool(enabled)
    action = "do_not_call_enabled" if client.do_not_call else "do_not_call_disabled"
    message = "Cliente marcado como NAO LIGAR." if client.do_not_call else "Cliente removido da lista NAO LIGAR."
    log_audit(
        session,
        current_user.username,
        action,
        target_type="client",
        target_id=normalized_cpf,
        message=message,
        metadata={"cpf": normalized_cpf, "enabled": client.do_not_call},
    )
    session.commit()
    target_url = safe_next_path(return_url, "/clients")
    separator = "&" if "?" in target_url else "?"
    return RedirectResponse(url=f"{target_url}{separator}message=" + quote_plus(message), status_code=303)


@app.post("/clients/favorites/{favorite_id}/delete")
def delete_client_filter(
    favorite_id: int,
    current_user: AppUser = Depends(require_permission("can_search")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    favorite = session.get(SavedFilter, favorite_id)
    if favorite is None or favorite.user_id != current_user.id:
        return RedirectResponse(url="/clients?error=" + quote_plus("Filtro favorito nao encontrado."), status_code=303)
    favorite_name = favorite.name
    session.delete(favorite)
    log_audit(session, current_user.username, "saved_filter_deleted", target_type="saved_filter", target_id=str(favorite_id), message="Filtro favorito removido.", metadata={"nome": favorite_name})
    session.commit()
    return RedirectResponse(url="/clients?message=" + quote_plus("Filtro favorito removido."), status_code=303)


@app.get("/users", response_class=HTMLResponse)
def users_page(
    request: Request,
    current_user: AppUser = Depends(require_permission("can_manage_users")),
    session: Session = Depends(get_session),
) -> HTMLResponse:
    users = session.execute(select(AppUser).order_by(AppUser.username.asc())).scalars().all()
    recent_audits = session.execute(select(AuditLog).order_by(AuditLog.created_at.desc()).limit(20)).scalars().all()
    return render_template(
        request,
        "users.html",
        {
            "title": "Usuarios",
            "users": users,
            "permission_labels": PERMISSION_LABELS,
            "recent_audits": recent_audits,
            "message": request.query_params.get("message", ""),
            "error": request.query_params.get("error", ""),
            "current_manager_id": current_user.id,
        },
    )


@app.post("/users")
def create_user(
    username: str = Form(...),
    full_name: str = Form(""),
    password: str = Form(...),
    is_active: bool = Form(False),
    is_admin: bool = Form(False),
    can_view_dashboard: bool = Form(False),
    can_import: bool = Form(False),
    can_search: bool = Form(False),
    can_export: bool = Form(False),
    can_view_history: bool = Form(False),
    can_delete_batches: bool = Form(False),
    can_manage_users: bool = Form(False),
    current_user: AppUser = Depends(require_permission("can_manage_users")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    normalized_username = username.strip().lower()
    if not normalized_username:
        return RedirectResponse(url="/users?error=" + quote_plus("Informe um usuario valido."), status_code=303)
    existing = session.scalar(select(AppUser).where(AppUser.username == normalized_username))
    if existing is not None:
        return RedirectResponse(url="/users?error=" + quote_plus("Ja existe um usuario com esse login."), status_code=303)
    password_hash, password_salt = hash_password(password)
    new_user = AppUser(
        username=normalized_username,
        full_name=full_name.strip(),
        password_hash=password_hash,
        password_salt=password_salt,
        is_active=is_active,
        is_admin=is_admin,
        can_view_dashboard=can_view_dashboard,
        can_import=can_import,
        can_search=can_search,
        can_export=can_export,
        can_view_history=can_view_history,
        can_delete_batches=can_delete_batches,
        can_manage_users=can_manage_users,
    )
    session.add(new_user)
    log_audit(session, current_user.username, "user_created", target_type="user", target_id=normalized_username, message="Usuario criado.")
    session.commit()
    return RedirectResponse(url="/users?message=" + quote_plus("Usuario criado com sucesso."), status_code=303)


@app.post("/users/{user_id}")
def update_user(
    user_id: int,
    username: str = Form(...),
    full_name: str = Form(""),
    password: str = Form(""),
    is_active: bool = Form(False),
    is_admin: bool = Form(False),
    can_view_dashboard: bool = Form(False),
    can_import: bool = Form(False),
    can_search: bool = Form(False),
    can_export: bool = Form(False),
    can_view_history: bool = Form(False),
    can_delete_batches: bool = Form(False),
    can_manage_users: bool = Form(False),
    current_user: AppUser = Depends(require_permission("can_manage_users")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    user = session.get(AppUser, user_id)
    if user is None:
        return RedirectResponse(url="/users?error=" + quote_plus("Usuario nao encontrado."), status_code=303)

    normalized_username = username.strip().lower()
    duplicate = session.scalar(select(AppUser).where(AppUser.username == normalized_username, AppUser.id != user_id))
    if duplicate is not None:
        return RedirectResponse(url="/users?error=" + quote_plus("Ja existe outro usuario com esse login."), status_code=303)

    user.username = normalized_username
    user.full_name = full_name.strip()
    user.is_active = is_active
    user.is_admin = is_admin
    user.can_view_dashboard = can_view_dashboard
    user.can_import = can_import
    user.can_search = can_search
    user.can_export = can_export
    user.can_view_history = can_view_history
    user.can_delete_batches = can_delete_batches
    user.can_manage_users = can_manage_users
    if password.strip():
        user.password_hash, user.password_salt = hash_password(password.strip())
        user.last_password_reset_at = current_utc()
    if user.id == current_user.id and not user.is_active:
        return RedirectResponse(url="/users?error=" + quote_plus("Voce nao pode desativar seu proprio usuario."), status_code=303)
    log_audit(session, current_user.username, "user_updated", target_type="user", target_id=str(user.id), message="Usuario atualizado.", metadata={"username": normalized_username})
    session.commit()
    return RedirectResponse(url="/users?message=" + quote_plus("Usuario atualizado com sucesso."), status_code=303)


@app.post("/users/{user_id}/reset-password")
def reset_user_password(
    user_id: int,
    new_password: str = Form(...),
    current_user: AppUser = Depends(require_permission("can_manage_users")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    user = session.get(AppUser, user_id)
    if user is None:
        return RedirectResponse(url="/users?error=" + quote_plus("Usuario nao encontrado."), status_code=303)
    password = new_password.strip()
    if len(password) < 4:
        return RedirectResponse(url="/users?error=" + quote_plus("Informe uma senha valida para reset."), status_code=303)
    user.password_hash, user.password_salt = hash_password(password)
    user.last_password_reset_at = current_utc()
    user.locked_until = None
    log_audit(session, current_user.username, "user_password_reset", target_type="user", target_id=str(user.id), message="Senha redefinida.", metadata={"username": user.username})
    session.commit()
    return RedirectResponse(url="/users?message=" + quote_plus(f"Senha redefinida para {user.username}."), status_code=303)


@app.post("/users/{user_id}/lock")
def lock_user_temporarily(
    user_id: int,
    minutes: int = Form(60),
    current_user: AppUser = Depends(require_permission("can_manage_users")),
    session: Session = Depends(get_session),
) -> RedirectResponse:
    user = session.get(AppUser, user_id)
    if user is None:
        return RedirectResponse(url="/users?error=" + quote_plus("Usuario nao encontrado."), status_code=303)
    if user.id == current_user.id and minutes > 0:
        return RedirectResponse(url="/users?error=" + quote_plus("Voce nao pode bloquear seu proprio usuario."), status_code=303)
    if minutes <= 0:
        user.locked_until = None
        action = "user_unlocked"
        message = f"Bloqueio removido de {user.username}."
    else:
        from datetime import timedelta

        user.locked_until = current_utc() + timedelta(minutes=minutes)
        action = "user_locked"
        message = f"Usuario {user.username} bloqueado por {minutes} minuto(s)."
    log_audit(session, current_user.username, action, target_type="user", target_id=str(user.id), message=message, metadata={"username": user.username, "minutes": minutes})
    session.commit()
    return RedirectResponse(url="/users?message=" + quote_plus(message), status_code=303)


@app.get("/health")
def healthcheck(session: Session = Depends(get_session)) -> JSONResponse:
    db_ok = False
    try:
        db_ok = bool(session.scalar(select(1)))
    except Exception:
        db_ok = False
    payload = {"status": "ok" if db_ok else "degraded", "database": "ok" if db_ok else "error"}
    return JSONResponse(payload, status_code=200 if db_ok else 503)


@app.get("/api/c6/worker-loan/health")
def c6_worker_health(_current_user: AppUser = Depends(require_permission("can_search"))) -> JSONResponse:
    payload = {
        "enabled": settings.c6_worker_enabled,
        "base_url": settings.c6_base_url,
        "has_credentials": bool(settings.c6_username and settings.c6_password),
    }
    return JSONResponse(payload, status_code=200)


@app.post("/api/c6/worker-loan/offer")
def c6_generate_worker_offer(
    request: Request,
    body: dict[str, object] = Body(...),
    current_user: AppUser = Depends(require_permission("can_search")),
    session: Session = Depends(get_session),
) -> JSONResponse:
    try:
        response = c6_worker_loan_client.generate_offer(body)
    except C6WorkerLoanError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    log_audit(
        session,
        current_user.username,
        "c6_worker_offer_generated",
        target_type="integration",
        message="Oferta do consignado trabalhador solicitada.",
        metadata={"path": str(request.url.path)},
    )
    session.commit()
    return JSONResponse({"status_code": response.status_code, "data": response.payload}, status_code=200)


@app.post("/api/c6/worker-loan/simulation")
def c6_simulate_worker_loan(
    request: Request,
    body: dict[str, object] = Body(...),
    version: str = "v2",
    current_user: AppUser = Depends(require_permission("can_search")),
    session: Session = Depends(get_session),
) -> JSONResponse:
    try:
        response = c6_worker_loan_client.simulate_proposal(body, version=version)
    except C6WorkerLoanError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    log_audit(
        session,
        current_user.username,
        "c6_worker_simulation_requested",
        target_type="integration",
        message="Simulacao do consignado trabalhador solicitada.",
        metadata={"path": str(request.url.path), "version": version},
    )
    session.commit()
    return JSONResponse({"status_code": response.status_code, "data": response.payload}, status_code=200)


@app.post("/api/c6/worker-loan/include")
def c6_include_worker_loan(
    request: Request,
    body: dict[str, object] = Body(...),
    current_user: AppUser = Depends(require_permission("can_search")),
    session: Session = Depends(get_session),
) -> JSONResponse:
    try:
        response = c6_worker_loan_client.include_proposal(body)
    except C6WorkerLoanError as exc:
        raise HTTPException(status_code=502, detail=str(exc)) from exc

    log_audit(
        session,
        current_user.username,
        "c6_worker_proposal_included",
        target_type="integration",
        message="Inclusao do consignado trabalhador solicitada.",
        metadata={"path": str(request.url.path)},
    )
    session.commit()
    return JSONResponse({"status_code": response.status_code, "data": response.payload}, status_code=200)
