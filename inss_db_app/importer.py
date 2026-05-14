from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import shutil
import subprocess
import tempfile
import threading
import unicodedata
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from sqlalchemy import Select, and_, delete, func, insert, or_, select
from sqlalchemy.exc import SQLAlchemyError
from sqlalchemy.orm import Session

from .address_lookup import lookup_brasilapi_ddd, lookup_cep_address
from .config import settings
from .database import SessionLocal, serialized_write
from .models import Client, ClientOccurrence, EnrichmentImportItem, ExportHistory, ImportBatch, PhoneEnrichment, PublicCampaignUpdate, SourceFile
from .schemas import ImportRequest, ImportSummary


EXPORT_HEADERS = [
    "TELEFONE",
    "NOME",
    "CPF",
    "NU-NB",
    "VL-RMI",
    "VL MARGEM",
    "VL RMC",
    "DDB",
    "UF",
    "CIDADE",
    "ESP",
]
PUBLIC_EXPORT_HEADERS = [
    "TELEFONE",
    "NOME",
    "CPF",
    "MATRICULA",
    "SERVICO",
    "SITUACAO",
    "ENTIDADE",
    "CBO TITULO",
    "MARGEM",
    "MARGEM TOTAL",
    "CARGO",
    "TIPO_VINCULO",
    "CONVENIO",
    "FONTE",
    "CAMPANHA",
    "UF",
    "CIDADE",
]
ENRICHMENT_CPF_HEADERS = ["CPF", "NOME"]
WRONG_NUMBER_EXPORT_HEADERS = ["CPF", "NOME", "TELEFONE", "MOTIVO"]
AUDIT_HEADERS = EXPORT_HEADERS + ["ARQUIVO_ORIGEM", "SEMANA", "ANO", "MES", "TELEFONE2", "TELEFONE3", "EMAIL"]
PUBLIC_AUDIT_HEADERS = PUBLIC_EXPORT_HEADERS + ["ARQUIVO_ORIGEM", "SEMANA", "ANO", "MES", "TELEFONE2", "TELEFONE3", "EMAIL", "FONTE"]
UPDATED_BASE_HEADERS = [
    "TELEFONE",
    "TELEFONE2",
    "TELEFONE3",
    "WHATSAPP",
    "WHATSAPP2",
    "NOME",
    "CPF",
    "SEGMENTO",
    "FONTE",
    "EMAIL",
    "EMAIL2",
    "EMAIL3",
    "CIDADE",
    "UF",
    "ENDERECO",
    "BAIRRO",
    "LOGRADOURO",
    "NUMERO",
    "COMPLEMENTO",
    "CEP",
    "DT_NASC",
    "IDADE",
    "NOME_MAE",
    "SEXO",
    "MATRICULA",
    "SERVICO",
    "SITUACAO",
    "ENTIDADE",
    "TIPO_VINCULO",
    "CARGO",
    "SALARIO",
    "MARGEM_DISPONIVEL",
    "MARGEM_RMC",
    "MARGEM_TOTAL",
    "DDB",
    "BENEFICIO",
    "ESPECIE",
]
BENEFIT_BOUND_FIELDS = {
    "nu_nb",
    "esp",
    "ddb",
    "vl_rmi",
    "vl_margem",
    "vl_rmc",
    "margem_total",
    "margem_liquida",
    "margem_bruta",
    "margem_utilizada",
    "salario",
    "referencia",
}
SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xlsm", ".xls"}
SQLITE_VARIABLE_CHUNK_SIZE = 900
HEADER_ALIASES = {
    "CPF": {
        "CPF",
        "NU CPF",
        "NUCPF",
        "NUMERO CPF",
        "DOCUMENTO CPF",
    },
    "NOME": {
        "NOME",
        "NOME BENEFICIARIO",
        "NOME DO BENEFICIARIO",
        "NOME DO SERVIDOR",
        "SERVIDOR",
        "NOME X",
        "NOME_X",
    },
    "NOME_HIGIENIZADO": {"NOME HIGIENIZADO"},
    "MATRICULA": {"MATRICULA", "MATRÃCULA"},
    "NU-NB": {
        "NU NB",
        "NU-NB",
        "NB",
        "NUMERO BENEFICIO",
        "NUMERO DO BENEFICIO",
        "NUMERO DE BENEFICIO",
        "BENEFICIO",
    },
    "ESP": {
        "ESP",
        "ESPECIE",
        "CODIGO ESPECIE",
        "COD ESPECIE",
    },
    "DDB": {
        "DDB",
        "DATA DE BENEFICIO",
        "DT BENEFICIO",
        "DT-BENEFICIO",
        "CONCESSAO",
        "DATA CONCESSAO",
        "DATA DE CONCESSAO",
    },
    "VL-RMI": {
        "VL RMI",
        "VL-RMI",
        "RMI",
        "SALARIO",
        "VALOR DO BENEFICIO",
        "VALOR BENEFICIO",
        "MARGEM TOTAL",
    },
    "VL MARGEM": {
        "VL MARGEM",
        "MARGEM",
        "MARGEM DISPONIVEL",
        "MARGEM CONSIGNAVEL",
        "VALOR MARGEM",
    },
    "VL RMC": {
        "VL RMC",
        "RMC",
        "MARGEM CARTAO",
    },
    "UF": {"UF", "ESTADO"},
    "PIS": {"PIS", "PIS/PASEP"},
    "ENTIDADE": {"ENTIDADE", "ESTADO_ENTIDADE", "ORGAO", "ORGÃƒO"},
    "SECRETARIA": {"SECRETARIA"},
    "CONVENIO": {"CONVENIO", "CONVENIO ", "CONVÃŠNIO"},
    "REGIME_CONTRATACAO": {"REGIME DE CONTRATACAO", "REGIME DE CONTRATAÃ‡ÃƒO", "GRUPO"},
    "SALARIO": {"SALARIO", "RENDA", "REMUNERACAO DO MES", "REMUNERAÃ‡ÃƒO DO MÃŠS"},
    "MARGEM_LIQUIDA": {"MARGEM LIQUIDA", " MARGEM LIQUIDA "},
    "MARGEM_BRUTA": {"MARGEM BRUTA", " MARGEM BRUTA "},
    "MARGEM_UTILIZADA": {"MARGEM UTILIZADA"},
    "CIDADE": {"CIDADE", "MUNICIPIO"},
    "BAIRRO": {"BAIRRO"},
    "LOGRADOURO": {"LOGRADOURO", "ENDERECO", "ENDERECO LOGRADOURO", "ENDEREÃ‡O"},
    "NUMERO": {"NUMERO", "NRO", "NUM", "N"},
    "COMPLEMENTO": {"COMPLEMENTO"},
    "CEP": {"CEP"},
    "TELEFONE1": {
        "TELEFONE1",
        "TELEFONE_01",
        "TELEFONE",
        "TEL",
        "CELULAR",
        "CELULAR ATUAL",
        "TELEFONE (CADASTRO GOV BA)",
        "CEL",
        "FIXO1",
        "CEL1",
        "LEMIT",
        "NOVA VIDA",
        "NOVAVIDA",
        "FONE",
        "DDD + FONE",
        "DDD+FONE",
        "DDD FONE",
    },
    "TELEFONE2": {"TELEFONE2", "TELEFONE_02", "CELULAR2", "TEL2", "CEL2", "FIXO2", "COMERCIAL2"},
    "TELEFONE3": {"TELEFONE3", "TELEFONE_03", "CELULAR3", "TEL3", "CEL3", "FIXO3", "TELEFONE_04", "TELEFONE_05", "TELEFONE_06"},
    "DT-NASC": {
        "DT-NASC",
        "DT NASC",
        "DT NASCIMENTO",
        "DT_NASC",
        "DT_NASCIMENTO",
        "DATA NASCIMENTO",
        "DATA DE NASCIMENTO",
        "NASC",
        "DATA NASC",
    },
    "NOME_MAE": {"NOME_MAE", "NOME MAE", "MAE", "NOME DA MAE"},
    "SEXO": {"SEXO", "GENERO"},
    "EMAIL": {"EMAIL", "E-MAIL", "MAIL"},
    "CARGO": {"CARGO"},
    "DDD": {"DDD"},
}
PHONE_HEADER_HINTS = ("TELEFONE", "TEL", "CELULAR", "CEL", "FONE", "LEMIT", "NOVAVIDA", "FIXO", "COMERCIAL")
VALID_BASE_SEGMENTS = {"INSS", "CREFAZ", "GOVERNO", "PREFEITURA"}
PUBLIC_BASE_SEGMENTS = {"GOVERNO", "PREFEITURA"}
OCCURRENCE_STR_MAX_LENGTHS: dict[str, int] = {
    "cpf": 11,
    "cpf_original": 20,
    "base_segment": 12,
    "base_source": 120,
    "nome": 255,
    "nome_higienizado": 255,
    "matricula": 80,
    "servico": 255,
    "situacao": 120,
    "nu_nb": 30,
    "pis": 20,
    "entidade": 255,
    "secretaria": 255,
    "convenio": 255,
    "cbo_titulo": 255,
    "esp": 10,
    "ddb": 40,
    "uf": 2,
    "cidade": 120,
    "bairro": 120,
    "logradouro": 255,
    "numero": 40,
    "complemento": 120,
    "cep": 20,
    "telefone1": 20,
    "telefone2": 20,
    "telefone3": 20,
    "dt_nasc": 40,
    "nome_mae": 255,
    "sexo": 20,
    "email": 255,
    "row_hash": 64,
}


@dataclass
class FilterSet:
    year: int | None = None
    month: int | None = None
    week: str = ""
    ddb_start: str = ""
    ddb_end: str = ""
    uf: str = ""
    city: str = ""
    esp: str = ""
    margin_min: Decimal | None = None
    margin_max: Decimal | None = None
    age_min: int | None = None
    age_max: int | None = None
    has_phone: bool | None = None
    source_file: str = ""
    base_segment: str = "INSS"
    base_source: str = ""
    entidade: str = ""
    convenio: str = ""
    servico_publico: str = ""
    situacao: str = ""
    cargo: str = ""
    tipo_vinculo: str = ""
    salary_min: Decimal | None = None
    salary_max: Decimal | None = None
    campaign_profile: str = ""
    cpf: str = ""
    phone: str = ""
    ddd: str = ""
    name: str = ""
    quick_mode: str = "cpf"
    quick_value: str = ""
    include_do_not_call: bool = False


@dataclass
class PhoneEnrichmentSummary:
    files_processed: int = 0
    rows_processed: int = 0
    clients_matched: int = 0
    clients_updated: int = 0
    data_fields_updated: int = 0
    phones_added: int = 0
    skipped_without_cpf: int = 0
    skipped_without_phone: int = 0
    skipped_not_found: int = 0


@dataclass
class PublicCampaignUpdateSummary:
    files_processed: int = 0
    rows_processed: int = 0
    rows_imported: int = 0
    contracts_saved: int = 0
    clients_matched: int = 0
    clients_updated: int = 0
    skipped_without_cpf: int = 0
    skipped_not_found: int = 0


@dataclass
class PublicCampaignAggregate:
    cpf: str
    matricula: str = ""
    nome: str = ""
    tipo_servico_servidor: str = ""
    servico_servidor: str = ""
    entidade: str = ""
    cbo_titulo: str = ""
    margem_disponivel: Decimal | None = None
    margem_total: Decimal | None = None
    situacao_principal: str = ""
    consignataria_principal: str = ""
    qtd_contratos: int = 0
    qtd_ativos: int = 0
    bancos: set[str] = field(default_factory=set)
    arquivos: set[str] = field(default_factory=set)


@dataclass
class AddressBackfillSummary:
    scanned_rows: int = 0
    updated_rows: int = 0
    rebuilt_clients: int = 0
    skipped_invalid_cep: int = 0


@dataclass
class AddressBackfillPlan:
    occurrence_id: int
    cpf: str
    phones: list[str]
    normalized: dict[str, str | Decimal | None]

def normalize_base_segment(value: str | None, default: str = "INSS") -> str:
    normalized = normalize_upper(value or default)
    aliases = {
        "GOV": "GOVERNO",
        "GOVERNO": "GOVERNO",
        "PREF": "PREFEITURA",
        "PREFEITURA": "PREFEITURA",
        "INSS": "INSS",
        "CREF": "CREFAZ",
        "CREFAS": "CREFAZ",
        "CREFAZ": "CREFAZ",
    }
    resolved = aliases.get(normalized, normalized)
    return resolved if resolved in VALID_BASE_SEGMENTS else default


def norm_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip().replace("\ufeff", "")


def parse_multi_text_filter(value: str) -> list[str]:
    if not value:
        return []
    return [item.strip() for item in value.split("|") if item.strip()]


def strip_accents(text: str) -> str:
    text = norm_text(text)
    return "".join(c for c in unicodedata.normalize("NFKD", text) if not unicodedata.combining(c))


def normalize_header_name(text: str) -> str:
    cleaned = strip_accents(text).upper()
    cleaned = re.sub(r"[^A-Z0-9]+", " ", cleaned)
    return re.sub(r"\s+", " ", cleaned).strip()


def strip_excel_trailing_decimal(text: str) -> str:
    cleaned = norm_text(text)
    if re.fullmatch(r"-?\d+\.0+", cleaned):
        return cleaned.split(".", 1)[0]
    return cleaned


def only_digits(text: str) -> str:
    cleaned = strip_excel_trailing_decimal(text)
    # Planilhas CREFAZ podem trazer CPF/CNPJ em notacao cientifica (ex.: 1,48424E+13).
    if re.fullmatch(r"-?\d+[.,]\d+e[+-]?\d+", cleaned, flags=re.IGNORECASE):
        try:
            numeric = Decimal(cleaned.replace(",", "."))
            as_text = format(numeric, "f")
            integer_part = as_text.split(".", 1)[0]
            return re.sub(r"\D", "", integer_part)
        except (InvalidOperation, ValueError):
            pass
    return re.sub(r"\D", "", cleaned)


def normalize_cpf(text: str) -> str:
    digits = only_digits(text)
    if not digits:
        return ""
    if len(digits) == 14:
        # Quando o campo "CPF" vem com CNPJ, geramos uma chave numerica estavel
        # para manter compatibilidade com o schema atual (11 digitos no campo cpf).
        digest = hashlib.sha1(digits.encode("utf-8")).hexdigest()
        stable_tail = int(digest[:12], 16) % 10_000_000_000
        return f"9{stable_tail:010d}"
    if len(digits) > 11:
        digits = digits[-11:]
    if set(digits) == {"0"}:
        return ""
    return digits.zfill(11)


def normalize_document_original(text: str) -> str:
    digits = only_digits(text)
    if not digits:
        return ""
    return digits


def is_valid_cpf_digits(digits: str) -> bool:
    if not re.fullmatch(r"\d{11}", digits):
        return False
    if len(set(digits)) == 1:
        return False
    total = sum(int(digits[i]) * (10 - i) for i in range(9))
    check_1 = (total * 10) % 11
    if check_1 == 10:
        check_1 = 0
    if check_1 != int(digits[9]):
        return False
    total = sum(int(digits[i]) * (11 - i) for i in range(10))
    check_2 = (total * 10) % 11
    if check_2 == 10:
        check_2 = 0
    return check_2 == int(digits[10])


def is_valid_cnpj_digits(digits: str) -> bool:
    if not re.fullmatch(r"\d{14}", digits):
        return False
    if len(set(digits)) == 1:
        return False
    weights_1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    weights_2 = [6] + weights_1
    total_1 = sum(int(digits[i]) * weights_1[i] for i in range(12))
    check_1 = 11 - (total_1 % 11)
    if check_1 >= 10:
        check_1 = 0
    if check_1 != int(digits[12]):
        return False
    total_2 = sum(int(digits[i]) * weights_2[i] for i in range(13))
    check_2 = 11 - (total_2 % 11)
    if check_2 >= 10:
        check_2 = 0
    return check_2 == int(digits[13])


def normalize_phone(text: str) -> str:
    digits = only_digits(text)
    if len(digits) < 10:
        return ""
    if len(digits) > 11:
        digits = digits[-11:]
    if len(set(digits)) == 1:
        return ""
    return digits


def normalize_ddd(text: str) -> str:
    digits = only_digits(text)
    if len(digits) < 2:
        return ""
    return digits[:2]


def parse_ddd_filters(value: str) -> list[str]:
    ddds: list[str] = []
    for item in re.split(r"[,\s;]+", norm_text(value)):
        ddd = normalize_ddd(item)
        if ddd and ddd not in ddds:
            ddds.append(ddd)
    return ddds


def normalize_species(text: str) -> str:
    digits = only_digits(text)
    if not digits:
        return ""
    return digits.lstrip("0") or "0"


def normalize_species_list(values: list[str]) -> list[str]:
    normalized: list[str] = []
    for value in values:
        for item in re.split(r"[,\s;]+", norm_text(value)):
            species = normalize_species(item)
            if species and species not in normalized:
                normalized.append(species)
    return normalized


def parse_species_filter(value: str) -> list[str]:
    return normalize_species_list([value]) if value else []


def build_document_search_terms(raw_digits: str) -> list[str]:
    digits = only_digits(raw_digits)
    if not digits:
        return []
    variants: list[str] = []
    for candidate in (
        digits,
        digits.lstrip("0"),
        digits.zfill(11) if len(digits) <= 11 else "",
        digits.zfill(14) if len(digits) <= 14 else "",
    ):
        value = candidate.strip()
        if value and value not in variants:
            variants.append(value)
    return variants


def normalize_upper(text: str) -> str:
    return norm_text(text).upper()


def normalize_uf(text: str) -> str:
    value = normalize_upper(text)
    letters = re.sub(r"[^A-Z]", "", value)
    if len(letters) >= 2:
        return letters[:2]
    return letters


def clamp_occurrence_string_lengths(normalized: dict[str, str | Decimal | None]) -> None:
    for field, max_len in OCCURRENCE_STR_MAX_LENGTHS.items():
        value = normalized.get(field)
        if value is None:
            continue
        text = str(value)
        if len(text) > max_len:
            normalized[field] = text[:max_len]


def write_import_error_report(batch_id: int, rows: list[dict[str, str]]) -> Path:
    ensure_directories()
    target = settings.exports_dir / f"import_errors_batch_{batch_id}_{datetime.now(UTC):%Y%m%d_%H%M%S}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "erros_importacao"
    headers = [
        "BATCH_ID",
        "ARQUIVO",
        "LINHA_ARQUIVO",
        "ERRO",
        "CPF_CHAVE",
        "CPF_CNPJ_ORIGINAL",
        "NOME",
        "TELEFONE",
        "UF",
        "CIDADE",
        "DADOS_BRUTOS_JSON",
    ]
    ws.append(headers)
    for item in rows:
        ws.append([item.get(key, "") for key in headers])
    wb.save(target)
    return target


def normalize_token(text: object) -> str:
    return normalize_header_name(norm_text(text))


def parse_money(text: str) -> Decimal | None:
    cleaned = norm_text(text).replace("R$", "").replace(" ", "")
    if not cleaned:
        return None
    if "," in cleaned and "." in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    elif "," in cleaned:
        cleaned = cleaned.replace(".", "").replace(",", ".")
    elif "." in cleaned:
        parts = cleaned.split(".")
        if len(parts) > 2:
            cleaned = "".join(parts[:-1]) + "." + parts[-1]
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def parse_int(text: str) -> int | None:
    digits = only_digits(text)
    if not digits:
        return None
    try:
        return int(digits)
    except ValueError:
        return None


def get_normalized_row_value(row: dict[str, str], *aliases: str) -> str:
    normalized_pairs: list[tuple[str, str]] = []
    for key, value in row.items():
        cleaned_key = key.replace("EXTRA::", "", 1)
        normalized_pairs.append((normalize_token(cleaned_key), norm_text(value)))
    for alias in aliases:
        target = normalize_token(alias)
        for normalized_key, value in normalized_pairs:
            if normalized_key == target and value:
                return value
    return ""


def classify_margin_bucket(value: Decimal | None) -> str:
    if value is None or value <= Decimal("0"):
        return "MARGEM_0"
    if value <= Decimal("300"):
        return "MARGEM_BAIXA"
    return "MARGEM_LIVRE"


def classify_public_role(cargo: str, servico: str, entidade: str, convenio: str, situacao: str) -> str:
    context = " ".join([cargo, servico, entidade, convenio, situacao])
    normalized = normalize_token(context)
    if any(token in normalized for token in ("APOSENT", "PENSION", "INATIV")):
        return "APOSENTADO"
    if any(token in normalized for token in ("PROFESSOR", "DOCENTE", "EDUCACAO", "MAGISTERIO")):
        return "PROFESSOR"
    if any(token in normalized for token in ("SAUDE", "ENFERM", "MEDIC", "HOSPITAL", "ODONTO", "AGENTE DE SAUDE")):
        return "SAUDE"
    if any(token in normalized for token in ("GUARDA", "POLIC", "SEGURANCA", "BOMBEIR")):
        return "SEGURANCA"
    if any(token in normalized for token in ("ADMIN", "ASSISTENTE", "AUXILIAR", "ESCRITUR", "ATENDENTE")):
        return "ADMINISTRATIVO"
    return "SERVIDOR"


def build_public_campaign_name(
    segment: str,
    uf: str,
    cargo: str,
    servico: str,
    entidade: str,
    convenio: str,
    situacao: str,
    margem: Decimal | None,
) -> str:
    prefix = "GOV" if segment == "GOVERNO" else "PREF"
    role = classify_public_role(cargo, servico, entidade, convenio, situacao)
    margin_bucket = classify_margin_bucket(margem)
    return f"{prefix}_{(uf or 'BR').upper()}_{role}_{margin_bucket}"


def matches_public_campaign_profile(
    profile: str,
    campaign_name: str,
    cargo: str,
    servico: str,
    entidade: str,
    convenio: str,
    situacao: str,
    margem: Decimal | None,
) -> bool:
    normalized_profile = normalize_token(profile)
    if not normalized_profile or normalized_profile in {"TODOS", "LIVRE"}:
        return True
    role = classify_public_role(cargo, servico, entidade, convenio, situacao)
    margin_bucket = classify_margin_bucket(margem)
    if normalized_profile == "REFIN":
        return margin_bucket in {"MARGEM_0", "MARGEM_BAIXA"}
    return normalized_profile in {role, margin_bucket, normalize_token(campaign_name)}


def calculate_margin_from_rmi(vl_rmi: Decimal | None, esp: str) -> Decimal | None:
    if vl_rmi is None:
        return None
    rate = Decimal("0.30") if esp in {"87", "88"} else Decimal("0.35")
    return (vl_rmi * rate).quantize(Decimal("0.01"))


def calculate_rmc_from_rmi(vl_rmi: Decimal | None) -> Decimal | None:
    if vl_rmi is None:
        return None
    return (vl_rmi * Decimal("0.05")).quantize(Decimal("0.01"))


def resolve_updated_base_salary(row: dict[str, object]) -> Decimal | None:
    salary = row.get("salario")
    if isinstance(salary, Decimal):
        return salary
    benefit_value = row.get("vl_rmi")
    return benefit_value if isinstance(benefit_value, Decimal) else None


def resolve_updated_base_margin_total(row: dict[str, object]) -> Decimal | None:
    margin_total = row.get("margem_total")
    if isinstance(margin_total, Decimal):
        return margin_total
    margin = row.get("vl_margem")
    rmc = row.get("vl_rmc")
    if isinstance(margin, Decimal) and isinstance(rmc, Decimal):
        return (margin + rmc).quantize(Decimal("0.01"))
    return margin if isinstance(margin, Decimal) else None


def normalize_date_text(value: str) -> str:
    cleaned = norm_text(value)
    if not cleaned:
        return ""
    match = re.fullmatch(r"(\d{1,2}/\d{1,2}/\d{4})\s+\d{1,2}:\d{2}:\d{2}", cleaned)
    if match:
        cleaned = match.group(1)
    match = re.fullmatch(r"(\d{4}-\d{2}-\d{2})[ T]\d{1,2}:\d{2}:\d{2}", cleaned)
    if match:
        cleaned = match.group(1)
    match = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", cleaned)
    if match:
        day, month, year = match.groups()
        return f"{int(day):02d}/{int(month):02d}/{year}"
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", cleaned):
        year, month, day = cleaned.split("-")
        return f"{day}/{month}/{year}"
    try:
        serial = float(cleaned.replace(",", "."))
    except ValueError:
        return cleaned
    if serial <= 0:
        return cleaned
    excel_epoch = datetime(1899, 12, 30)
    parsed = excel_epoch + timedelta(days=serial)
    return parsed.strftime("%d/%m/%Y")


def normalize_date_filter(value: str) -> str:
    normalized = normalize_date_text(value)
    match = re.fullmatch(r"(\d{2})/(\d{2})/(\d{4})", normalized)
    if not match:
        return ""
    day, month, year = match.groups()
    return f"{year}-{month}-{day}"


def ddb_date_expr():
    return (
        func.substr(ClientOccurrence.ddb, 7, 4)
        + "-"
        + func.substr(ClientOccurrence.ddb, 4, 2)
        + "-"
        + func.substr(ClientOccurrence.ddb, 1, 2)
    )


def birth_date_expr():
    return (
        func.substr(ClientOccurrence.dt_nasc, 7, 4)
        + "-"
        + func.substr(ClientOccurrence.dt_nasc, 4, 2)
        + "-"
        + func.substr(ClientOccurrence.dt_nasc, 1, 2)
    )


def parse_date_value(value: str) -> date | None:
    normalized = normalize_date_text(value)
    match = re.fullmatch(r"(\d{2})/(\d{2})/(\d{4})", normalized)
    if not match:
        return None
    day, month, year = match.groups()
    try:
        return date(int(year), int(month), int(day))
    except ValueError:
        return None


def add_years_safe(value: date, years: int) -> date:
    try:
        return value.replace(year=value.year + years)
    except ValueError:
        return value.replace(month=2, day=28, year=value.year + years)


def compute_age_from_birth(value: str, today: date | None = None) -> int | None:
    birth_date = parse_date_value(value)
    if birth_date is None:
        return None
    reference = today or date.today()
    age = reference.year - birth_date.year
    if (reference.month, reference.day) < (birth_date.month, birth_date.day):
        age -= 1
    return age


def filter_rows_by_ddb_range(
    rows: list[tuple[ClientOccurrence, SourceFile, ImportBatch]],
    ddb_start: str,
    ddb_end: str,
) -> list[tuple[ClientOccurrence, SourceFile, ImportBatch]]:
    if not ddb_start and not ddb_end:
        return rows
    start_date = parse_date_value(ddb_start)
    end_date = parse_date_value(ddb_end)
    if not start_date and not end_date:
        return rows

    filtered: list[tuple[ClientOccurrence, SourceFile, ImportBatch]] = []
    for occurrence, source_file, batch in rows:
        occurrence_date = parse_date_value(occurrence.ddb)
        if occurrence_date is None:
            continue
        if start_date and occurrence_date < start_date:
            continue
        if end_date and occurrence_date > end_date:
            continue
        filtered.append((occurrence, source_file, batch))
    return filtered


def format_money(value: Decimal | None) -> str:
    if value is None:
        return ""
    quantized = value.quantize(Decimal("0.01"))
    integer, decimals = f"{quantized:.2f}".split(".")
    groups = []
    while integer:
        groups.insert(0, integer[-3:])
        integer = integer[:-3]
    return f"R$ {'.'.join(groups)},{decimals}"


def decode_text_file(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def validate_header_row(fieldnames: list[str] | None) -> list[str]:
    if not fieldnames:
        return ["Cabecalho ausente ou planilha vazia."]
    normalized_headers = {normalize_header_name(item) for item in fieldnames if norm_text(item)}
    issues: list[str] = []
    if not any(candidate in normalized_headers for candidate in ("CPF", "NU CPF", "NUCPF", "NUMERO CPF", "DOCUMENTO CPF")):
        issues.append("Coluna obrigatoria CPF nao encontrada.")
    if not any(
        candidate in normalized_headers
        for candidate in ("NOME", "NOME BENEFICIARIO", "NOME DO BENEFICIARIO", "NOME DO SERVIDOR", "SERVIDOR", "NOME HIGIENIZADO", "NAME")
    ):
        issues.append("Coluna de nome nao encontrada.")
    return issues


def build_header_alias_lookup() -> dict[str, str]:
    lookup: dict[str, str] = {}
    for target, aliases in HEADER_ALIASES.items():
        lookup[normalize_header_name(target)] = target
        for alias in aliases:
            lookup[normalize_header_name(alias)] = target
    # Algumas planilhas chegam com caracteres perdidos na exportacao do cabecalho.
    lookup.update(
        {
            "MATR CULA": "MATRICULA",
            "CONV NIO": "CONVENIO",
            "ENDERE O": "LOGRADOURO",
            "NOME_COMPLETO": "NOME",
            "NOME COMPLETO": "NOME",
            "CATEGORIA": "REGIME_CONTRATACAO",
            "REGIME DE CONTRATA O": "REGIME_CONTRATACAO",
            "TIPO": "REGIME_CONTRATACAO",
            "TITULO": "CARGO",
            "EMAIL1": "EMAIL",
            "EMAIL 1": "EMAIL",
            "EMAIL2": "EMAIL2",
            "EMAIL 2": "EMAIL2",
            "EMAIL3": "EMAIL3",
            "EMAIL 3": "EMAIL3",
        }
    )
    return lookup


HEADER_ALIAS_LOOKUP = build_header_alias_lookup()


def canonicalize_row_keys(row: dict[str, str]) -> dict[str, str]:
    canonical: dict[str, str] = {}
    dynamic_phone_columns: list[tuple[str, str]] = []
    for raw_key, raw_value in row.items():
        key = norm_text(raw_key)
        value = norm_text(raw_value)
        if not key:
            continue
        normalized_key = normalize_header_name(key)
        canonical_key = HEADER_ALIAS_LOOKUP.get(normalized_key)
        if canonical_key:
            if value or canonical_key not in canonical:
                canonical[canonical_key] = value
            continue
        if any(hint in normalized_key for hint in PHONE_HEADER_HINTS) and not normalized_key.startswith(("FLGWHATS", "PROCON")):
            dynamic_phone_columns.append((normalized_key, value))
            continue
        if value:
            canonical[f"EXTRA::{key.upper()}"] = value
    for _normalized_key, value in dynamic_phone_columns:
        if not value:
            continue
        if not canonical.get("TELEFONE1"):
            canonical["TELEFONE1"] = value
        elif not canonical.get("TELEFONE2") and canonical["TELEFONE1"] != value:
            canonical["TELEFONE2"] = value
        elif not canonical.get("TELEFONE3") and canonical["TELEFONE1"] != value and canonical.get("TELEFONE2") != value:
            canonical["TELEFONE3"] = value
    return canonical


def get_first_value(row: dict[str, str], *keys: str) -> str:
    for key in keys:
        value = norm_text(row.get(key, ""))
        if value:
            return value
    # Fallback para lidar com colunas extras com acentos/perda de charset no cabecalho.
    normalized_items = [(normalize_token(existing_key), norm_text(existing_value)) for existing_key, existing_value in row.items()]
    for key in keys:
        target = normalize_token(key)
        for normalized_key, normalized_value in normalized_items:
            if normalized_key == target and normalized_value:
                return normalized_value
    return ""


def normalize_contact_flag(value: str) -> str:
    normalized = normalize_header_name(value)
    if normalized in {"1", "S", "SIM", "TRUE", "VERDADEIRO", "COM WHATSAPP", "WHATSAPP", "WHATS"}:
        return "SIM"
    if normalized in {"0", "N", "NAO", "FALSE", "FALSO", "SEM WHATSAPP", "SEM WHATS"}:
        return "NAO"
    return norm_text(value).upper()


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    text = decode_text_file(path)
    sample = text[:4096]
    delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    return [canonicalize_row_keys({norm_text(k): norm_text(v) for k, v in row.items() if k is not None}) for row in reader]


def read_csv_rows_with_validation(path: Path) -> tuple[list[dict[str, str]], list[str]]:
    text = decode_text_file(path)
    sample = text[:4096]
    delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    issues = validate_header_row(reader.fieldnames)
    rows = [canonicalize_row_keys({norm_text(k): norm_text(v) for k, v in row.items() if k is not None}) for row in reader]
    if not rows:
        issues.append("Nenhuma linha de dados encontrada.")
    return rows, issues


def export_xlsx_to_csv(source_path: Path, temp_dir: Path) -> Path:
    target_path = temp_dir / f"{source_path.stem}.csv"
    powershell_script = rf"""
$ErrorActionPreference = 'Stop'
$excel = New-Object -ComObject Excel.Application
$excel.Visible = $false
$excel.DisplayAlerts = $false
try {{
  $wb = $excel.Workbooks.Open('{source_path}')
  if (Test-Path '{target_path}') {{ Remove-Item -LiteralPath '{target_path}' -Force }}
  $wb.SaveAs('{target_path}', 62)
  $wb.Close($false)
}} finally {{
  $excel.Quit()
  [System.Runtime.Interopservices.Marshal]::ReleaseComObject($excel) | Out-Null
}}
"""
    subprocess.run(["powershell", "-NoProfile", "-Command", powershell_script], check=True, capture_output=True, text=True)
    return target_path


def iter_rows_from_file(path: Path, temp_dir: Path) -> list[dict[str, str]]:
    if path.suffix.lower() == ".csv":
        return read_csv_rows(path)
    if path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
        return read_csv_rows(export_xlsx_to_csv(path, temp_dir))
    return []


def iter_rows_from_file_with_validation(path: Path, temp_dir: Path) -> tuple[list[dict[str, str]], list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv_rows_with_validation(path)
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        try:
            exported_csv = export_xlsx_to_csv(path, temp_dir)
        except subprocess.CalledProcessError as exc:
            return [], [f"Falha ao converter planilha Excel: {exc.stderr.strip() or exc.stdout.strip() or exc}"]
        except Exception as exc:
            return [], [f"Falha ao abrir planilha Excel: {exc}"]
        return read_csv_rows_with_validation(exported_csv)
    return [], [f"Extensao nao suportada: {suffix or 'sem extensao'}"]


def detect_phones(row: dict[str, str]) -> list[str]:
    phones: list[str] = []
    for key in ("CELULAR1", "CELULAR2", "CELULAR3", "TELEFONE1", "TELEFONE2", "TELEFONE3"):
        phone = normalize_phone(row.get(key, ""))
        if phone and phone not in phones:
            phones.append(phone)
    # Layout CREFAZ costuma vir com DDD separado de FONE.
    ddd = normalize_ddd(row.get("DDD", "") or row.get("EXTRA::DDD", ""))
    fone = normalize_phone(row.get("FONE", "") or row.get("EXTRA::FONE", ""))
    if ddd and len(ddd) == 2 and fone and len(fone) in {8, 9}:
        combined = normalize_phone(ddd + fone)
        if combined and combined not in phones:
            phones.insert(0, combined)
    return phones[:3]


def merge_phone_sequence(priority_phones: list[str], existing_phones: list[str]) -> list[str]:
    ordered: list[str] = []
    for candidate in [*priority_phones, *existing_phones]:
        phone = normalize_phone(candidate)
        if phone and phone not in ordered:
            ordered.append(phone)
    return ordered[:3]


def merge_extras_json(existing_payload: str, incoming_payload: str) -> str:
    existing = _parse_extras_json(existing_payload)
    incoming = _parse_extras_json(incoming_payload)
    changed = False
    for key, value in incoming.items():
        if value in (None, ""):
            continue
        if existing.get(key) != value:
            existing[key] = value
            changed = True
    if not existing:
        return ""
    if not changed and norm_text(existing_payload):
        return existing_payload
    return json.dumps(existing, ensure_ascii=False, sort_keys=True)


def _apply_string_update(current: str, incoming: str) -> tuple[str, bool]:
    if incoming and current != incoming:
        return incoming, True
    return current, False


def _apply_decimal_update(current: Decimal | None, incoming: Decimal | None) -> tuple[Decimal | None, bool]:
    if incoming is not None and current != incoming:
        return incoming, True
    return current, False


def apply_enrichment_to_occurrence(occurrence: ClientOccurrence, normalized: dict[str, str | Decimal | None], phones: list[str]) -> tuple[bool, bool]:
    updated_any_phone = False
    updated_any_data = False

    current_sequence = [occurrence.telefone1, occurrence.telefone2, occurrence.telefone3]
    next_sequence = merge_phone_sequence(phones, current_sequence)
    if next_sequence != current_sequence[:3]:
        updated_any_phone = True
    occurrence.telefone1 = next_sequence[0] if len(next_sequence) > 0 else ""
    occurrence.telefone2 = next_sequence[1] if len(next_sequence) > 1 else ""
    occurrence.telefone3 = next_sequence[2] if len(next_sequence) > 2 else ""

    string_fields = (
        "nome",
        "nome_higienizado",
        "cpf_original",
        "matricula",
        "servico",
        "situacao",
        "nu_nb",
        "pis",
        "entidade",
        "secretaria",
        "convenio",
        "cbo_titulo",
        "esp",
        "ddb",
        "uf",
        "cidade",
        "bairro",
        "logradouro",
        "numero",
        "complemento",
        "cep",
        "dt_nasc",
        "nome_mae",
        "sexo",
        "email",
    )
    decimal_fields = (
        "salario",
        "vl_rmi",
        "vl_margem",
        "vl_rmc",
        "margem_total",
        "margem_liquida",
        "margem_bruta",
        "margem_utilizada",
    )

    for field in string_fields:
        next_value, changed = _apply_string_update(getattr(occurrence, field), str(normalized.get(field, "") or ""))
        if changed:
            setattr(occurrence, field, next_value)
            updated_any_data = True

    for field in decimal_fields:
        next_value, changed = _apply_decimal_update(getattr(occurrence, field), normalized.get(field))  # type: ignore[arg-type]
        if changed:
            setattr(occurrence, field, next_value)
            updated_any_data = True

    merged_extras = merge_extras_json(occurrence.extras_json, str(normalized.get("extras_json", "") or ""))
    if merged_extras != occurrence.extras_json:
        occurrence.extras_json = merged_extras
        updated_any_data = True

    return updated_any_phone, updated_any_data


def materialize_public_update_occurrence(
    template: ClientOccurrence,
    aggregate: PublicCampaignAggregate,
    source_name: str,
    file_name: str,
) -> ClientOccurrence:
    row_hash = compute_row_hash(
        {
            "cpf": template.cpf,
            "base_segment": template.base_segment,
            "base_source": template.base_source,
            "matricula": aggregate.matricula,
            "servico": aggregate.servico_servidor,
            "source_file_id": str(template.source_file_id),
            "public_update_file": file_name,
        }
    )
    clone = ClientOccurrence(
        source_file_id=template.source_file_id,
        cpf=template.cpf,
        cpf_original=template.cpf_original or template.cpf,
        base_segment=template.base_segment,
        base_source=template.base_source,
        nome=template.nome,
        nome_higienizado=template.nome_higienizado or template.nome,
        matricula=aggregate.matricula or "",
        servico=template.servico,
        situacao=template.situacao,
        nu_nb=template.nu_nb,
        pis=template.pis,
        entidade=template.entidade,
        secretaria=template.secretaria,
        convenio=template.convenio,
        cbo_titulo=template.cbo_titulo,
        salario=template.salario,
        esp=template.esp,
        ddb=template.ddb,
        vl_rmi=template.vl_rmi,
        vl_margem=template.vl_margem,
        vl_rmc=template.vl_rmc,
        margem_total=template.margem_total,
        margem_liquida=template.margem_liquida,
        margem_bruta=template.margem_bruta,
        margem_utilizada=template.margem_utilizada,
        uf=template.uf,
        cidade=template.cidade,
        bairro=template.bairro,
        logradouro=template.logradouro,
        numero=template.numero,
        complemento=template.complemento,
        cep=template.cep,
        telefone1=template.telefone1,
        telefone2=template.telefone2,
        telefone3=template.telefone3,
        dt_nasc=template.dt_nasc,
        nome_mae=template.nome_mae,
        sexo=template.sexo,
        email=template.email,
        extras_json=template.extras_json,
        row_hash=row_hash,
    )
    payload = build_public_update_payload(aggregate, source_name, file_name)
    apply_enrichment_to_occurrence(clone, payload, [])
    return clone


def week_label_from_ddb(value: str) -> str:
    parsed = parse_date_value(value)
    if parsed is None:
        return ""
    day = parsed.day
    if day <= 10:
        return "01 a 10"
    if day <= 17:
        return "11 a 17"
    if day <= 24:
        return "18 a 24"
    return "25 a 31"


def detect_week_label(rows: list[dict[str, str]]) -> str:
    labels: list[str] = []
    for row in rows:
        label = week_label_from_ddb(get_first_value(row, "DDB"))
        if label and label not in labels:
            labels.append(label)
    if not labels:
        return "SEMANA INDEFINIDA"
    if len(labels) == 1:
        return labels[0]
    return " / ".join(labels)


def compute_file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while True:
            chunk = handle.read(1024 * 1024)
            if not chunk:
                break
            digest.update(chunk)
    return digest.hexdigest()


def compute_row_hash(payload: dict[str, str]) -> str:
    serialized = json.dumps(payload, ensure_ascii=True, sort_keys=True)
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()


def normalize_row(row: dict[str, str]) -> dict[str, str | Decimal | None]:
    phones = detect_phones(row)
    cpf_original = normalize_document_original(get_first_value(row, "CPF"))
    nome_higienizado = get_first_value(row, "NOME_HIGIENIZADO").upper()
    nome_base = get_first_value(row, "NOME", "NOME BENEFICIARIO", "NOME DO SERVIDOR", "SERVIDOR", "EXTRA::NAME").upper() or nome_higienizado
    esp = normalize_species(get_first_value(row, "ESP", "ESPECIE"))
    vl_rmi = parse_money(get_first_value(row, "VL-RMI", "RMI", "SALARIO", "VALOR DO BENEFICIO", "MARGEM TOTAL"))
    vl_margem = parse_money(get_first_value(row, "VL MARGEM", "MARGEM", "MARGEM DISPONIVEL", "VALOR MARGEM", "EXTRA::MARGEM_DISPONIVEL"))
    vl_rmc = parse_money(get_first_value(row, "VL RMC", "RMC", "MARGEM CARTAO", " VL RMC "))
    salario = parse_money(get_first_value(row, "SALARIO", "VL-RMI", "RMI"))
    margem_liquida = parse_money(get_first_value(row, "MARGEM_LIQUIDA"))
    margem_bruta = parse_money(get_first_value(row, "MARGEM_BRUTA"))
    margem_utilizada = parse_money(get_first_value(row, "MARGEM_UTILIZADA"))
    servico = normalize_upper(get_first_value(row, "SERVICO", "EXTRA::SERVICO", "EXTRA::SERVIÃ‡O"))
    situacao = normalize_upper(get_first_value(row, "SITUACAO", "EXTRA::SITUACAO", "EXTRA::SITUAÃ‡ÃƒO", "EXTRA::SITUACA", "EXTRA::STATUS"))
    cbo_titulo = normalize_upper(
        get_first_value(row, "CBO_TITULO", "EXTRA::CBO_TITULO", "EXTRA::CBO TITULO", "EXTRA::CBO_TIT", "EXTRA::CBO TIT")
    )
    margem_total = parse_money(get_first_value(row, "MARGEM_TOTAL", "EXTRA::MARGEM TOTAL", "EXTRA::MARGEM_TOTAL", "EXTRA::MARGEM TOTAL (R$)"))
    if margem_total is None and servico and vl_rmi is not None and not get_first_value(row, "SALARIO", "RMI", "VALOR DO BENEFICIO"):
        margem_total = vl_rmi
    if vl_margem is None:
        vl_margem = calculate_margin_from_rmi(vl_rmi, esp)
    if vl_rmc is None:
        vl_rmc = calculate_rmc_from_rmi(vl_rmi)
    extras = {key.replace("EXTRA::", "", 1): value for key, value in row.items() if key.startswith("EXTRA::") and value}
    tipo_vinculo = normalize_upper(get_first_value(row, "REGIME_CONTRATACAO", "EXTRA::REGIME DE CONTRATACAO", "EXTRA::TIPO DE VINCULO", "EXTRA::TIPO_VINCULO", "EXTRA::VINCULO"))
    cargo_funcao = normalize_upper(get_first_value(row, "CARGO", "EXTRA::CARGO", "EXTRA::FUNCAO", "EXTRA::FUNÃ‡ÃƒO", "EXTRA::CARGO/FUNCAO", "EXTRA::CARGO FUNCAO"))
    data_admissao = normalize_date_text(
        get_first_value(
            row,
            "EXTRA::DATA ADMISSAO",
            "EXTRA::DT ADMISSAO",
            "EXTRA::DT_ADMISSAO",
            "EXTRA::ADMISSAO",
            "EXTRA::DATA_ADMISSAO",
        )
    )
    if tipo_vinculo:
        extras["REGIME_CONTRATACAO"] = tipo_vinculo
    if cargo_funcao:
        extras["CARGO"] = cargo_funcao
    if data_admissao:
        extras["DATA ADMISSAO"] = data_admissao
    whatsapp_telefone1 = normalize_contact_flag(get_first_value(row, "EXTRA::FLGWHATSCEL1"))
    whatsapp_telefone2 = normalize_contact_flag(get_first_value(row, "EXTRA::FLGWHATSCEL2"))
    email2 = get_first_value(row, "EMAIL2", "EXTRA::EMAIL2", "EXTRA::EMAIL_02").lower()
    email3 = get_first_value(row, "EMAIL3", "EXTRA::EMAIL3", "EXTRA::EMAIL_03").lower()
    if whatsapp_telefone1:
        extras["FLGWHATSCEL1"] = whatsapp_telefone1
    if whatsapp_telefone2:
        extras["FLGWHATSCEL2"] = whatsapp_telefone2
    if email2:
        extras["EMAIL2"] = email2
    if email3:
        extras["EMAIL3"] = email3
    return {
        "cpf": normalize_cpf(get_first_value(row, "CPF")),
        "cpf_original": cpf_original,
        "nome": nome_base,
        "nome_higienizado": nome_higienizado,
        "matricula": norm_text(get_first_value(row, "MATRICULA", "EXTRA::MATRIC")),
        "servico": servico,
        "situacao": situacao,
        "nu_nb": only_digits(get_first_value(row, "NU-NB", "NB")),
        "pis": only_digits(get_first_value(row, "PIS")),
        "entidade": normalize_upper(get_first_value(row, "ENTIDADE")),
        "secretaria": normalize_upper(get_first_value(row, "SECRETARIA")),
        "convenio": normalize_upper(get_first_value(row, "CONVENIO")),
        "cbo_titulo": cbo_titulo,
        "esp": esp,
        "ddb": normalize_date_text(get_first_value(row, "DDB")),
        "salario": salario,
        "vl_rmi": vl_rmi,
        "vl_margem": vl_margem,
        "vl_rmc": vl_rmc,
        "margem_total": margem_total,
        "margem_liquida": margem_liquida,
        "margem_bruta": margem_bruta,
        "margem_utilizada": margem_utilizada,
        "uf": normalize_uf(get_first_value(row, "UF")),
        "cidade": normalize_upper(get_first_value(row, "CIDADE")),
        "bairro": normalize_upper(get_first_value(row, "BAIRRO")),
        "logradouro": normalize_upper(get_first_value(row, "LOGRADOURO")),
        "numero": get_first_value(row, "NUMERO"),
        "complemento": normalize_upper(get_first_value(row, "COMPLEMENTO")),
        "cep": only_digits(get_first_value(row, "CEP")),
        "telefone1": phones[0] if len(phones) > 0 else "",
        "telefone2": phones[1] if len(phones) > 1 else "",
        "telefone3": phones[2] if len(phones) > 2 else "",
        "dt_nasc": normalize_date_text(get_first_value(row, "DT-NASC", "DT_NASCIMENTO", "DT_NASC")),
        "nome_mae": normalize_upper(get_first_value(row, "NOME_MAE")),
        "sexo": normalize_upper(get_first_value(row, "SEXO")),
        "email": get_first_value(row, "EMAIL", "EXTRA::EMAIL1", "EXTRA::EMAIL_01").lower(),
        "extras_json": json.dumps(extras, ensure_ascii=False, sort_keys=True) if extras else "",
    }


def parse_extras_json(payload: str) -> dict[str, object]:
    if not payload:
        return {}
    try:
        data = json.loads(payload)
    except json.JSONDecodeError:
        return {}
    return data if isinstance(data, dict) else {}


def update_normalized_extras(normalized: dict[str, str | Decimal | None], updates: dict[str, object]) -> None:
    extras = parse_extras_json(str(normalized.get("extras_json") or ""))
    extras.update({key: value for key, value in updates.items() if value not in ("", None, (), [])})
    normalized["extras_json"] = json.dumps(extras, ensure_ascii=False, sort_keys=True) if extras else ""


def needs_cep_enrichment(normalized: dict[str, str | Decimal | None]) -> bool:
    cep = str(normalized.get("cep") or "")
    if len(cep) != 8:
        return False
    return any(not str(normalized.get(field) or "") for field in ("uf", "cidade", "bairro", "logradouro"))


def enrich_normalized_row_with_cep(
    normalized: dict[str, str | Decimal | None],
    cep_cache: dict[str, object | None],
) -> dict[str, str | Decimal | None]:
    if not settings.enable_cep_enrichment or not needs_cep_enrichment(normalized):
        return normalized

    cep = str(normalized.get("cep") or "")
    if len(cep) != 8:
        return normalized
    if cep not in cep_cache:
        cep_cache[cep] = lookup_cep_address(cep)
    address = cep_cache.get(cep)
    if address is None:
        return normalized

    field_map = {
        "cep": only_digits(getattr(address, "cep", "")),
        "uf": normalize_uf(getattr(address, "uf", "")),
        "cidade": normalize_upper(getattr(address, "cidade", "")),
        "bairro": normalize_upper(getattr(address, "bairro", "")),
        "logradouro": normalize_upper(getattr(address, "logradouro", "")),
        "complemento": normalize_upper(getattr(address, "complemento", "")),
    }
    for field, candidate in field_map.items():
        if candidate and not str(normalized.get(field) or ""):
            normalized[field] = candidate
    return normalized


def first_phone_ddd(normalized: dict[str, str | Decimal | None]) -> str:
    for field in ("telefone1", "telefone2", "telefone3"):
        phone = normalize_phone(str(normalized.get(field) or ""))
        if phone:
            return phone[:2]
    return ""


def enrich_normalized_row_with_ddd(
    normalized: dict[str, str | Decimal | None],
    ddd_cache: dict[str, object | None],
) -> dict[str, str | Decimal | None]:
    if not settings.enable_ddd_enrichment:
        return normalized

    ddd = first_phone_ddd(normalized)
    if len(ddd) != 2:
        return normalized

    if ddd not in ddd_cache:
        ddd_cache[ddd] = lookup_brasilapi_ddd(ddd)
    ddd_info = ddd_cache.get(ddd)
    if ddd_info is None:
        return normalized

    ddd_uf = normalize_uf(getattr(ddd_info, "uf", ""))
    ddd_cities = tuple(normalize_upper(city) for city in getattr(ddd_info, "cidades", ()) if city)
    updates: dict[str, object] = {
        "DDD_TELEFONE": ddd,
        "DDD_UF_BRASILAPI": ddd_uf,
    }
    if ddd_cities:
        updates["DDD_CIDADES_BRASILAPI"] = list(ddd_cities)

    current_uf = normalize_uf(str(normalized.get("uf") or ""))
    if not current_uf and ddd_uf:
        normalized["uf"] = ddd_uf
        updates["UF_INFERIDA_POR_DDD"] = ddd_uf
    elif current_uf and ddd_uf and current_uf != ddd_uf:
        updates["ALERTA_REGIONAL"] = f"DDD {ddd} indica {ddd_uf}, mas UF informada e {current_uf}"

    update_normalized_extras(normalized, updates)
    return normalized


def ensure_directories() -> None:
    settings.uploads_dir.mkdir(parents=True, exist_ok=True)
    settings.archive_dir.mkdir(parents=True, exist_ok=True)
    settings.exports_dir.mkdir(parents=True, exist_ok=True)


def save_uploaded_files(files: list[tuple[str, bytes]]) -> list[Path]:
    ensure_directories()
    stamp_dir = settings.uploads_dir / datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    stamp_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for name, content in files:
        path = stamp_dir / Path(name).name
        path.write_bytes(content)
        paths.append(path)
    return paths


def create_import_batch(session: Session, request: ImportRequest) -> ImportBatch:
    with serialized_write():
        batch = ImportBatch(
            mes_referencia=request.month,
            ano_referencia=request.year,
            origem_pasta=request.origin_folder,
            base_segment=normalize_base_segment(request.base_segment),
            base_source=normalize_upper(request.base_source),
            usuario=request.user_name or "operador",
            status="PROCESSANDO",
        )
        session.add(batch)
        session.flush()
        return batch


def latest_reference_label(year: int, month: int) -> str:
    return f"{year:04d}-{month:02d}"


def prefer_non_empty(current: str, candidate: str) -> str:
    return candidate if candidate else current


def build_normalized_from_occurrence(occurrence: ClientOccurrence) -> dict[str, str | Decimal | None]:
    return {
        "nome": occurrence.nome or "",
        "nome_higienizado": occurrence.nome_higienizado or "",
        "cpf_original": occurrence.cpf_original or "",
        "matricula": occurrence.matricula or "",
        "servico": occurrence.servico or "",
        "situacao": occurrence.situacao or "",
        "nu_nb": occurrence.nu_nb or "",
        "pis": occurrence.pis or "",
        "entidade": occurrence.entidade or "",
        "secretaria": occurrence.secretaria or "",
        "convenio": occurrence.convenio or "",
        "cbo_titulo": occurrence.cbo_titulo or "",
        "esp": occurrence.esp or "",
        "ddb": occurrence.ddb or "",
        "uf": occurrence.uf or "",
        "cidade": occurrence.cidade or "",
        "bairro": occurrence.bairro or "",
        "logradouro": occurrence.logradouro or "",
        "numero": occurrence.numero or "",
        "complemento": occurrence.complemento or "",
        "cep": occurrence.cep or "",
        "telefone1": occurrence.telefone1 or "",
        "telefone2": occurrence.telefone2 or "",
        "telefone3": occurrence.telefone3 or "",
        "dt_nasc": occurrence.dt_nasc or "",
        "nome_mae": occurrence.nome_mae or "",
        "sexo": occurrence.sexo or "",
        "email": occurrence.email or "",
        "salario": occurrence.salario,
        "vl_rmi": occurrence.vl_rmi,
        "vl_margem": occurrence.vl_margem,
        "vl_rmc": occurrence.vl_rmc,
        "margem_total": occurrence.margem_total,
        "margem_liquida": occurrence.margem_liquida,
        "margem_bruta": occurrence.margem_bruta,
        "margem_utilizada": occurrence.margem_utilizada,
        "extras_json": occurrence.extras_json or "",
    }


def chunked_values(values: set[str] | list[str], chunk_size: int = SQLITE_VARIABLE_CHUNK_SIZE) -> list[list[str]]:
    sequence = list(values)
    return [sequence[index:index + chunk_size] for index in range(0, len(sequence), chunk_size)]


def rebuild_clients_for_cpfs(session: Session, cpfs: set[str]) -> None:
    normalized_cpfs = {cpf for cpf in cpfs if cpf}
    if not normalized_cpfs:
        return

    preserved_status: dict[str, tuple[bool, str]] = {}
    for cpf_chunk in chunked_values(normalized_cpfs):
        current_rows = session.execute(select(Client.cpf, Client.do_not_call, Client.do_not_call_reason).where(Client.cpf.in_(cpf_chunk))).all()
        for current_cpf, do_not_call, do_not_call_reason in current_rows:
            preserved_status[current_cpf] = (bool(do_not_call), do_not_call_reason or "")

    for cpf_chunk in chunked_values(normalized_cpfs):
        session.execute(delete(Client).where(Client.cpf.in_(cpf_chunk)))
    session.flush()
    session.expire_all()

    grouped: dict[str, Client] = {}
    for cpf_chunk in chunked_values(normalized_cpfs):
        rows = session.execute(
            select(
                ClientOccurrence.cpf,
                ClientOccurrence.base_segment,
                ClientOccurrence.nome,
                ClientOccurrence.uf,
                ClientOccurrence.cidade,
                ClientOccurrence.pis,
                ClientOccurrence.matricula,
                ClientOccurrence.entidade,
                ClientOccurrence.telefone1,
                ClientOccurrence.telefone2,
                ClientOccurrence.telefone3,
                ClientOccurrence.vl_margem,
                ImportBatch.ano_referencia,
                ImportBatch.mes_referencia,
            )
            .join(SourceFile, ClientOccurrence.source_file_id == SourceFile.id)
            .join(ImportBatch, SourceFile.batch_id == ImportBatch.id)
            .where(ClientOccurrence.cpf.in_(cpf_chunk))
            .order_by(
                ClientOccurrence.cpf,
                ImportBatch.ano_referencia,
                ImportBatch.mes_referencia,
                SourceFile.id,
                ClientOccurrence.id,
            )
            .execution_options(yield_per=1000)
        )
        for row in rows:
            row_segment = normalize_base_segment(row.base_segment)
            client = grouped.get(row.cpf)
            if client is None:
                do_not_call, do_not_call_reason = preserved_status.get(row.cpf, (False, ""))
                client = Client(
                    cpf=row.cpf,
                    nome_atual=row.nome,
                    uf_atual=row.uf,
                    cidade_atual=row.cidade,
                    melhor_telefone=row.telefone1 or row.telefone2 or row.telefone3,
                    tem_telefone=bool(row.telefone1 or row.telefone2 or row.telefone3),
                    do_not_call=do_not_call,
                    do_not_call_reason=do_not_call_reason,
                    tem_inss=row_segment == "INSS",
                    tem_governo=row_segment in PUBLIC_BASE_SEGMENTS,
                    pis_atual=row.pis or "",
                    matricula_atual=row.matricula or "",
                    entidade_atual=row.entidade or "",
                    maior_margem=row.vl_margem,
                    ultima_referencia=latest_reference_label(row.ano_referencia, row.mes_referencia),
                    qtd_ocorrencias=0,
                )
                grouped[row.cpf] = client
            client.qtd_ocorrencias += 1
            client.nome_atual = prefer_non_empty(client.nome_atual, row.nome)
            client.uf_atual = prefer_non_empty(client.uf_atual, row.uf)
            client.cidade_atual = prefer_non_empty(client.cidade_atual, row.cidade)
            client.pis_atual = prefer_non_empty(client.pis_atual, row.pis or "")
            client.matricula_atual = prefer_non_empty(client.matricula_atual, row.matricula or "")
            client.entidade_atual = prefer_non_empty(client.entidade_atual, row.entidade or "")
            if row_segment == "INSS":
                client.tem_inss = True
            if row_segment in PUBLIC_BASE_SEGMENTS:
                client.tem_governo = True
            for candidate in (row.telefone1, row.telefone2, row.telefone3):
                if candidate:
                    client.melhor_telefone = candidate
                    client.tem_telefone = True
                    break
            if row.vl_margem is not None and (client.maior_margem is None or row.vl_margem > client.maior_margem):
                client.maior_margem = row.vl_margem
            client.ultima_referencia = latest_reference_label(row.ano_referencia, row.mes_referencia)
    for client in grouped.values():
        session.merge(client)
    if grouped:
        session.flush()


def backfill_missing_addresses(session: Session, batch_size: int | None = None) -> AddressBackfillSummary:
    summary = AddressBackfillSummary()
    cep_cache: dict[str, object | None] = {}
    ddd_cache: dict[str, object | None] = {}
    limit = max(batch_size or settings.startup_address_backfill_batch_size, 1)
    last_seen_id = 0
    has_any_phone = or_(
        ClientOccurrence.telefone1 != "",
        ClientOccurrence.telefone2 != "",
        ClientOccurrence.telefone3 != "",
    )
    needs_cep_backfill = and_(
        ClientOccurrence.cep != "",
        or_(
            ClientOccurrence.uf == "",
            ClientOccurrence.cidade == "",
            ClientOccurrence.logradouro == "",
        ),
    )
    needs_ddd_backfill = and_(
        ClientOccurrence.uf == "",
        has_any_phone,
    )

    summary.skipped_invalid_cep = session.scalar(
        select(func.count())
        .select_from(ClientOccurrence)
        .where(
            ClientOccurrence.cep != "",
            func.length(ClientOccurrence.cep) != 8,
            needs_cep_backfill,
        )
    ) or 0

    while True:
        candidates = session.execute(
            select(ClientOccurrence)
            .where(
                ClientOccurrence.id > last_seen_id,
                or_(
                    and_(needs_cep_backfill, func.length(ClientOccurrence.cep) == 8),
                    needs_ddd_backfill,
                ),
            )
            .order_by(ClientOccurrence.id)
            .limit(limit)
        ).scalars().all()
        if not candidates:
            break
        last_seen_id = candidates[-1].id

        plans: list[AddressBackfillPlan] = []
        for occurrence in candidates:
            summary.scanned_rows += 1
            normalized = build_normalized_from_occurrence(occurrence)
            normalized = enrich_normalized_row_with_cep(normalized, cep_cache)
            normalized = enrich_normalized_row_with_ddd(normalized, ddd_cache)
            plans.append(
                AddressBackfillPlan(
                    occurrence_id=occurrence.id,
                    cpf=occurrence.cpf,
                    phones=[occurrence.telefone1 or "", occurrence.telefone2 or "", occurrence.telefone3 or ""],
                    normalized=normalized,
                )
            )

        with serialized_write():
            batch_cpfs: set[str] = set()
            for plan in plans:
                occurrence = session.get(ClientOccurrence, plan.occurrence_id)
                if occurrence is None:
                    continue
                _updated_phone, updated_data = apply_enrichment_to_occurrence(occurrence, plan.normalized, plan.phones)
                if updated_data:
                    summary.updated_rows += 1
                    batch_cpfs.add(plan.cpf)

            if batch_cpfs:
                rebuild_clients_for_cpfs(session, batch_cpfs)
                summary.rebuilt_clients += len(batch_cpfs)
            session.commit()

        if len(candidates) < limit:
            break

    return summary


def delete_batch(session: Session, batch_id: int) -> tuple[bool, str]:
    affected_cpfs: set[str] = set()
    archived_file_paths: list[str] = []
    with serialized_write():
        batch = session.get(ImportBatch, batch_id)
        if batch is None:
            return False, "Lote nao encontrado."
        if batch.status == "PROCESSANDO":
            return False, "Lote em processamento. Aguarde a conclusao para apagar."

        source_files = session.execute(
            select(SourceFile.id, SourceFile.caminho_arquivo).where(SourceFile.batch_id == batch_id)
        ).all()
        source_file_ids = [row.id for row in source_files]
        archived_file_paths = [str(row.caminho_arquivo) for row in source_files if row.caminho_arquivo]
        affected_cpfs = set(
            session.scalars(
                select(ClientOccurrence.cpf)
                .where(ClientOccurrence.source_file_id.in_(source_file_ids))
                .distinct()
            ).all()
        ) if source_file_ids else set()

        if source_file_ids:
            session.execute(delete(ClientOccurrence).where(ClientOccurrence.source_file_id.in_(source_file_ids)))
            session.execute(delete(SourceFile).where(SourceFile.id.in_(source_file_ids)))
        session.execute(delete(ImportBatch).where(ImportBatch.id == batch_id))
        session.commit()

    def _cleanup_after_delete(cpfs: set[str], file_paths: list[str]) -> None:
        cleanup_session = SessionLocal()
        try:
            if cpfs:
                rebuild_clients_for_cpfs(cleanup_session, cpfs)
                cleanup_session.commit()
        finally:
            cleanup_session.close()
        for raw_path in file_paths:
            path = Path(raw_path)
            if path.exists():
                path.unlink(missing_ok=True)

    threading.Thread(
        target=_cleanup_after_delete,
        args=(set(affected_cpfs), list(archived_file_paths)),
        daemon=True,
    ).start()

    return True, "Campanha apagada com sucesso."


def process_import_batch(session: Session, batch: ImportBatch, request: ImportRequest) -> ImportSummary:
    ensure_directories()
    summary = ImportSummary(batch_id=batch.id)
    temp_dir = Path(tempfile.mkdtemp(prefix="inss_import_"))
    allowed_species = set(normalize_species_list(request.allowed_species))
    progress_flush_rows = 5000
    insert_chunk_size = 2000

    def update_batch_progress(message: str) -> None:
        current_batch = session.get(ImportBatch, batch.id)
        if current_batch is not None:
            current_batch.resumo = message
            session.commit()

    def summary_text() -> str:
        return (
            f"Arquivos importados: {summary.files_imported}; ignorados: {summary.files_skipped}; "
            f"linhas: {summary.rows_imported}; invalidas: {summary.invalid_rows}; "
            f"duplicadas: {summary.duplicate_rows}; novos CPFs: {summary.cpfs_new}; "
            f"CPFs atualizados: {summary.cpfs_updated}; sem telefone: {summary.rows_without_phone}; "
            f"sem cidade: {summary.rows_without_city}; fora da especie: {summary.species_filtered_rows}"
        )

    try:
        cep_cache: dict[str, object | None] = {}
        ddd_cache: dict[str, object | None] = {}
        cpf_existence_cache: dict[str, bool] = {}
        new_cpfs_seen: set[str] = set()
        updated_cpfs_seen: set[str] = set()
        row_errors: list[dict[str, str]] = []
        total_files = len(request.files)
        for file_index, file_path in enumerate(request.files, start=1):
            update_batch_progress(f"Preparando arquivo {file_index}/{total_files}: {file_path.name}")
            rows, issues = iter_rows_from_file_with_validation(file_path, temp_dir)
            if issues:
                summary.invalid_files.append(file_path.name)
                summary.validation_errors.extend(f"{file_path.name}: {issue}" for issue in issues)
                if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS or not rows:
                    continue
            file_hash = compute_file_hash(file_path)

            with serialized_write():
                if session.scalar(select(func.count()).select_from(SourceFile).where(SourceFile.hash_arquivo == file_hash)):
                    summary.files_skipped += 1
                    summary.duplicate_files.append(file_path.name)
                    update_batch_progress(
                        f"Arquivo duplicado ignorado ({file_index}/{total_files}): {file_path.name}. {summary_text()}"
                    )
                    continue

                archive_path = settings.archive_dir / f"{datetime.now(UTC):%Y%m%d_%H%M%S}_{file_path.name}"
                shutil.copy2(file_path, archive_path)
                source_file = SourceFile(
                    batch_id=batch.id,
                    nome_arquivo=file_path.name,
                    caminho_arquivo=str(archive_path),
                    base_segment=batch.base_segment,
                    base_source=batch.base_source,
                    semana_label=detect_week_label(rows),
                    tipo_arquivo=file_path.suffix.lower().lstrip("."),
                    hash_arquivo=file_hash,
                    total_linhas=len(rows),
                    duplicado=False,
                )
                session.add(source_file)
                session.flush()

                seen_row_hashes: set[str] = set()
                file_affected_cpfs: set[str] = set()
                is_crefaz_import = normalize_base_segment(batch.base_segment) == "CREFAZ"
                staged_occurrences: list[dict[str, object]] = []
                staged_context: list[tuple[int, dict[str, object]]] = []

                def flush_staged_occurrences() -> None:
                    if not staged_occurrences:
                        return

                    unresolved_cpfs = sorted(
                        {
                            str(payload.get("cpf") or "")
                            for payload in staged_occurrences
                            if str(payload.get("cpf") or "") and str(payload.get("cpf") or "") not in cpf_existence_cache
                        }
                    )
                    for cpf_chunk in chunked_values(unresolved_cpfs):
                        existing_cpfs = set(
                            session.scalars(
                                select(ClientOccurrence.cpf).where(ClientOccurrence.cpf.in_(cpf_chunk)).distinct()
                            ).all()
                        )
                        for cpf in cpf_chunk:
                            cpf_existence_cache[cpf] = cpf in existing_cpfs

                    for payload in staged_occurrences:
                        cpf_value = str(payload.get("cpf") or "")
                        if not cpf_value:
                            continue
                        if cpf_existence_cache.get(cpf_value, False):
                            if cpf_value not in updated_cpfs_seen:
                                summary.cpfs_updated += 1
                                updated_cpfs_seen.add(cpf_value)
                        else:
                            if cpf_value not in new_cpfs_seen:
                                summary.cpfs_new += 1
                                new_cpfs_seen.add(cpf_value)

                    try:
                        session.execute(insert(ClientOccurrence), staged_occurrences)
                        session.flush()
                        for payload in staged_occurrences:
                            cpf_value = str(payload.get("cpf") or "")
                            if cpf_value:
                                file_affected_cpfs.add(cpf_value)
                            summary.rows_imported += 1
                            cpf_existence_cache[cpf_value] = True
                    except SQLAlchemyError:
                        session.rollback()
                        for payload, (row_index, raw_row) in zip(staged_occurrences, staged_context):
                            try:
                                with session.begin_nested():
                                    session.add(ClientOccurrence(**payload))
                                    session.flush()
                            except SQLAlchemyError as exc:
                                summary.invalid_rows += 1
                                summary.row_errors += 1
                                row_errors.append(
                                    {
                                        "BATCH_ID": str(batch.id),
                                        "ARQUIVO": file_path.name,
                                        "LINHA_ARQUIVO": str(row_index),
                                        "ERRO": str(exc)[:1000],
                                        "CPF_CHAVE": str(payload.get("cpf") or ""),
                                        "CPF_CNPJ_ORIGINAL": str(payload.get("cpf_original") or ""),
                                        "NOME": str(payload.get("nome") or ""),
                                        "TELEFONE": str(payload.get("telefone1") or payload.get("telefone2") or payload.get("telefone3") or ""),
                                        "UF": str(payload.get("uf") or ""),
                                        "CIDADE": str(payload.get("cidade") or ""),
                                        "DADOS_BRUTOS_JSON": json.dumps(raw_row, ensure_ascii=False)[:30000],
                                    }
                                )
                                continue
                            cpf_value = str(payload.get("cpf") or "")
                            if cpf_value:
                                file_affected_cpfs.add(cpf_value)
                                cpf_existence_cache[cpf_value] = True
                            summary.rows_imported += 1

                    staged_occurrences.clear()
                    staged_context.clear()

                for row_index, raw_row in enumerate(rows, start=1):
                    normalized = normalize_row(raw_row)
                    if not is_crefaz_import:
                        normalized = enrich_normalized_row_with_cep(normalized, cep_cache)
                        normalized = enrich_normalized_row_with_ddd(normalized, ddd_cache)
                    normalized["base_segment"] = normalize_base_segment(batch.base_segment)
                    normalized["base_source"] = batch.base_source
                    normalized["uf"] = normalize_uf(str(normalized.get("uf") or ""))
                    clamp_occurrence_string_lengths(normalized)
                    if normalized["base_segment"] != "INSS":
                        if not get_first_value(raw_row, "SALARIO"):
                            normalized["salario"] = None
                            normalized["vl_rmi"] = None
                        if not get_first_value(raw_row, "VL MARGEM"):
                            normalized["vl_margem"] = None
                        normalized["vl_rmc"] = None
                    if not normalized["cpf"]:
                        summary.invalid_rows += 1
                        continue
                    if allowed_species and str(normalized["esp"]) not in allowed_species:
                        summary.species_filtered_rows += 1
                        continue
                    segment_value = str(normalized["base_segment"])
                    hash_payload = {
                        "source_file": source_file.nome_arquivo,
                        "cpf": str(normalized["cpf"]),
                        "matricula": str(normalized["matricula"]),
                        "servico": str(normalized["servico"]),
                        "nu_nb": str(normalized["nu_nb"]),
                        "vl_margem": str(normalized["vl_margem"]),
                        "margem_total": str(normalized["margem_total"]),
                        "telefone1": str(normalized["telefone1"]),
                        "telefone2": str(normalized["telefone2"]),
                        "telefone3": str(normalized["telefone3"]),
                    }
                    if segment_value == "CREFAZ":
                        extras = parse_extras_json(str(normalized.get("extras_json") or ""))
                        numero_cliente = norm_text(
                            str(
                                extras.get("NÂº DO CLIENTE")
                                or extras.get("N DO CLIENTE")
                                or extras.get("NRO CLIENTE")
                                or extras.get("NUMERO CLIENTE")
                                or ""
                            )
                        )
                        conta_contrato = norm_text(str(extras.get("CONTA CONTRATO") or extras.get("CONTA_CONTRATO") or ""))
                        numero_instalacao = norm_text(
                            str(
                                extras.get("NUMERO INSTALAÃ‡ÃƒO")
                                or extras.get("NUMERO INSTALACAO")
                                or extras.get("NÂº INSTALACAO")
                                or extras.get("N INSTALACAO")
                                or ""
                            )
                        )
                        fatura_atual = norm_text(str(extras.get("FATURA ATUAL") or ""))
                        # Para CREFAZ, a mesma pessoa pode ter multiplos imoveis.
                        # Entao a deduplicacao precisa considerar identificadores da unidade.
                        hash_payload = {
                            "doc_original": str(normalized["cpf_original"] or normalized["cpf"]),
                            "cpf": str(normalized["cpf"]),
                            "nome": str(normalized["nome"]),
                            "numero_cliente": numero_cliente,
                            "conta_contrato": conta_contrato,
                            "numero_instalacao": numero_instalacao,
                            "fatura_atual": fatura_atual,
                            "cidade": str(normalized["cidade"]),
                            "logradouro": str(normalized["logradouro"]),
                            "numero": str(normalized["numero"]),
                            "cep": str(normalized["cep"]),
                            "telefone1": str(normalized["telefone1"]),
                            "telefone2": str(normalized["telefone2"]),
                            "telefone3": str(normalized["telefone3"]),
                        }
                    row_hash = compute_row_hash(hash_payload)
                    if row_hash in seen_row_hashes:
                        summary.invalid_rows += 1
                        summary.duplicate_rows += 1
                        continue
                    seen_row_hashes.add(row_hash)
                    cpf_value = str(normalized["cpf"])
                    if not any(str(normalized[field]) for field in ("telefone1", "telefone2", "telefone3")):
                        summary.rows_without_phone += 1
                    if not str(normalized["cidade"]):
                        summary.rows_without_city += 1
                    staged_occurrences.append({"source_file_id": source_file.id, "row_hash": row_hash, **normalized})
                    staged_context.append((row_index, raw_row))
                    if len(staged_occurrences) >= insert_chunk_size:
                        flush_staged_occurrences()

                    if row_index % progress_flush_rows == 0:
                        flush_staged_occurrences()
                        current_batch = session.get(ImportBatch, batch.id)
                        if current_batch is not None:
                            current_batch.resumo = (
                                f"Processando {file_index}/{total_files}: {file_path.name} "
                                f"({row_index}/{len(rows)} linhas lidas). {summary_text()}"
                            )
                        session.commit()

                flush_staged_occurrences()
                rebuild_clients_for_cpfs(session, file_affected_cpfs)
                summary.files_imported += 1
                current_batch = session.get(ImportBatch, batch.id)
                if current_batch is not None:
                    current_batch.resumo = (
                        f"Arquivo {file_index}/{total_files} concluido: {file_path.name}. {summary_text()}"
                    )
                session.commit()

        with serialized_write():
            final_batch = session.get(ImportBatch, batch.id)
            if final_batch is not None:
                if row_errors:
                    report_path = write_import_error_report(batch.id, row_errors)
                    summary.error_report_file = str(report_path)
                    summary.validation_errors.append(
                        f"Linhas com erro: {summary.row_errors}. Relatorio: {report_path.name}"
                    )
                final_batch.status = "CONCLUIDO"
                final_batch.resumo = summary_text()
                if summary.validation_errors:
                    final_batch.resumo += " | Avisos: " + "; ".join(summary.validation_errors[:5])
                session.commit()
        return summary
    except Exception as exc:
        session.rollback()
        with serialized_write():
            failed_batch = session.get(ImportBatch, batch.id)
            if failed_batch is not None:
                failed_batch.status = "ERRO"
                failed_batch.resumo = f"Falha na importacao: {exc}"
                session.commit()
        raise
    finally:
        shutil.rmtree(temp_dir, ignore_errors=True)


def import_files(session: Session, request: ImportRequest) -> ImportSummary:
    batch = create_import_batch(session, request)
    return process_import_batch(session, batch, request)


def build_occurrence_query(filters: FilterSet) -> Select[tuple[ClientOccurrence, SourceFile, ImportBatch]]:
    stmt = select(ClientOccurrence, SourceFile, ImportBatch).join(SourceFile, ClientOccurrence.source_file_id == SourceFile.id).join(ImportBatch, SourceFile.batch_id == ImportBatch.id)
    conditions = []
    segment = normalize_base_segment(filters.base_segment)
    entidade_filters = parse_multi_text_filter(filters.entidade)
    convenio_filters = parse_multi_text_filter(filters.convenio)
    servico_filters = parse_multi_text_filter(filters.servico_publico)
    situacao_filters = parse_multi_text_filter(filters.situacao)
    cargo_filters = parse_multi_text_filter(filters.cargo)
    tipo_vinculo_filters = parse_multi_text_filter(filters.tipo_vinculo)
    if segment in VALID_BASE_SEGMENTS:
        conditions.append(ClientOccurrence.base_segment == segment)
    if filters.base_source:
        conditions.append(ClientOccurrence.base_source.ilike(f"%{filters.base_source.upper()}%"))
    if entidade_filters:
        conditions.append(or_(*[ClientOccurrence.entidade.ilike(f"%{value.upper()}%") for value in entidade_filters]))
    if convenio_filters:
        conditions.append(or_(*[ClientOccurrence.convenio.ilike(f"%{value.upper()}%") for value in convenio_filters]))
    if servico_filters:
        conditions.append(or_(*[ClientOccurrence.servico.ilike(f"%{value.upper()}%") for value in servico_filters]))
    if situacao_filters:
        conditions.append(or_(*[ClientOccurrence.situacao.ilike(f"%{value.upper()}%") for value in situacao_filters]))
    if cargo_filters:
        cargo_conditions = []
        for value in cargo_filters:
            cargo_like = f"%{value.upper()}%"
            cargo_conditions.append(ClientOccurrence.cbo_titulo.ilike(cargo_like))
            cargo_conditions.append(ClientOccurrence.extras_json.ilike(cargo_like))
        conditions.append(
            or_(*cargo_conditions)
        )
    if tipo_vinculo_filters:
        conditions.append(or_(*[ClientOccurrence.extras_json.ilike(f"%{value.upper()}%") for value in tipo_vinculo_filters]))
    quick_digits = only_digits(filters.quick_value)
    ddb_start = normalize_date_filter(filters.ddb_start)
    ddb_end = normalize_date_filter(filters.ddb_end)
    has_ddb_range = bool(ddb_start or ddb_end)
    if quick_digits:
        if filters.quick_mode == "telefone":
            conditions.append(
                or_(
                    ClientOccurrence.telefone1.ilike(f"%{quick_digits}%"),
                    ClientOccurrence.telefone2.ilike(f"%{quick_digits}%"),
                    ClientOccurrence.telefone3.ilike(f"%{quick_digits}%"),
                )
            )
        else:
            normalized_cpf = normalize_cpf(quick_digits)
            identifier_conditions = [ClientOccurrence.cpf == normalized_cpf]
            if len(quick_digits) < 11 or len(quick_digits) == 14:
                doc_terms = build_document_search_terms(quick_digits)
                identifier_conditions.extend([ClientOccurrence.cpf_original == value for value in doc_terms])
                identifier_conditions.extend([ClientOccurrence.cpf_original.ilike(f"%{value}%") for value in doc_terms])
            if segment == "INSS":
                identifier_conditions.append(ClientOccurrence.nu_nb == quick_digits)
            else:
                identifier_conditions.append(ClientOccurrence.matricula.ilike(f"%{quick_digits}%"))
            conditions.append(
                or_(*identifier_conditions)
            )
    if filters.year and not has_ddb_range:
        conditions.append(ImportBatch.ano_referencia == filters.year)
    if filters.month and not has_ddb_range:
        conditions.append(ImportBatch.mes_referencia == filters.month)
    if filters.week:
        conditions.append(SourceFile.semana_label.ilike(f"%{filters.week}%"))
    if filters.uf:
        conditions.append(ClientOccurrence.uf == filters.uf.upper())
    if filters.city:
        conditions.append(ClientOccurrence.cidade.ilike(f"%{filters.city.upper()}%"))
    if filters.esp:
        species_filters = parse_species_filter(filters.esp)
        if species_filters:
            conditions.append(ClientOccurrence.esp.in_(species_filters))
    if filters.margin_min is not None:
        conditions.append(ClientOccurrence.vl_margem >= filters.margin_min)
    if filters.margin_max is not None:
        conditions.append(ClientOccurrence.vl_margem <= filters.margin_max)
    if filters.salary_min is not None:
        conditions.append(ClientOccurrence.salario >= filters.salary_min)
    if filters.salary_max is not None:
        conditions.append(ClientOccurrence.salario <= filters.salary_max)
    if filters.age_min is not None or filters.age_max is not None:
        today = date.today()
        if filters.age_min is not None:
            max_birth_date = add_years_safe(today, -filters.age_min)
            conditions.append(birth_date_expr() <= max_birth_date.isoformat())
        if filters.age_max is not None:
            min_birth_date = add_years_safe(today, -(filters.age_max + 1)) + timedelta(days=1)
            conditions.append(birth_date_expr() >= min_birth_date.isoformat())
    if ddb_start:
        conditions.append(ddb_date_expr() >= ddb_start)
    if ddb_end:
        conditions.append(ddb_date_expr() <= ddb_end)
    if filters.has_phone is True:
        conditions.append(or_(ClientOccurrence.telefone1 != "", ClientOccurrence.telefone2 != "", ClientOccurrence.telefone3 != ""))
    if filters.has_phone is False:
        conditions.append(and_(ClientOccurrence.telefone1 == "", ClientOccurrence.telefone2 == "", ClientOccurrence.telefone3 == ""))
    if filters.source_file:
        conditions.append(SourceFile.nome_arquivo.ilike(f"%{filters.source_file}%"))
    if filters.cpf:
        raw_cpf = only_digits(filters.cpf)
        normalized_cpf = normalize_cpf(filters.cpf)
        cpf_conditions = [ClientOccurrence.cpf == normalized_cpf]
        if raw_cpf and (len(raw_cpf) < 11 or len(raw_cpf) == 14):
            doc_terms = build_document_search_terms(raw_cpf)
            cpf_conditions.extend([ClientOccurrence.cpf_original == value for value in doc_terms])
            cpf_conditions.extend([ClientOccurrence.cpf_original.ilike(f"%{value}%") for value in doc_terms])
        conditions.append(or_(*cpf_conditions))
    if filters.phone:
        phone_digits = only_digits(filters.phone)
        if phone_digits:
            conditions.append(
                or_(
                    ClientOccurrence.telefone1.ilike(f"%{phone_digits}%"),
                    ClientOccurrence.telefone2.ilike(f"%{phone_digits}%"),
                    ClientOccurrence.telefone3.ilike(f"%{phone_digits}%"),
                )
            )
    if filters.ddd:
        ddd_filters = parse_ddd_filters(filters.ddd)
        if ddd_filters:
            conditions.append(
                or_(*[
                    ClientOccurrence.telefone1.like(f"{ddd}%")
                    for ddd in ddd_filters
                ])
            )
    if filters.name:
        conditions.append(ClientOccurrence.nome.ilike(f"%{filters.name.upper()}%"))
    if conditions:
        stmt = stmt.where(and_(*conditions))
    return stmt.order_by(ImportBatch.ano_referencia.desc(), ImportBatch.mes_referencia.desc(), ClientOccurrence.nome.asc())


def parse_decimal(value: str) -> Decimal | None:
    cleaned = norm_text(value).replace(",", ".")
    if not cleaned:
        return None
    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def _is_public_operational_row(row: dict[str, object]) -> bool:
    return any(
        [
            norm_text(row.get("matricula", "")),
            norm_text(row.get("servico", "")),
            norm_text(row.get("situacao", "")),
            row.get("vl_margem") is not None,
            row.get("margem_total") is not None,
            row.get("margem_liquida") is not None,
            row.get("margem_bruta") is not None,
            row.get("margem_utilizada") is not None,
        ]
    )


def _first_non_empty_value(rows: list[dict[str, object]], field: str) -> object:
    for row in rows:
        value = row.get(field)
        if isinstance(value, str):
            if norm_text(value):
                return value
        elif value is not None:
            return value
    return None


def _load_clients_by_cpfs(session: Session, cpfs: set[str]) -> dict[str, Client]:
    normalized_cpfs = {cpf for cpf in cpfs if cpf}
    if not normalized_cpfs:
        return {}
    loaded: dict[str, Client] = {}
    for cpf_chunk in chunked_values(normalized_cpfs):
        rows = session.execute(select(Client).where(Client.cpf.in_(cpf_chunk))).scalars().all()
        for row in rows:
            loaded[row.cpf] = row
    return loaded


def _reference_sort_key(value: object) -> tuple[int, int]:
    text = norm_text(value)
    if re.fullmatch(r"\d{4}-\d{2}", text):
        year_text, month_text = text.split("-", 1)
        return int(year_text), int(month_text)
    return 0, 0


def _row_currentness_key(row: dict[str, object]) -> tuple[int, tuple[int, int], date, int, int, int]:
    base_segment = normalize_base_segment(str(row.get("base_segment", "")))
    reference_key = _reference_sort_key(row.get("referencia", ""))
    ddb_value = parse_date_value(norm_text(row.get("ddb", ""))) or date.min
    has_benefit = 1 if norm_text(row.get("nu_nb", "")) else 0
    has_species = 1 if norm_text(row.get("esp", "")) else 0
    has_margin = 1 if row.get("vl_margem") is not None else 0
    return (
        1 if base_segment == "INSS" else 0,
        reference_key,
        ddb_value,
        has_benefit,
        has_species,
        has_margin,
    )


def _parse_extras_json(payload: str) -> dict[str, object]:
    if not payload:
        return {}
    try:
        data = json.loads(payload)
    except json.JSONDecodeError:
        return {}
    return data if isinstance(data, dict) else {}


def _first_extra_value(extras: dict[str, object], *keys: str) -> str:
    for key in keys:
        value = extras.get(key)
        if norm_text(value):
            return norm_text(value)
    return ""


def _consolidate_public_results(results: list[dict[str, object]]) -> list[dict[str, object]]:
    if not results:
        return results

    grouped: dict[str, list[dict[str, object]]] = {}
    for row in results:
        grouped.setdefault(str(row.get("cpf", "")), []).append(row)

    consolidated: list[dict[str, object]] = []
    for row in results:
        cpf = str(row.get("cpf", ""))
        group = grouped.get(cpf, [])
        has_operational = any(_is_public_operational_row(item) for item in group)
        is_operational = _is_public_operational_row(row)
        if has_operational and not is_operational:
            continue

        merged = dict(row)
        if group:
            field_names: set[str] = set()
            for item in group:
                field_names.update(item.keys())
            field_names.discard("selection_url")
            for field in field_names:
                current = merged.get(field)
                if (isinstance(current, str) and not norm_text(current)) or current is None:
                    fallback = _first_non_empty_value(group, field)
                    if fallback is not None:
                        merged[field] = fallback
        consolidated.append(merged)
    return consolidated


def _consolidate_results_by_cpf(results: list[dict[str, object]]) -> list[dict[str, object]]:
    if not results:
        return results

    grouped: dict[str, list[dict[str, object]]] = {}
    for row in results:
        cpf = str(row.get("cpf", ""))
        if cpf:
            grouped.setdefault(cpf, []).append(row)

    consolidated: list[dict[str, object]] = []
    for cpf, group in grouped.items():
        ranked = sorted(group, key=_row_currentness_key, reverse=True)
        latest_row = ranked[0]
        merged = dict(latest_row)
        # Regra de negocio: DDB e campos ligados ao beneficio devem sempre vir
        # da ocorrencia com DDB mais recente; registros antigos entram apenas no historico.
        newest_ddb = max((parse_date_value(norm_text(str(item.get("ddb", "") or ""))) or date.min) for item in ranked)
        benefit_candidates = [
            item
            for item in ranked
            if (parse_date_value(norm_text(str(item.get("ddb", "") or ""))) or date.min) == newest_ddb
        ] or ranked
        benefit_source = benefit_candidates[0]
        field_names: set[str] = set()
        for item in ranked:
            field_names.update(item.keys())
        for field in field_names:
            if field in BENEFIT_BOUND_FIELDS:
                benefit_value = benefit_source.get(field)
                if (isinstance(benefit_value, str) and norm_text(benefit_value)) or benefit_value is not None:
                    merged[field] = benefit_value
                continue
            current = merged.get(field)
            if (isinstance(current, str) and not norm_text(current)) or current is None:
                fallback = _first_non_empty_value(ranked, field)
                if fallback is not None:
                    merged[field] = fallback
        consolidated.append(merged)
    return consolidated


def _serialize_occurrence_row(
    occurrence: ClientOccurrence,
    source_file: SourceFile,
    batch: ImportBatch,
    client: Client | None = None,
) -> dict[str, object]:
    segment = normalize_base_segment(occurrence.base_segment)
    extras = _parse_extras_json(occurrence.extras_json)
    display_vl_margem = occurrence.vl_margem
    display_vl_rmc = occurrence.vl_rmc
    resolved_uf = occurrence.uf or (client.uf_atual if client is not None else "")
    resolved_city = occurrence.cidade or (client.cidade_atual if client is not None else "")
    if segment == "INSS":
        if display_vl_margem is None:
            display_vl_margem = calculate_margin_from_rmi(occurrence.vl_rmi, occurrence.esp)
        if display_vl_rmc is None:
            display_vl_rmc = calculate_rmc_from_rmi(occurrence.vl_rmi)
    address_parts = [
        norm_text(occurrence.logradouro),
        norm_text(occurrence.numero),
        norm_text(occurrence.complemento),
        norm_text(occurrence.bairro),
    ]
    endereco = ", ".join(part for part in address_parts if part)
    tipo_vinculo = _first_extra_value(extras, "REGIME_CONTRATACAO", "TIPO DE VINCULO", "TIPO_VINCULO", "VINCULO", "REGIME DE CONTRATACAO")
    cargo_funcao = _first_extra_value(extras, "CARGO", "FUNCAO", "FUNÃƒÆ’Ã¢â‚¬Â¡ÃƒÆ’Ã†â€™O", "CARGO/FUNCAO", "CARGO FUNCAO") or occurrence.cbo_titulo
    campaign_name = ""
    if segment in PUBLIC_BASE_SEGMENTS:
        campaign_name = build_public_campaign_name(
            segment,
            resolved_uf,
            cargo_funcao,
            occurrence.servico,
            occurrence.entidade,
            occurrence.convenio,
            occurrence.situacao,
            display_vl_margem,
        )
    return {
        "cpf": occurrence.cpf,
        "cpf_original": occurrence.cpf_original,
        "base_segment": occurrence.base_segment,
        "base_source": occurrence.base_source,
        "nome": occurrence.nome,
        "matricula": occurrence.matricula,
        "servico": occurrence.servico,
        "situacao": occurrence.situacao,
        "pis": occurrence.pis,
        "entidade": occurrence.entidade,
        "secretaria": occurrence.secretaria,
        "convenio": occurrence.convenio,
        "cbo_titulo": occurrence.cbo_titulo,
        "nu_nb": occurrence.nu_nb,
        "ddb": occurrence.ddb,
        "dt_nasc": occurrence.dt_nasc,
        "idade": compute_age_from_birth(occurrence.dt_nasc),
        "sexo": occurrence.sexo,
        "nome_mae": occurrence.nome_mae,
        "esp": occurrence.esp,
        "cidade": resolved_city,
        "uf": resolved_uf,
        "bairro": occurrence.bairro,
        "logradouro": occurrence.logradouro,
        "numero": occurrence.numero,
        "complemento": occurrence.complemento,
        "cep": occurrence.cep,
        "endereco": endereco,
        "email": occurrence.email,
        "email2": _first_extra_value(extras, "EMAIL2", "EMAIL_02"),
        "email3": _first_extra_value(extras, "EMAIL3", "EMAIL_03"),
        "salario": occurrence.salario,
        "vl_rmi": occurrence.vl_rmi,
        "vl_margem": display_vl_margem,
        "vl_rmc": display_vl_rmc,
        "margem_total": occurrence.margem_total,
        "margem_liquida": occurrence.margem_liquida,
        "margem_bruta": occurrence.margem_bruta,
        "margem_utilizada": occurrence.margem_utilizada,
        "telefone": occurrence.telefone1 or occurrence.telefone2 or occurrence.telefone3,
        "telefone2": occurrence.telefone2,
        "telefone3": occurrence.telefone3,
        "whatsapp_telefone1": _first_extra_value(extras, "FLGWHATSCEL1"),
        "whatsapp_telefone2": _first_extra_value(extras, "FLGWHATSCEL2"),
        "tipo_vinculo": tipo_vinculo,
        "cargo_funcao": cargo_funcao,
        "data_admissao": normalize_date_text(_first_extra_value(extras, "DATA ADMISSAO", "DT ADMISSAO", "ADMISSAO", "DATA_ADMISSAO")),
        "campaign_name": campaign_name,
        "numero_cliente": _first_extra_value(extras, "N DO CLIENTE", "NÂº DO CLIENTE", "NRO CLIENTE", "NUMERO CLIENTE"),
        "conta_contrato": _first_extra_value(extras, "CONTA CONTRATO", "CONTA_CONTRATO"),
        "numero_instalacao": _first_extra_value(extras, "NUMERO INSTALACAO", "NUMERO INSTALAÃ‡ÃƒO", "N INSTALACAO", "NÂº INSTALACAO"),
        "do_not_call": bool(client.do_not_call) if client is not None else False,
        "do_not_call_reason": (client.do_not_call_reason or "") if client is not None else "",
        "arquivo": source_file.nome_arquivo,
        "semana": source_file.semana_label,
        "referencia": latest_reference_label(batch.ano_referencia, batch.mes_referencia),
    }


def count_client_results(session: Session, filters: FilterSet) -> int:
    stmt = (
        build_occurrence_query(filters)
        .with_only_columns(func.count(func.distinct(ClientOccurrence.cpf)))
        .order_by(None)
    )
    return int(session.scalar(stmt) or 0)


def query_clients(session: Session, filters: FilterSet, limit: int = 200, offset: int = 0) -> list[dict[str, object]]:
    rows = session.execute(
        build_occurrence_query(filters)
        .offset(max(offset, 0))
        .limit(limit)
    ).all()
    clients_by_cpf = _load_clients_by_cpfs(session, {occurrence.cpf for occurrence, _source_file, _batch in rows})
    results = [
        _serialize_occurrence_row(occurrence, source_file, batch, clients_by_cpf.get(occurrence.cpf))
        for occurrence, source_file, batch in rows
    ]
    if normalize_base_segment(filters.base_segment) in PUBLIC_BASE_SEGMENTS:
        return _consolidate_public_results(results)
    return results


def _export_rows(
    session: Session,
    rows: list[tuple[ClientOccurrence, SourceFile, ImportBatch]],
    include_audit: bool,
    filters: FilterSet,
) -> list[dict[str, object]]:
    output: list[dict[str, object]] = []
    ddd_filters = parse_ddd_filters(filters.ddd) if filters.ddd else []
    clients_by_cpf = _load_clients_by_cpfs(session, {occurrence.cpf for occurrence, _source_file, _batch in rows})
    for occurrence, source_file, batch in rows:
        client = clients_by_cpf.get(occurrence.cpf)
        if client is not None and bool(client.do_not_call) and not filters.include_do_not_call:
            continue
        serialized = _serialize_occurrence_row(occurrence, source_file, batch, client)
        export_phone = occurrence.telefone1 or occurrence.telefone2 or occurrence.telefone3
        if ddd_filters and export_phone and not any(export_phone.startswith(ddd) for ddd in ddd_filters):
            continue
        segment = normalize_base_segment(occurrence.base_segment)
        export_vl_margem = serialized.get("vl_margem")
        export_vl_rmc = serialized.get("vl_rmc")
        if segment in PUBLIC_BASE_SEGMENTS and not matches_public_campaign_profile(
            filters.campaign_profile,
            str(serialized.get("campaign_name", "")),
            str(serialized.get("cargo_funcao", "")),
            str(serialized.get("servico", "")),
            str(serialized.get("entidade", "")),
            str(serialized.get("convenio", "")),
            str(serialized.get("situacao", "")),
            export_vl_margem if isinstance(export_vl_margem, Decimal) else None,
        ):
            continue
        if segment == "INSS":
            row = {
                "TELEFONE": export_phone,
                "NOME": occurrence.nome,
                "CPF": occurrence.cpf,
                "NU-NB": occurrence.nu_nb,
                "VL-RMI": format_money(occurrence.vl_rmi),
                "VL MARGEM": format_money(export_vl_margem),
                "VL RMC": format_money(export_vl_rmc),
                "DDB": occurrence.ddb,
                "UF": serialized.get("uf", ""),
                "CIDADE": serialized.get("cidade", ""),
                "ESP": occurrence.esp,
            }
        else:
            row = {
                "TELEFONE": export_phone,
                "NOME": occurrence.nome,
                "CPF": occurrence.cpf,
                "MATRICULA": occurrence.matricula,
                "SERVICO": occurrence.servico,
                "SITUACAO": occurrence.situacao,
                "ENTIDADE": occurrence.entidade,
                "CBO TITULO": occurrence.cbo_titulo,
                "MARGEM": format_money(export_vl_margem if isinstance(export_vl_margem, Decimal) else occurrence.vl_margem),
                "MARGEM TOTAL": format_money(occurrence.margem_total),
                "CARGO": serialized.get("cargo_funcao", ""),
                "TIPO_VINCULO": serialized.get("tipo_vinculo", ""),
                "CONVENIO": occurrence.convenio,
                "FONTE": occurrence.base_source,
                "CAMPANHA": serialized.get("campaign_name", ""),
                "UF": serialized.get("uf", ""),
                "CIDADE": serialized.get("cidade", ""),
            }
        if include_audit:
            audit_fields = {
                "ARQUIVO_ORIGEM": source_file.nome_arquivo,
                "SEMANA": source_file.semana_label,
                "ANO": batch.ano_referencia,
                "MES": batch.mes_referencia,
                "TELEFONE2": occurrence.telefone2,
                "TELEFONE3": occurrence.telefone3,
                "EMAIL": occurrence.email,
            }
            if segment in PUBLIC_BASE_SEGMENTS:
                audit_fields["FONTE"] = occurrence.base_source
            row.update(audit_fields)
        output.append(row)
    return output


def export_clients(session: Session, filters: FilterSet, include_audit: bool, file_format: str = "csv") -> Path:
    ensure_directories()
    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    suffix = "auditoria" if include_audit else "operacional"
    rows = session.execute(build_occurrence_query(filters)).all()
    selected_segment = normalize_base_segment(filters.base_segment)
    if selected_segment == "INSS":
        headers = AUDIT_HEADERS if include_audit else EXPORT_HEADERS
    else:
        headers = PUBLIC_AUDIT_HEADERS if include_audit else PUBLIC_EXPORT_HEADERS
    export_rows = _export_rows(session, rows, include_audit, filters)
    if file_format == "xlsx":
        export_path = settings.exports_dir / f"base_{suffix}_{timestamp}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Base"
        sheet.append(headers)
        for row in export_rows:
            sheet.append([row.get(header, "") for header in headers])
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(column[0].value or "")) + 2, 14), 28)
        workbook.save(export_path)
    else:
        export_path = settings.exports_dir / f"base_{suffix}_{timestamp}.csv"
        with export_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers, delimiter=";")
            writer.writeheader()
            writer.writerows(export_rows)
    with serialized_write():
        session.add(ExportHistory(filtros_json=json.dumps({**filters.__dict__, "file_format": file_format}, ensure_ascii=False, default=str), arquivo_saida=str(export_path), total_registros=len(export_rows)))
        session.commit()
    return export_path


def export_cpfs_for_enrichment(session: Session, filters: FilterSet, file_format: str = "xlsx") -> Path:
    ensure_directories()
    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    raw_rows = session.execute(build_occurrence_query(filters)).all()
    clients_by_cpf = _load_clients_by_cpfs(session, {occurrence.cpf for occurrence, _source_file, _batch in raw_rows})
    deduped: dict[str, dict[str, str]] = {}
    for occurrence, _source_file, _batch in raw_rows:
        client = clients_by_cpf.get(occurrence.cpf)
        if client is not None and bool(client.do_not_call) and not filters.include_do_not_call:
            continue
        if occurrence.cpf not in deduped:
            deduped[occurrence.cpf] = {"CPF": occurrence.cpf, "NOME": occurrence.nome}

    export_rows = list(deduped.values())
    if file_format == "csv":
        export_path = settings.exports_dir / f"cpfs_enriquecimento_{timestamp}.csv"
        with export_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=ENRICHMENT_CPF_HEADERS, delimiter=";")
            writer.writeheader()
            writer.writerows(export_rows)
    else:
        export_path = settings.exports_dir / f"cpfs_enriquecimento_{timestamp}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "CPFs"
        sheet.append(ENRICHMENT_CPF_HEADERS)
        for row in export_rows:
            sheet.append([row.get(header, "") for header in ENRICHMENT_CPF_HEADERS])
        workbook.save(export_path)

    with serialized_write():
        session.add(
            ExportHistory(
                filtros_json=json.dumps({**filters.__dict__, "purpose": "phone_enrichment", "file_format": file_format}, ensure_ascii=False, default=str),
                arquivo_saida=str(export_path),
                total_registros=len(export_rows),
            )
        )
        session.commit()
    return export_path


def export_wrong_number_clients_for_novavida(session: Session, filters: FilterSet, file_format: str = "xlsx") -> Path:
    ensure_directories()
    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    raw_rows = session.execute(build_occurrence_query(filters)).all()
    clients_by_cpf = _load_clients_by_cpfs(session, {occurrence.cpf for occurrence, _source_file, _batch in raw_rows})
    deduped: dict[str, dict[str, str]] = {}
    for occurrence, _source_file, _batch in raw_rows:
        client = clients_by_cpf.get(occurrence.cpf)
        if client is None or not bool(client.do_not_call):
            continue
        if (client.do_not_call_reason or "") != "nao_e_o_cliente":
            continue
        if occurrence.cpf not in deduped:
            deduped[occurrence.cpf] = {
                "CPF": occurrence.cpf,
                "NOME": occurrence.nome,
                "TELEFONE": occurrence.telefone1 or occurrence.telefone2 or occurrence.telefone3 or "",
                "MOTIVO": "Nao e o cliente",
            }

    export_rows = list(deduped.values())
    if file_format == "csv":
        export_path = settings.exports_dir / f"novavida_numero_errado_{timestamp}.csv"
        with export_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=WRONG_NUMBER_EXPORT_HEADERS, delimiter=";")
            writer.writeheader()
            writer.writerows(export_rows)
    else:
        export_path = settings.exports_dir / f"novavida_numero_errado_{timestamp}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Numero Errado"
        sheet.append(WRONG_NUMBER_EXPORT_HEADERS)
        for row in export_rows:
            sheet.append([row.get(header, "") for header in WRONG_NUMBER_EXPORT_HEADERS])
        workbook.save(export_path)

    with serialized_write():
        session.add(
            ExportHistory(
                filtros_json=json.dumps({**filters.__dict__, "purpose": "novavida_wrong_number", "file_format": file_format}, ensure_ascii=False, default=str),
                arquivo_saida=str(export_path),
                total_registros=len(export_rows),
            )
        )
        session.commit()
    return export_path


def export_updated_clients(session: Session, filters: FilterSet, file_format: str = "xlsx") -> Path:
    ensure_directories()
    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    rows = query_clients(session, filters, limit=100000)
    consolidated_rows = _consolidate_results_by_cpf(rows)

    export_rows: list[dict[str, object]] = []
    for row in consolidated_rows:
        if bool(row.get("do_not_call")) and not filters.include_do_not_call:
            continue
        export_salary = resolve_updated_base_salary(row)
        export_margin_total = resolve_updated_base_margin_total(row)
        export_rows.append(
            {
                "TELEFONE": row.get("telefone", ""),
                "TELEFONE2": row.get("telefone2", ""),
                "TELEFONE3": row.get("telefone3", ""),
                "WHATSAPP": row.get("whatsapp_telefone1", ""),
                "WHATSAPP2": row.get("whatsapp_telefone2", ""),
                "NOME": row.get("nome", ""),
                "CPF": row.get("cpf", ""),
                "SEGMENTO": row.get("base_segment", ""),
                "FONTE": row.get("base_source", ""),
                "EMAIL": row.get("email", ""),
                "EMAIL2": row.get("email2", ""),
                "EMAIL3": row.get("email3", ""),
                "CIDADE": row.get("cidade", ""),
                "UF": row.get("uf", ""),
                "ENDERECO": row.get("endereco", ""),
                "BAIRRO": row.get("bairro", ""),
                "LOGRADOURO": row.get("logradouro", ""),
                "NUMERO": row.get("numero", ""),
                "COMPLEMENTO": row.get("complemento", ""),
                "CEP": row.get("cep", ""),
                "DT_NASC": row.get("dt_nasc", ""),
                "IDADE": row.get("idade", ""),
                "NOME_MAE": row.get("nome_mae", ""),
                "SEXO": row.get("sexo", ""),
                "MATRICULA": row.get("matricula", ""),
                "SERVICO": row.get("servico", ""),
                "SITUACAO": row.get("situacao", ""),
                "ENTIDADE": row.get("entidade", ""),
                "TIPO_VINCULO": row.get("tipo_vinculo", ""),
                "CARGO": row.get("cargo_funcao", ""),
                "SALARIO": format_money(export_salary) if export_salary is not None else "",
                "MARGEM_DISPONIVEL": format_money(row.get("vl_margem")) if row.get("vl_margem") is not None else "",
                "MARGEM_RMC": format_money(row.get("vl_rmc")) if row.get("vl_rmc") is not None else "",
                "MARGEM_TOTAL": format_money(export_margin_total) if export_margin_total is not None else "",
                "DDB": row.get("ddb", ""),
                "BENEFICIO": row.get("nu_nb", ""),
                "ESPECIE": row.get("esp", ""),
            }
        )

    if file_format == "csv":
        export_path = settings.exports_dir / f"base_atualizada_{timestamp}.csv"
        with export_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=UPDATED_BASE_HEADERS, delimiter=";")
            writer.writeheader()
            writer.writerows(export_rows)
    else:
        export_path = settings.exports_dir / f"base_atualizada_{timestamp}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Base Atualizada"
        sheet.append(UPDATED_BASE_HEADERS)
        for row in export_rows:
            sheet.append([row.get(header, "") for header in UPDATED_BASE_HEADERS])
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(column[0].value or "")) + 2, 14), 28)
        workbook.save(export_path)

    with serialized_write():
        session.add(
            ExportHistory(
                filtros_json=json.dumps({**filters.__dict__, "purpose": "updated_base", "file_format": file_format}, ensure_ascii=False, default=str),
                arquivo_saida=str(export_path),
                total_registros=len(export_rows),
            )
        )
        session.commit()
    return export_path


def export_updated_clients_from_latest_enrichment(
    session: Session,
    base_segment: str,
    file_format: str = "xlsx",
) -> Path:
    latest_entry = session.execute(
        select(EnrichmentImportItem).order_by(EnrichmentImportItem.criado_em.desc(), EnrichmentImportItem.id.desc()).limit(1)
    ).scalar_one_or_none()
    if latest_entry is None:
        raise ValueError("Nenhum retorno enriquecido foi importado ainda.")

    latest_file = latest_entry.arquivo_origem
    cpfs = {
        cpf
        for (cpf,) in session.execute(
            select(EnrichmentImportItem.cpf).where(EnrichmentImportItem.arquivo_origem == latest_file).distinct()
        ).all()
        if cpf
    }
    if not cpfs:
        raise ValueError("O ultimo retorno nao possui CPFs validos para exportacao.")

    filters = FilterSet(base_segment=normalize_base_segment(base_segment))
    targeted_rows: list[dict[str, object]] = []
    for cpf_chunk in chunked_values(list(cpfs), chunk_size=5000):
        stmt = build_occurrence_query(filters).where(ClientOccurrence.cpf.in_(cpf_chunk))
        for occurrence, source_file, batch in session.execute(stmt).all():
            targeted_rows.append(_serialize_occurrence_row(occurrence, source_file, batch))
    filtered_rows = targeted_rows
    consolidated_rows = _consolidate_results_by_cpf(filtered_rows)

    ensure_directories()
    timestamp = datetime.now(UTC).strftime("%Y%m%d_%H%M%S")
    export_rows: list[dict[str, object]] = []
    for row in consolidated_rows:
        export_salary = resolve_updated_base_salary(row)
        export_margin_total = resolve_updated_base_margin_total(row)
        export_rows.append(
            {
                "TELEFONE": row.get("telefone", ""),
                "TELEFONE2": row.get("telefone2", ""),
                "TELEFONE3": row.get("telefone3", ""),
                "WHATSAPP": row.get("whatsapp_telefone1", ""),
                "WHATSAPP2": row.get("whatsapp_telefone2", ""),
                "NOME": row.get("nome", ""),
                "CPF": row.get("cpf", ""),
                "SEGMENTO": row.get("base_segment", ""),
                "FONTE": row.get("base_source", ""),
                "EMAIL": row.get("email", ""),
                "EMAIL2": row.get("email2", ""),
                "EMAIL3": row.get("email3", ""),
                "CIDADE": row.get("cidade", ""),
                "UF": row.get("uf", ""),
                "ENDERECO": row.get("endereco", ""),
                "BAIRRO": row.get("bairro", ""),
                "LOGRADOURO": row.get("logradouro", ""),
                "NUMERO": row.get("numero", ""),
                "COMPLEMENTO": row.get("complemento", ""),
                "CEP": row.get("cep", ""),
                "DT_NASC": row.get("dt_nasc", ""),
                "IDADE": row.get("idade", ""),
                "NOME_MAE": row.get("nome_mae", ""),
                "SEXO": row.get("sexo", ""),
                "MATRICULA": row.get("matricula", ""),
                "SERVICO": row.get("servico", ""),
                "SITUACAO": row.get("situacao", ""),
                "ENTIDADE": row.get("entidade", ""),
                "TIPO_VINCULO": row.get("tipo_vinculo", ""),
                "CARGO": row.get("cargo_funcao", ""),
                "SALARIO": format_money(export_salary) if export_salary is not None else "",
                "MARGEM_DISPONIVEL": format_money(row.get("vl_margem")) if row.get("vl_margem") is not None else "",
                "MARGEM_RMC": format_money(row.get("vl_rmc")) if row.get("vl_rmc") is not None else "",
                "MARGEM_TOTAL": format_money(export_margin_total) if export_margin_total is not None else "",
                "DDB": row.get("ddb", ""),
                "BENEFICIO": row.get("nu_nb", ""),
                "ESPECIE": row.get("esp", ""),
            }
        )

    if file_format == "csv":
        export_path = settings.exports_dir / f"base_atualizada_ultimo_retorno_{timestamp}.csv"
        with export_path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=UPDATED_BASE_HEADERS, delimiter=";")
            writer.writeheader()
            writer.writerows(export_rows)
    else:
        export_path = settings.exports_dir / f"base_atualizada_ultimo_retorno_{timestamp}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Base Atualizada"
        sheet.append(UPDATED_BASE_HEADERS)
        for row in export_rows:
            sheet.append([row.get(header, "") for header in UPDATED_BASE_HEADERS])
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(max(len(str(column[0].value or "")) + 2, 14), 28)
        workbook.save(export_path)

    with serialized_write():
        session.add(
            ExportHistory(
                filtros_json=json.dumps(
                    {
                        "base_segment": normalize_base_segment(base_segment),
                        "purpose": "updated_base_latest_enrichment",
                        "latest_file": latest_file,
                        "file_format": file_format,
                    },
                    ensure_ascii=False,
                    default=str,
                ),
                arquivo_saida=str(export_path),
                total_registros=len(export_rows),
            )
        )
        session.commit()
    return export_path


def normalize_public_campaign_update_row(row: dict[str, str]) -> dict[str, str | Decimal | int | None]:
    cpf = normalize_cpf(get_normalized_row_value(row, "CPF"))
    matricula = only_digits(get_normalized_row_value(row, "MATRICULA"))
    nome = normalize_upper(get_normalized_row_value(row, "NOME", "NAME"))
    tipo_servico_servidor = normalize_upper(get_normalized_row_value(row, "TIPO SERVICO (SERVIDOR)", "TIPO SERVIÃ‡O (SERVIDOR)"))
    servico_servidor = normalize_upper(get_normalized_row_value(row, "SERVICO (SERVIDOR)", "SERVIÃ‡O (SERVIDOR)", "SERVICO"))
    margem_disponivel = parse_money(
        get_normalized_row_value(row, "MARGEM DISPONIVEL (R$)", "MARGEM DISPONIVEL", "MARGEM DISPONÃVEL (R$)", "VL MARGEM")
    )
    margem_total = parse_money(get_normalized_row_value(row, "MARGEM TOTAL (R$)", "MARGEM TOTAL", "VL-RMI"))
    entidade = normalize_upper(get_normalized_row_value(row, "ENTIDADE"))
    cbo_titulo = normalize_upper(get_normalized_row_value(row, "CBO_TITULO", "CBO TITULO", "CARGO", "CARGO_FUNCAO", "CARGO/FUNCAO"))
    consignataria = normalize_upper(get_normalized_row_value(row, "CONSIGNATARIA", "CONSIGNATÃRIA"))
    situacao = normalize_upper(get_normalized_row_value(row, "SITUACAO", "SITUAÃ‡ÃƒO"))
    ade = only_digits(get_normalized_row_value(row, "ADE"))
    servico_consignataria = normalize_upper(get_normalized_row_value(row, "SERVICO (CONSIGNATARIA)", "SERVIÃ‡O (CONSIGNATÃRIA)"))
    prestacoes = parse_int(get_normalized_row_value(row, "PRESTACOES", "PRESTAÃ‡Ã•ES"))
    pagas = parse_int(get_normalized_row_value(row, "PAGAS"))
    valor = parse_money(get_normalized_row_value(row, "VALOR"))
    deferimento = normalize_date_text(get_normalized_row_value(row, "DEFERIMENTO"))
    quitacao = normalize_date_text(get_normalized_row_value(row, "QUITACAO", "QUITAÃ‡ÃƒO"))
    ultimo_desconto = normalize_date_text(get_normalized_row_value(row, "ULTIMO DESCONTO", "ÃšLTIMO DESCONTO"))
    ultima_parcela = normalize_date_text(get_normalized_row_value(row, "ULTIMA PARCELA", "ÃšLTIMA PARCELA"))
    if not ade:
        row_fingerprint = "|".join(
            [
                cpf,
                matricula,
                servico_consignataria,
                consignataria,
                format_money(valor) if isinstance(valor, Decimal) else "",
                deferimento,
                ultima_parcela,
            ]
        )
        ade = hashlib.sha1(row_fingerprint.encode("utf-8")).hexdigest()[:20]
    return {
        "cpf": cpf,
        "matricula": matricula,
        "nome": nome,
        "tipo_servico_servidor": tipo_servico_servidor,
        "servico_servidor": servico_servidor,
        "entidade": entidade,
        "cbo_titulo": cbo_titulo,
        "margem_disponivel": margem_disponivel,
        "margem_total": margem_total,
        "consignataria": consignataria,
        "situacao": situacao,
        "ade": ade,
        "servico_consignataria": servico_consignataria,
        "prestacoes": prestacoes,
        "pagas": pagas,
        "valor": valor,
        "deferimento": deferimento,
        "quitacao": quitacao,
        "ultimo_desconto": ultimo_desconto,
        "ultima_parcela": ultima_parcela,
    }


def is_active_public_contract_status(status: str) -> bool:
    normalized = normalize_token(status)
    return normalized.startswith("DEFERIDA") or normalized.startswith("ATIVA") or normalized.startswith("AVERBADA")


def build_public_campaign_aggregates(rows: list[dict[str, str | Decimal | int | None]]) -> dict[str, list[PublicCampaignAggregate]]:
    grouped: dict[tuple[str, str], PublicCampaignAggregate] = {}
    for row in rows:
        cpf = str(row.get("cpf") or "")
        if not cpf:
            continue
        matricula = str(row.get("matricula") or "")
        key = (cpf, matricula)
        item = grouped.get(key)
        if item is None:
            item = PublicCampaignAggregate(cpf=cpf, matricula=matricula)
            grouped[key] = item
        item.qtd_contratos += 1
        if row.get("nome"):
            item.nome = str(row["nome"])
        if row.get("tipo_servico_servidor"):
            item.tipo_servico_servidor = str(row["tipo_servico_servidor"])
        if row.get("servico_servidor"):
            item.servico_servidor = str(row["servico_servidor"])
        if row.get("entidade"):
            item.entidade = str(row["entidade"])
        if row.get("cbo_titulo"):
            item.cbo_titulo = str(row["cbo_titulo"])
        margem_disponivel = row.get("margem_disponivel")
        if isinstance(margem_disponivel, Decimal) and (item.margem_disponivel is None or margem_disponivel > item.margem_disponivel):
            item.margem_disponivel = margem_disponivel
        margem_total = row.get("margem_total")
        if isinstance(margem_total, Decimal) and (item.margem_total is None or margem_total > item.margem_total):
            item.margem_total = margem_total
        situacao = str(row.get("situacao") or "")
        if situacao and not item.situacao_principal:
            item.situacao_principal = situacao
        if is_active_public_contract_status(situacao):
            item.qtd_ativos += 1
            item.situacao_principal = situacao
            if row.get("consignataria"):
                item.consignataria_principal = str(row["consignataria"])
        consignataria = str(row.get("consignataria") or "")
        if consignataria:
            item.bancos.add(consignataria)
            if not item.consignataria_principal:
                item.consignataria_principal = consignataria
    by_cpf: dict[str, list[PublicCampaignAggregate]] = {}
    for item in grouped.values():
        by_cpf.setdefault(item.cpf, []).append(item)
    for cpf, items in by_cpf.items():
        by_cpf[cpf] = sorted(
            items,
            key=lambda current: (
                1 if current.matricula else 0,
                1 if current.margem_disponivel is not None else 0,
                current.margem_disponivel or Decimal("0"),
                current.qtd_ativos,
                current.qtd_contratos,
            ),
            reverse=True,
        )
    return by_cpf


def select_public_aggregate_for_occurrence(occurrence: ClientOccurrence, options: list[PublicCampaignAggregate]) -> PublicCampaignAggregate | None:
    if not options:
        return None
    if occurrence.matricula:
        for option in options:
            if option.matricula and option.matricula == occurrence.matricula:
                return option
    if len(options) == 1:
        return options[0]
    return options[0]


def build_public_update_payload(
    aggregate: PublicCampaignAggregate,
    source_name: str,
    file_name: str,
) -> dict[str, str | Decimal | None]:
    extras = {
        "PUBLIC_UPDATE_SOURCE": source_name,
        "PUBLIC_UPDATE_FILE": file_name,
        "PUBLIC_UPDATE_QTD_CONTRATOS": str(aggregate.qtd_contratos),
        "PUBLIC_UPDATE_QTD_ATIVOS": str(aggregate.qtd_ativos),
        "PUBLIC_UPDATE_BANCOS": " | ".join(sorted(aggregate.bancos)),
        "PUBLIC_UPDATE_CONSIGNATARIA_PRINCIPAL": aggregate.consignataria_principal,
        "PUBLIC_UPDATE_TIPO_SERVICO": aggregate.tipo_servico_servidor,
    }
    return {
        "nome": aggregate.nome,
        "matricula": aggregate.matricula,
        "servico": aggregate.servico_servidor,
        "situacao": aggregate.situacao_principal,
        "entidade": aggregate.entidade,
        "cbo_titulo": aggregate.cbo_titulo,
        "vl_margem": aggregate.margem_disponivel,
        "margem_total": aggregate.margem_total,
        "extras_json": json.dumps({key: value for key, value in extras.items() if value}, ensure_ascii=False, sort_keys=True),
    }


def import_public_campaign_updates(
    session: Session,
    files: list[Path],
    base_segment: str,
    source_name: str,
    imported_by: str,
) -> PublicCampaignUpdateSummary:
    summary = PublicCampaignUpdateSummary()
    ensure_directories()
    normalized_segment = normalize_base_segment(base_segment, default="GOVERNO")
    normalized_source = normalize_upper(source_name)
    affected_cpfs: set[str] = set()

    with tempfile.TemporaryDirectory() as temp_dir_name:
        temp_dir = Path(temp_dir_name)
        with serialized_write():
            for file_path in files:
                if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                    continue
                summary.files_processed += 1
                rows = iter_rows_from_file(file_path, temp_dir)
                normalized_rows: list[dict[str, str | Decimal | int | None]] = []
                seen_contract_keys: set[tuple[str, str, str, str, str]] = set()
                for row in rows:
                    summary.rows_processed += 1
                    normalized = normalize_public_campaign_update_row(row)
                    cpf = str(normalized.get("cpf") or "")
                    if not cpf:
                        summary.skipped_without_cpf += 1
                        continue
                    ade = str(normalized.get("ade") or "")
                    contract_key = (
                        cpf,
                        str(normalized.get("matricula") or ""),
                        ade,
                        normalized_source,
                        file_path.name,
                    )
                    if contract_key in seen_contract_keys:
                        continue
                    seen_contract_keys.add(contract_key)
                    normalized_rows.append(normalized)
                    summary.rows_imported += 1
                    exists = session.scalar(
                        select(func.count())
                        .select_from(PublicCampaignUpdate)
                        .where(
                            PublicCampaignUpdate.cpf == cpf,
                            PublicCampaignUpdate.matricula == str(normalized.get("matricula") or ""),
                            PublicCampaignUpdate.ade == ade,
                            PublicCampaignUpdate.fonte == normalized_source,
                            PublicCampaignUpdate.arquivo_origem == file_path.name,
                        )
                    )
                    if not exists:
                        session.add(
                            PublicCampaignUpdate(
                                base_segment=normalized_segment,
                                cpf=cpf,
                                matricula=str(normalized.get("matricula") or ""),
                                tipo_servico_servidor=str(normalized.get("tipo_servico_servidor") or ""),
                                servico_servidor=str(normalized.get("servico_servidor") or ""),
                                margem_disponivel=normalized.get("margem_disponivel"),  # type: ignore[arg-type]
                                margem_total=normalized.get("margem_total"),  # type: ignore[arg-type]
                                consignataria=str(normalized.get("consignataria") or ""),
                                situacao=str(normalized.get("situacao") or ""),
                                ade=ade,
                                servico_consignataria=str(normalized.get("servico_consignataria") or ""),
                                prestacoes=normalized.get("prestacoes"),  # type: ignore[arg-type]
                                pagas=normalized.get("pagas"),  # type: ignore[arg-type]
                                valor=normalized.get("valor"),  # type: ignore[arg-type]
                                deferimento=str(normalized.get("deferimento") or ""),
                                quitacao=str(normalized.get("quitacao") or ""),
                                ultimo_desconto=str(normalized.get("ultimo_desconto") or ""),
                                ultima_parcela=str(normalized.get("ultima_parcela") or ""),
                                fonte=normalized_source,
                                arquivo_origem=file_path.name,
                                importado_por=imported_by or "operador",
                            )
                        )
                        summary.contracts_saved += 1

                aggregates_by_cpf = build_public_campaign_aggregates(normalized_rows)
                for cpf, options in aggregates_by_cpf.items():
                    occurrences = session.execute(
                        select(ClientOccurrence).where(
                            ClientOccurrence.cpf == cpf,
                            ClientOccurrence.base_segment == normalized_segment,
                        )
                    ).scalars().all()
                    if not occurrences:
                        summary.skipped_not_found += 1
                        continue
                    summary.clients_matched += 1
                    affected_cpfs.add(cpf)
                    updated_any = False
                    existing_matriculas = {norm_text(item.matricula) for item in occurrences if norm_text(item.matricula)}
                    template_occurrence = occurrences[0]
                    for aggregate in options:
                        aggregate_matricula = norm_text(aggregate.matricula)
                        if not aggregate_matricula or aggregate_matricula in existing_matriculas:
                            continue
                        clone = materialize_public_update_occurrence(template_occurrence, aggregate, normalized_source, file_path.name)
                        session.add(clone)
                        session.flush()
                        occurrences.append(clone)
                        existing_matriculas.add(aggregate_matricula)
                        updated_any = True
                    for occurrence in occurrences:
                        aggregate = select_public_aggregate_for_occurrence(occurrence, options)
                        if aggregate is None:
                            continue
                        payload = build_public_update_payload(aggregate, normalized_source, file_path.name)
                        _phone_changed, data_changed = apply_enrichment_to_occurrence(occurrence, payload, [])
                        updated_any = updated_any or data_changed
                    if updated_any:
                        summary.clients_updated += 1

            if affected_cpfs:
                rebuild_clients_for_cpfs(session, affected_cpfs)
            session.commit()
    return summary


def import_phone_enrichments(
    session: Session,
    files: list[Path],
    source_name: str,
    imported_by: str,
) -> PhoneEnrichmentSummary:
    summary = PhoneEnrichmentSummary()
    ensure_directories()
    source_label = normalize_upper(source_name) or "ENRIQUECIMENTO"
    affected_cpfs: set[str] = set()

    with tempfile.TemporaryDirectory() as temp_dir_name:
        temp_dir = Path(temp_dir_name)
        with serialized_write():
            for file_path in files:
                if file_path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                    continue
                summary.files_processed += 1
                rows = iter_rows_from_file(file_path, temp_dir)
                for row in rows:
                    summary.rows_processed += 1
                    normalized = normalize_row(row)
                    cpf = str(normalized["cpf"])
                    if not cpf:
                        summary.skipped_without_cpf += 1
                        continue
                    phones = detect_phones(row)
                    has_non_phone_data = any(
                        normalized.get(field)
                        for field in (
                            "nome",
                            "nome_higienizado",
                            "cpf_original",
                            "matricula",
                            "servico",
                            "situacao",
                            "nu_nb",
                            "pis",
                            "entidade",
                            "secretaria",
                            "convenio",
                            "cbo_titulo",
                            "esp",
                            "ddb",
                            "uf",
                            "cidade",
                            "bairro",
                            "logradouro",
                            "numero",
                            "complemento",
                            "cep",
                            "dt_nasc",
                            "nome_mae",
                            "sexo",
                            "email",
                            "extras_json",
                        )
                    ) or any(
                        normalized.get(field) is not None
                        for field in (
                            "salario",
                            "vl_rmi",
                            "vl_margem",
                            "vl_rmc",
                            "margem_total",
                            "margem_liquida",
                            "margem_bruta",
                            "margem_utilizada",
                        )
                    )
                    if not phones and not has_non_phone_data:
                        summary.skipped_without_phone += 1
                        continue

                    occurrences = session.execute(select(ClientOccurrence).where(ClientOccurrence.cpf == cpf)).scalars().all()
                    client = session.get(Client, cpf)
                    if client is None and occurrences:
                        seed = occurrences[0]
                        client = Client(
                            cpf=cpf,
                            nome_atual=seed.nome,
                            uf_atual=seed.uf,
                            cidade_atual=seed.cidade,
                            melhor_telefone=seed.telefone1 or seed.telefone2 or seed.telefone3,
                            tem_telefone=bool(seed.telefone1 or seed.telefone2 or seed.telefone3),
                            maior_margem=seed.vl_margem,
                            ultima_referencia="",
                            qtd_ocorrencias=len(occurrences),
                        )
                        session.add(client)
                        session.flush()
                    if client is None:
                        summary.skipped_not_found += 1
                        continue

                    summary.clients_matched += 1
                    affected_cpfs.add(cpf)
                    import_item_exists = session.scalar(
                        select(func.count())
                        .select_from(EnrichmentImportItem)
                        .where(
                            EnrichmentImportItem.cpf == cpf,
                            EnrichmentImportItem.arquivo_origem == file_path.name,
                            EnrichmentImportItem.fonte == source_label,
                        )
                    )
                    if not import_item_exists:
                        session.add(
                            EnrichmentImportItem(
                                cpf=cpf,
                                fonte=source_label,
                                arquivo_origem=file_path.name,
                                importado_por=imported_by or "operador",
                            )
                        )
                    previous_primary = client.melhor_telefone
                    existing_phones = [client.melhor_telefone]
                    if occurrences:
                        sample = occurrences[0]
                        existing_phones.extend([sample.telefone1, sample.telefone2, sample.telefone3])
                    known_phone_set = {only_digits(phone) for phone in existing_phones if only_digits(phone)}
                    incoming_phone_set = {only_digits(phone) for phone in phones if only_digits(phone)}
                    has_new_incoming_phone = bool(incoming_phone_set - known_phone_set)
                    merged_phones = merge_phone_sequence(phones, existing_phones) if phones else merge_phone_sequence([], existing_phones)
                    if merged_phones:
                        client.melhor_telefone = merged_phones[0]
                        client.tem_telefone = True
                    # Regra operacional: se o cliente estava bloqueado por "Nao e o cliente" e chegou
                    # telefone novo no retorno, remover bloqueio automaticamente para voltar ao funil.
                    if bool(client.do_not_call) and (client.do_not_call_reason or "") == "nao_e_o_cliente" and has_new_incoming_phone:
                        client.do_not_call = False
                        client.do_not_call_reason = ""
                    updated_any_phone = False
                    updated_any_data = False
                    for occurrence in occurrences:
                        phone_changed, data_changed = apply_enrichment_to_occurrence(occurrence, normalized, phones)
                        updated_any_phone = updated_any_phone or phone_changed
                        updated_any_data = updated_any_data or data_changed

                    if updated_any_data:
                        summary.data_fields_updated += 1
                    if previous_primary != client.melhor_telefone or updated_any_phone or updated_any_data:
                        summary.clients_updated += 1

                    for phone in phones:
                        exists = session.scalar(
                            select(func.count())
                            .select_from(PhoneEnrichment)
                            .where(
                                PhoneEnrichment.cpf == cpf,
                                PhoneEnrichment.telefone == phone,
                                PhoneEnrichment.fonte == source_label,
                            )
                        )
                        if exists:
                            continue
                        session.add(
                            PhoneEnrichment(
                                cpf=cpf,
                                telefone=phone,
                                fonte=source_label,
                                arquivo_origem=file_path.name,
                                importado_por=imported_by or "operador",
                            )
                        )
                        summary.phones_added += 1

            if affected_cpfs:
                rebuild_clients_for_cpfs(session, affected_cpfs)
            session.commit()
    return summary
